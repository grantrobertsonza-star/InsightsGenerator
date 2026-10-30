"""Generators for L2 codebook cases: each file is written FROM a declared truth (a list of Var), never
the other way round. Nothing here reads the app's code."""
import csv
import io
import json

import openpyxl


class Var:
    """One codebook variable as the codebook author meant it: name, label, answer codes and labels (codes
    as written), the section heading it sits under."""

    def __init__(self, name, label="", opts=(), section=None, note=None):
        self.name = name
        self.label = label
        self.opts = [(str(c), str(l)) for c, l in opts]
        self.section = section
        self.note = note  # a values cell that is not an answer list ("as is", "1 thru 58")

    def expect(self, **extra):
        e = {"v": self.name, "label": self.label, "opts": list(self.opts)}
        e.update(extra)
        return e


def values_cell(v, pair="{c}={l}", sep=", "):
    if v.note is not None:
        return v.note
    return sep.join(pair.format(c=c, l=l) for c, l in v.opts)


def wide_rows(truth, header=("Variable", "Label", "Values"), order=(0, 1, 2), pair="{c}={l}", sep=", ", titles=(), sections=True):
    """One row per variable: name, label, answer list in one cell. `order` places the three columns."""
    rows = [list(t) for t in titles]
    if header is not None:
        rows.append([header[i] for i in order])
    current = None
    for v in truth:
        if sections and v.section and v.section != current:
            rows.append([v.section])
            current = v.section
        cells = [v.name, v.label, values_cell(v, pair, sep)]
        rows.append([cells[i] for i in order])
    return rows


def long_rows(truth, header=("Variable", "Variable label", "Value", "Value label"), repeat=True, order=None, titles=()):
    """One row per answer code; a variable with no codes gets one row with an empty code."""
    rows = [list(t) for t in titles]
    if header is not None:
        rows.append(list(header) if order is None else [header[i] for i in order])
    for v in truth:
        opts = v.opts or [("", "")]
        for k, (c, l) in enumerate(opts):
            name = v.name if (repeat or k == 0) else ""
            lab = v.label if (repeat or k == 0) else ""
            cells = [name, lab, c, l] if len(header or [1, 2, 3, 4]) == 4 else [name, c, l]
            rows.append(cells if order is None else [cells[i] for i in order])
    return rows


def csv_bytes(rows, delimiter=",", encoding="utf-8", bom=False, newline="\n"):
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=delimiter, lineterminator=newline, quoting=csv.QUOTE_MINIMAL)
    for r in rows:
        w.writerow(r)
    text = buf.getvalue()
    data = text.encode(encoding)
    if bom:
        data = b"\xef\xbb\xbf" + data
    return data


def xlsx_bytes(sheets, merges=None):
    """sheets: [(sheet name, rows)]; cells may be str, int or float (written as Excel numbers)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append([None if c == "" else c for c in r])
        for rng in (merges or {}).get(name, []):
            ws.merge_cells(rng)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def json_bytes(obj):
    return json.dumps(obj, ensure_ascii=False).encode("utf-8")
