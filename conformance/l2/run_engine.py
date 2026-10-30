"""L2 conformance: how the engine reads a codebook (and an XLSForm questionnaire), case by case.

Run:   cd /home/claude/b66 && python3 conformance/l2/run_engine.py [case-id-prefix ...]
Writes conformance/l2/results/engine.json and prints one line per case.

Every expected value comes from the truth the file is generated from (gen.py) or from a hand reading
(the IGS cases, transcribed from the workbook cell by cell), and from the contracts in claude/. The
app's output is never the expected value.

Case kinds
  contract  the contracts settle the answer; a mismatch is FAIL.
  gap       a shape the contracts do not support. Acceptable: the correct reading, a refusal, or the
            affected items flagged unread. A silent wrong or partial reading reported complete is FAIL
            (standing requirement 1 and 7: unknown never silently becomes empty).
  question  the contracts do not settle it; the hand-read truth is recorded and the actual result is
            reported as QUESTION (agrees / differs), never counted as pass or fail.
  feature   a feature the contracts do not offer (e.g. multiple-response sets in a codebook); reported
            as GAP-supported / GAP-unsupported, listed under feature gaps, never counted as a fault.
  known     a documented planned feature (grid detection, batch 53); the result is recorded as KNOWN.
"""
import importlib.util
import json
import os
import re
import sys
import traceback

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, HERE)
from gen import Var, csv_bytes, json_bytes, long_rows, wide_rows, xlsx_bytes  # noqa: E402

IGS_DIR = "/mnt/user-data/uploads/InsighsAI/TestProjects/IGS"


def load_engine():
    spec = importlib.util.spec_from_file_location("parse_stats_file", os.path.join(ROOT, "api", "parse-stats-file.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# ---------------------------------------------------------------- truths

YN = [("1", "Yes"), ("2", "No")]
BASIC = [
    Var("id", "Respondent identifier", note="as is"),
    Var("q1", "Would you recommend us to a friend?", YN),
    Var("q2", "Overall satisfaction with the service", [("1", "Very satisfied"), ("2", "Satisfied"), ("3", "Dissatisfied"), ("4", "Very dissatisfied"), ("9", "Don't know")]),
    Var("region", "Region of residence", [("1", "North"), ("2", "South"), ("3", "East"), ("4", "West")]),
]


def items_of(truth):
    out = []
    for v in truth:
        e = {"v": v.name, "label": v.label, "opts": list(v.opts), "unread": False}
        out.append(e)
    return out


def sectioned(truth):
    return [dict(e, section=v.section) for e, v in zip(items_of(truth), truth)]


def I(spec, strict=True, sheets=None, skipped=None):  # noqa: E743
    return ("items", spec, {"strict": strict, "sheets": sheets, "skipped": skipped})


def R(pattern):
    return ("refused", pattern)


def F(names):
    return ("flagged", names)


CASES = []


def case(cid, title, kind, build, expect, clause="", severity="S2", note=""):
    CASES.append({"id": cid, "title": title, "kind": kind, "build": build, "expect": expect if isinstance(expect, list) else [expect], "clause": clause, "severity": severity, "note": note})


CB = "claude/codebook-import.md"

# ---- formats: one truth, every supported container (cross-format agreement)
case("E001", "wide codebook, CSV, standard header", "contract", lambda: (csv_bytes(wide_rows(BASIC)), "cb.csv"), I(items_of(BASIC)), CB + " Reading: CSV read by the engine")
case("E002", "wide codebook, XLSX", "contract", lambda: (xlsx_bytes([("Codebook", wide_rows(BASIC))]), "cb.xlsx"), I(items_of(BASIC), sheets=["Codebook"]), CB + " Reading: Excel")
case("E003", "wide codebook, TSV", "contract", lambda: (csv_bytes(wide_rows(BASIC), "\t"), "cb.tsv"), I(items_of(BASIC)), CB + " Reading: TSV")
case("E004", "wide codebook, tab-delimited .txt", "contract", lambda: (csv_bytes(wide_rows(BASIC), "\t"), "cb.txt"), I(items_of(BASIC)), CB + " Reading: delimited text")
case("E005", "wide codebook, semicolon-delimited .csv (answer lists hold commas)", "contract", lambda: (csv_bytes(wide_rows(BASIC), ";"), "cb.csv"), I(items_of(BASIC)), "claude/one-reader-import.md Delimiter: documents take the candidate that splits the most rows")
case("E006", "wide codebook, pipe-delimited .txt", "contract", lambda: (csv_bytes(wide_rows(BASIC), "|"), "cb.txt"), I(items_of(BASIC)), "claude/one-reader-import.md Delimiter")
case("E007", "wide codebook, CSV with UTF-8 byte-order mark", "contract", lambda: (csv_bytes(wide_rows(BASIC), bom=True), "cb.csv"), I(items_of(BASIC)), "claude/one-reader-import.md Encoding")
case("E008", "wide codebook, CSV in UTF-16 with BOM", "contract", lambda: (csv_bytes(wide_rows(BASIC), encoding="utf-16"), "cb.csv"), I(items_of(BASIC)), "claude/one-reader-import.md Encoding")
W1252 = [Var("q1", "Café visit", [("1", "Don’t know"), ("2", "Señor")])]
case("E009", "wide codebook, Windows-1252 CSV (curly apostrophe, accents)", "contract", lambda: (csv_bytes(wide_rows(W1252), encoding="cp1252"), "cb.csv"), I(items_of(W1252)), "claude/one-reader-import.md Encoding")
case("E010", "wide codebook, CSV with CRLF line ends", "contract", lambda: (csv_bytes(wide_rows(BASIC), newline="\r\n"), "cb.csv"), I(items_of(BASIC)), "claude/one-reader-import.md Quoting")

# ---- JSON shapes
def jlist():
    return [{"name": v.name, "label": v.label, "values": {c: l for c, l in v.opts}} for v in BASIC if v.opts] + [{"name": "id", "label": "Respondent identifier"}]


JBASIC = [e for e in items_of(BASIC) if e["opts"]] + [{"v": "id", "label": "Respondent identifier", "opts": [], "unread": False}]
case("E011", "JSON list of variable objects, values as {code: label}", "contract", lambda: (json_bytes(jlist()), "cb.json"), I(JBASIC), CB + " JSON: a list of variable objects")
case("E012", "JSON {variables: [...]}, values as [{value, label}] with numeric values", "contract",
     lambda: (json_bytes({"variables": [{"variable": v.name, "description": v.label, "value_labels": [{"value": int(c), "label": l} for c, l in v.opts]} for v in BASIC if v.opts]}), "cb.json"),
     I([e for e in items_of(BASIC) if e["opts"]]), CB + " JSON: values as [{value, label}]")
case("E013", "JSON object keyed by variable name, values as [[code, label]]", "contract",
     lambda: (json_bytes({v.name: {"label": v.label, "values": [[int(c), l] for c, l in v.opts]} for v in BASIC if v.opts}), "cb.json"),
     I([e for e in items_of(BASIC) if e["opts"]]), CB + " JSON: object keyed by variable name")
case("E014", "JSON Frictionless Table Schema (fields, title, categories)", "contract",
     lambda: (json_bytes({"fields": [{"name": v.name, "title": v.label, "categories": [{"value": int(c), "label": l} for c, l in v.opts]} for v in BASIC if v.opts]}), "cb.json"),
     I([e for e in items_of(BASIC) if e["opts"]]), CB + " JSON: fields (Frictionless)")
case("E015", "JSON keyed by name with a plain label string", "contract", lambda: (json_bytes({"q1": "Would you recommend us?", "age": "Age in years"}), "cb.json"),
     I([{"v": "q1", "label": "Would you recommend us?", "opts": []}, {"v": "age", "label": "Age in years", "opts": []}]), CB + " JSON: or a label")
case("E016", "JSON codes '01' and '1' are two codes", "contract", lambda: (json_bytes([{"name": "q1", "label": "L", "values": {"01": "Refused", "1": "Yes"}}]), "cb.json"),
     I([{"v": "q1", "opts": [("01", "Refused"), ("1", "Yes")]}]), CB + " IS-F02/B36-F01 codes as written")
case("E017", "JSON large integer code 9007199254740993 not rounded", "contract", lambda: (b'[{"name":"q1","label":"L","values":[{"value":9007199254740993,"label":"Big"}]}]', "cb.json"),
     I([{"v": "q1", "opts": [("9007199254740993", "Big")]}]), CB + " IS-F02 large codes")
case("E018", "JSON numeric code 1.0 is the code 1 (a number, not text)", "contract", lambda: (b'[{"name":"q1","label":"L","values":[{"value":1.0,"label":"Yes"},{"value":2,"label":"No"}]}]', "cb.json"),
     I([{"v": "q1", "opts": [("1", "Yes"), ("2", "No")]}]), CB + " _cb_code: a cell that is a number is written as that number")
case("E019", "JSON negative and decimal codes", "contract", lambda: (b'[{"name":"q1","label":"L","values":{"-1":"Not asked","1.5":"Half"}}]', "cb.json"),
     I([{"v": "q1", "opts": [("-1", "Not asked"), ("1.5", "Half")]}]), CB + " codes as written")
case("E020", "JSON that is not JSON is refused", "contract", lambda: (b'{"variables": [', "cb.json"), R(r"could not be read"), CB + " A file with no codebook table is refused")
case("E021", "JSON with no variables is refused, naming what was expected", "contract", lambda: (b'{"meta": [1, 2, 3]}', "cb.json"), R(r"No variables were found"), CB + " refused with what was expected")
case("E022", "JSON duplicate variable: first description wins", "contract", lambda: (json_bytes([{"name": "q1", "label": "First", "values": {"1": "A"}}, {"name": "q1", "label": "Second", "values": {"1": "B"}}]), "cb.json"),
     I([{"v": "q1", "label": "First", "opts": [("1", "A")]}]), CB + " The first description of a variable wins")
case("E023", "JSON section/group key kept as section", "contract", lambda: (json_bytes([{"name": "q1", "label": "L", "group": "Demographics"}]), "cb.json"),
     I([{"v": "q1", "section": "Demographics"}]), CB + " JSON shapes (section)")

# ---- long layout (one row per value)
LONG = [
    Var("q1", "Would you recommend us?", YN),
    Var("q2", "Satisfaction", [("1", "Very satisfied"), ("2", "Satisfied"), ("98", "Don't know"), ("99", "Refused")]),
    Var("age", "Age in years"),
]
case("E030", "long codebook, variable repeated on each row (CSV)", "contract", lambda: (csv_bytes(long_rows(LONG)), "cb.csv"), I(items_of(LONG)), CB + " One row per value ... read by its header names")
case("E031", "long codebook, variable written once then blank (CSV)", "contract", lambda: (csv_bytes(long_rows(LONG, repeat=False)), "cb.csv"), I(items_of(LONG)), CB + " with the variable repeated or blank")
case("E032", "long codebook, XLSX with numeric code cells (1, 2, 98, -1, 1.5)", "contract",
     lambda: (xlsx_bytes([("Values", [["Variable", "Variable label", "Value", "Value label"], ["q1", "L", 1, "Yes"], ["q1", "L", 2, "No"], ["q1", "L", 98, "Don't know"], ["q1", "L", -1, "Not asked"], ["q1", "L", 1.5, "Half"]])]), "cb.xlsx"),
     I([{"v": "q1", "label": "L", "opts": [("1", "Yes"), ("2", "No"), ("98", "Don't know"), ("-1", "Not asked"), ("1.5", "Half")]}]), CB + " codes as written; an Excel number cell is that number")
case("E033", "long codebook, header Name/Code/Label (Label is the value label)", "contract",
     lambda: (csv_bytes([["Name", "Code", "Label"], ["q1", "1", "Yes"], ["q1", "2", "No"], ["q2", "1", "Male"], ["q2", "2", "Female"]]), "cb.csv"),
     I([{"v": "q1", "opts": YN}, {"v": "q2", "opts": [("1", "Male"), ("2", "Female")]}]), CB + " long layout by header names")
case("E034", "long codebook, leading-zero codes as text", "contract",
     lambda: (csv_bytes([["Variable", "Value", "Value label"], ["q1", "01", "Refused"], ["q1", "1", "Yes"], ["q1", "001", "Other"]]), "cb.csv"),
     I([{"v": "q1", "opts": [("01", "Refused"), ("1", "Yes"), ("001", "Other")]}]), CB + " B36-F01 codes as written")
case("E035", "long codebook under two title rows", "contract", lambda: (csv_bytes(long_rows(LONG, titles=[["Brand tracker codebook"], ["Wave 3, May 2026"]])), "cb.csv"), I(items_of(LONG)), CB + " IS-F07 title rows never variables")
case("E036", "long codebook, columns in a different order (Value label, Value, Variable label, Variable)", "contract",
     lambda: (csv_bytes(long_rows(LONG, order=(3, 2, 1, 0))), "cb.csv"), I(items_of(LONG)), CB + " read by header names, not position")
case("E037", "long codebook with no header row", "gap",
     lambda: (csv_bytes(long_rows(LONG, header=None)), "cb.csv"), [I(items_of(LONG)), R(r".")], CB + " (headerless long layout not described)")
case("E038", "long codebook, header 'Variable Name, Value Code, Value Label'", "gap",
     lambda: (csv_bytes([["Variable Name", "Value Code", "Value Label"], ["q1", "1", "Yes"], ["q1", "2", "No"], ["q2", "1", "Male"], ["q2", "2", "Female"], ["q3", "1", "Agree"], ["q3", "2", "Disagree"]]), "cb.csv"),
     [I([{"v": "q1", "opts": YN}, {"v": "q2", "opts": [("1", "Male"), ("2", "Female")]}, {"v": "q3", "opts": [("1", "Agree"), ("2", "Disagree")]}]), R(r".")],
     CB + " long layout read by header names; requirement 1: unknown never silently empty")
case("E039", "long codebook in XLSX with the variable cell merged down its value rows", "contract",
     lambda: (xlsx_bytes([("CB", [["Variable", "Variable label", "Value", "Value label"], ["q1", "Recommend", 1, "Yes"], ["", "", 2, "No"], ["q2", "Gender", 1, "Male"], ["", "", 2, "Female"]])], merges={"CB": ["A2:A3", "B2:B3", "A4:A5", "B4:B5"]}), "cb.xlsx"),
     I([{"v": "q1", "label": "Recommend", "opts": YN}, {"v": "q2", "label": "Gender", "opts": [("1", "Male"), ("2", "Female")]}]), CB + " variable repeated or blank")
case("E040", "long codebook, SPSS DISPLAY DICTIONARY 'Variable Values' export (blank first header cell)", "gap",
     lambda: (xlsx_bytes([("Variable Values", [["", "Value", "Label"], ["q1", 1, "Yes"], ["", 2, "No"], ["q2", 1, "Male"], ["", 2, "Female"], ["q3", 1, "Agree"], ["", 2, "Disagree"]])]), "cb.xlsx"),
     [I([{"v": "q1", "opts": YN}, {"v": "q2", "opts": [("1", "Male"), ("2", "Female")]}, {"v": "q3", "opts": [("1", "Agree"), ("2", "Disagree")]}], strict=False), R(r".")],
     "requirement 1 and 7; common SPSS export shape")
case("E041", "long codebook, variable listed with no codes (empty value cell) keeps no options", "contract",
     lambda: (csv_bytes([["Variable", "Variable label", "Value", "Value label"], ["age", "Age in years", "", ""], ["q1", "Recommend", "1", "Yes"], ["q1", "Recommend", "2", "No"]]), "cb.csv"),
     I([{"v": "age", "label": "Age in years", "opts": []}, {"v": "q1", "label": "Recommend", "opts": YN}]), CB + " long layout")

# ---- wide layout: headers and titles
case("E050", "wide codebook with no header row (3+ rows, name column by content)", "contract", lambda: (csv_bytes(wide_rows(BASIC, header=None)), "cb.csv"), I(items_of(BASIC)), CB + " Without such a header, the column of name-like cells")
TWO = BASIC[1:3]
case("E051", "wide codebook with no header and only two rows is refused", "contract", lambda: (csv_bytes(wide_rows(TWO, header=None)), "cb.csv"), R(r"No codebook table was found"), CB + " three or more rows are needed without a header")
case("E052", "one variable under a 'VARIABLE NAME, VARIABLE LABEL' header is read", "contract",
     lambda: (csv_bytes([["VARIABLE NAME", "VARIABLE LABEL", "VALUES"], ["q_2028", "Satisfaction", "1=Low, 2=High"]]), "cb.csv"),
     I([{"v": "q_2028", "label": "Satisfaction", "opts": [("1", "Low"), ("2", "High")]}]), CB + " B36-F05 / IS-F07")
case("E053", "two title rows above the header are not variables", "contract",
     lambda: (csv_bytes(wide_rows(BASIC, titles=[["Customer survey 2026: codebook"], ["Prepared by Fieldwork Ltd", "28 October 2026"]])), "cb.csv"), I(items_of(BASIC)), CB + " IS-F07")
case("E054", "seven title rows above the header: header and titles are not variables", "contract",
     lambda: (csv_bytes(wide_rows(BASIC, titles=[["Customer survey 2026"], ["Codebook"], ["Version 3"], ["Prepared by Fieldwork Ltd"], ["Contact: data team"], ["Confidential"], ["See notes sheet"]])), "cb.csv"),
     I(items_of(BASIC)), CB + " IS-F07: the header row, and any title rows above it, are never variables")
case("E055", "header written the wrong way round (IGS: 'VARIABLE LABEL' over the names)", "contract",
     lambda: (csv_bytes(wide_rows(BASIC, header=("VARIABLE LABEL", "VARIABLE NAME", "VALUES"))), "cb.csv"), I(items_of(BASIC)), CB + " IS-F05 columns")
case("E056", "columns in another order: Label, Values, Variable", "contract",
     lambda: (csv_bytes(wide_rows(BASIC, order=(1, 2, 0))), "cb.csv"), I(items_of(BASIC)), CB + " the names column is the one a header row names")
case("E057", "header 'Name, Description, Codes'", "contract",
     lambda: (csv_bytes(wide_rows(BASIC, header=("Name", "Description", "Codes"))), "cb.csv"), I(items_of(BASIC)), CB + " columns by header")
case("E058", "header 'Question ID, Question text, Response options'", "contract",
     lambda: (csv_bytes(wide_rows(BASIC, header=("Question ID", "Question text", "Response options"))), "cb.csv"), I(items_of(BASIC)), CB + " columns by header")
case("E059", "unrecognised header 'Varname, Description, Value labels' is not read as a variable", "contract",
     lambda: (csv_bytes(wide_rows(BASIC, header=("Varname", "Description", "Value labels"))), "cb.csv"), I(items_of(BASIC)), CB + " IS-F07: the header row is never a variable")
case("E060", "unrecognised header 'Var, Lab, Vals' is not read as a variable", "contract",
     lambda: (csv_bytes(wide_rows(BASIC, header=("Var", "Lab", "Vals"))), "cb.csv"), I(items_of(BASIC)), CB + " IS-F07")
SEC = [Var("q1", "Recommend", YN, section="Section A: Service"), Var("q2", "Satisfied", YN, section="Section A: Service"), Var("age", "Age in years", section="Section B: About you"), Var("sex", "Gender", [("1", "Male"), ("2", "Female")], section="Section B: About you")]
case("E061", "section heading rows (one filled cell, not name-like) become sections", "contract", lambda: (csv_bytes(wide_rows(SEC)), "cb.csv"), I(sectioned(SEC)), CB + " A row with one filled cell is a section heading")
case("E062", "a one-word upper-case section heading ('DEMOGRAPHICS') between variables", "question",
     lambda: (csv_bytes([["Variable", "Label", "Values"], ["q1", "Recommend", "1=Yes, 2=No"], ["DEMOGRAPHICS"], ["age", "Age in years", "as is"], ["sex", "Gender", "1=Male, 2=Female"]]), "cb.csv"),
     I([{"v": "q1"}, {"v": "age", "section": "DEMOGRAPHICS"}, {"v": "sex", "section": "DEMOGRAPHICS"}]), CB + " 'One filled cell that is not name-like is a section heading' (silent on name-like headings)", "S3")
case("E063", "one name cell listing several names 'white, black, asian' gives each the row's label and values", "contract",
     lambda: (csv_bytes([["Variable", "Label", "Values"], ["q1", "Recommend", "1=Yes, 2=No"], ["white, black, asian", "Race (summary)", "0=No, 1=Yes"], ["age", "Age", "as is"]]), "cb.csv"),
     I([{"v": "q1"}, {"v": "white", "label": "Race (summary)", "opts": [("0", "No"), ("1", "Yes")]}, {"v": "black", "label": "Race (summary)", "opts": [("0", "No"), ("1", "Yes")]}, {"v": "asian", "label": "Race (summary)", "opts": [("0", "No"), ("1", "Yes")]}, {"v": "age"}]),
     CB + " a name cell listing several names gives each the row's label and values")
case("E064", "a header naming the label column takes it even when a notes column is longer", "contract",
     lambda: (csv_bytes([["Variable", "Label", "Values", "Notes"], ["q1", "Recommend", "1=Yes, 2=No", "Asked of all respondents who completed the service module in wave 3 and wave 4"], ["q2", "Satisfied", "1=Yes, 2=No", "Asked of all respondents who completed the service module in wave 3 and wave 4"], ["q3", "Return", "1=Yes, 2=No", "Asked of all respondents who completed the service module in wave 3 and wave 4"]]), "cb.csv"),
     I([{"v": "q1", "label": "Recommend", "opts": YN}, {"v": "q2", "label": "Satisfied", "opts": YN}, {"v": "q3", "label": "Return", "opts": YN}]), CB + " The label column is the header's")

# ---- answer lists in one cell
def one(values, opts=None, unread=False, cid=None, title=None, kind="contract", clause=CB + " Answer lists (_cb_list)", alts=None, sev="S2"):
    rows = [["Variable", "Label", "Values"], ["q1", "Question one", values], ["q2", "Question two", "1=Yes, 2=No"], ["age", "Age", "as is"]]
    exp = []
    if opts is not None:
        exp.append(I([{"v": "q1", "opts": opts, "unread": False}, {"v": "q2", "opts": YN}, {"v": "age", "opts": []}]))
    if unread:
        exp.append(F(["q1"]))
    case(cid, title, kind, lambda: (csv_bytes(rows), "cb.csv"), exp + (alts or []), clause, sev)


one("1=Yes, 2=No, 9=Don't know", YN + [("9", "Don't know")], cid="E070", title="'1=Yes, 2=No, 9=Don't know' (apostrophe in label)")
one("1 = Yes , 2 = No", YN, cid="E071", title="spaces around '=' and before commas")
one("1=Yes,2=No", YN, cid="E072", title="no space after commas")
one("1=Yes, 2=No,", YN, cid="E073", title="trailing comma")
one("1=Yes; 2=No", YN, unread=True, cid="E074", title="'1=Yes; 2=No' semicolon-separated list", kind="gap", clause=CB + " unread lists; semicolons not described")
one("1=Yes\n2=No", YN, unread=True, cid="E075", title="one code per line inside the cell ('1=Yes' newline '2=No')", kind="gap", clause=CB + " unread lists; line-separated lists not described")
one("1=Yes 2=No", None, unread=True, cid="E076", title="'1=Yes 2=No' without separators is an unread list", clause=CB + " IS-F04 '1=Yes 2=No' is an unread list")
one("=Yes", None, unread=True, cid="E077", title="'=Yes' is an unread list", clause=CB + " IS-F04")
one('"A"=Agree, "D"=Disagree', [("A", "Agree"), ("D", "Disagree")], cid="E078", title="double-quoted text codes")
one("'A'=Agree, 'D'=Disagree", [("A", "Agree"), ("D", "Disagree")], cid="E079", title="single-quoted text codes")
one('1="Yes, often", 2="No, never"', [("1", "Yes, often"), ("2", "No, never")], cid="E080", title="quoted labels holding commas")
one("01=Yes, 02=No, 10=Other", [("01", "Yes"), ("02", "No"), ("10", "Other")], cid="E081", title="leading-zero codes kept as written")
one("01=Refused, 1=Yes", [("01", "Refused"), ("1", "Yes")], cid="E082", title="'01' and '1' are two codes", clause=CB + " IS-F02")
one("9007199254740993=Big, 1=Small", [("9007199254740993", "Big"), ("1", "Small")], cid="E083", title="large integer code not rounded", clause=CB + " IS-F02")
one("-1=Not asked, -9=Refused, 1=Yes", [("-1", "Not asked"), ("-9", "Refused"), ("1", "Yes")], cid="E084", title="negative codes")
one("0.5=Half, 1.5=One and a half", [("0.5", "Half"), ("1.5", "One and a half")], cid="E085", title="decimal codes")
one("1e0=One, 2=Two", [("1e0", "One"), ("2", "Two")], cid="E086", title="exponent spelling kept as its own code", clause=CB + " IS-F04 '1e0'")
one("1=Yes, 2=No, 98=Don't know, 99=Refused", YN + [("98", "Don't know"), ("99", "Refused")], cid="E087", title="missing/non-answer codes 98, 99 kept as answer codes")
one("DK=Don't know, RF=Refused, Y=Yes", [("DK", "Don't know"), ("RF", "Refused"), ("Y", "Yes")], cid="E088", title="text codes DK, RF, Y")
one("A_B=Yes, C-D=No", [("A_B", "Yes"), ("C-D", "No")], cid="E089", title="text codes with _ and -", clause=CB + " IS-F04")
one("1=Yes, I agree, 2=No, I don't", [("1", "Yes, I agree"), ("2", "No, I don't")], cid="E090", title="unquoted labels holding commas and an apostrophe")
one("1=Score >= 5, 2=Below 5", [("1", "Score >= 5"), ("2", "Below 5")], unread=True, cid="E091", title="label holding '=' ('Score >= 5')", kind="gap", clause=CB + " a label may hold a comma but not ', <code>='")
one("as is", [], cid="E092", title="'as is' is not a list (no options, not unread)", clause=CB + " Text without '=' ('as is') is not a list")
one("1 thru 58", [], cid="E093", title="'1 thru 58' gives no answer list and is not unread", clause=CB + " values written another way ('1 thru 58') give no answer list")
one("0-7", [], cid="E094", title="range '0-7' gives no answer list", clause=CB + " values written another way")
one("{1, Yes}, {2, No}", YN, cid="E095", title="SPSS Variable View style '{1, Yes}, {2, No}'", kind="gap", clause="requirement 1: unknown never silently empty",
    alts=[F(["q1"]), R(r".")])
one("1. Yes, 2. No", YN, cid="E096", title="'1. Yes, 2. No' (dot instead of '=')", kind="gap", clause="requirement 1", alts=[F(["q1"]), R(r".")])
one("1: Yes, 2: No", YN, cid="E097", title="'1: Yes, 2: No' (colon instead of '=')", kind="gap", clause="requirement 1", alts=[F(["q1"]), R(r".")])
one("1 Yes, 2 No", YN, cid="E098", title="'1 Yes, 2 No' (no separator)", kind="gap", clause="requirement 1", alts=[F(["q1"]), R(r".")], sev="S3")

# ---- variable names
NAMES = [Var("Age years", "Age in years"), Var("q1.a", "Part a", YN), Var("_Q1", "Underscore", YN), Var("回答", "Answer", YN), Var("123", "Numeric name", YN), Var("Q 1", "Spaced", YN)]
case("E100", "names with spaces, dots, underscore, CJK, digits kept as written", "contract", lambda: (csv_bytes(wide_rows(NAMES)), "cb.csv"), I(items_of(NAMES)), CB + " IS-F05 names kept exactly as written")
case("E101", "names with leading/trailing spaces are trimmed as the data reader trims", "contract",
     lambda: (csv_bytes([["Variable", "Label", "Values"], [" q1 ", "One", "1=Yes, 2=No"], ["q2  ", "Two", "1=Yes, 2=No"], ["q3", "Three", "1=Yes, 2=No"]]), "cb.csv"),
     I([{"v": "q1"}, {"v": "q2"}, {"v": "q3"}]), "claude/one-reader-import.md Names: trimmed")
case("E102", "duplicate name in one sheet: first description wins", "contract",
     lambda: (csv_bytes([["Variable", "Label", "Values"], ["q1", "First", "1=A, 2=B"], ["q2", "Two", "1=Yes, 2=No"], ["q1", "Second", "1=C, 2=D"]]), "cb.csv"),
     I([{"v": "q1", "label": "First", "opts": [("1", "A"), ("2", "B")]}, {"v": "q2"}]), CB + " The first description of a variable wins")
case("E103", "names differing only in case (Q1, q1) are two variables", "contract",
     lambda: (csv_bytes([["Variable", "Label", "Values"], ["Q1", "Upper", "1=Yes, 2=No"], ["q1", "Lower", "1=Yes, 2=No"], ["q2", "Two", "1=Yes, 2=No"]]), "cb.csv"),
     I([{"v": "Q1", "label": "Upper"}, {"v": "q1", "label": "Lower"}, {"v": "q2"}]), CB + " names kept exactly as written")
case("E104", "XLSX numeric name cell 123 is the name '123'", "contract",
     lambda: (xlsx_bytes([("CB", [["Variable", "Label", "Values"], [123, "Numeric name", "1=Yes, 2=No"], ["q2", "Two", "1=Yes, 2=No"]])]), "cb.xlsx"),
     I([{"v": "123", "label": "Numeric name", "opts": YN}, {"v": "q2"}]), "claude/one-reader-import.md Excel number cells as names are written without '.0'")
case("E105", "a range in a name cell 'Q5_1 to Q5_8'", "question",
     lambda: (csv_bytes([["Variable", "Label", "Values"], ["Q4", "Four", "1=Yes, 2=No"], ["Q5_1 to Q5_8", "Opinion of candidates", "1=Favourable, 2=Unfavourable"], ["Q6", "Six", "1=Yes, 2=No"]]), "cb.csv"),
     I([{"v": "Q4"}] + [{"v": f"Q5_{k}", "opts": [("1", "Favourable"), ("2", "Unfavourable")]} for k in range(1, 9)] + [{"v": "Q6"}]), CB + " names kept as written (IS-F05) vs. a range of names", "S2")

# ---- several sheets
VARS_SHEET = [["Variable", "Label"], ["q1", "Would you recommend us?"], ["q2", "Gender"], ["age", "Age in years"]]
VALS_SHEET = [["Variable", "Value", "Value label"], ["q1", "1", "Yes"], ["q1", "2", "No"], ["q2", "1", "Male"], ["q2", "2", "Female"]]
SPLIT_TRUTH = [{"v": "q1", "label": "Would you recommend us?", "opts": YN}, {"v": "q2", "label": "Gender", "opts": [("1", "Male"), ("2", "Female")]}, {"v": "age", "label": "Age in years", "opts": []}]
case("E110", "variables sheet then values sheet: labels from one, answer codes from the other", "question",
     lambda: (xlsx_bytes([("Variables", VARS_SHEET), ("Values", VALS_SHEET)]), "cb.xlsx"), I(SPLIT_TRUTH, sheets=["Variables", "Values"]),
     CB + " 'The first description of a variable wins' vs. a two-sheet dictionary", "S2")
case("E111", "values sheet first, then variables sheet", "question",
     lambda: (xlsx_bytes([("Values", VALS_SHEET), ("Variables", VARS_SHEET)]), "cb.xlsx"), I([SPLIT_TRUTH[0], SPLIT_TRUTH[1], SPLIT_TRUTH[2]], strict=False),
     CB + " first description wins", "S2")
case("E112", "a notes sheet of prose beside the codebook sheet is skipped and named", "contract",
     lambda: (xlsx_bytes([("Notes", [["This codebook describes the 2026 customer survey."], ["Fieldwork ran from 1 to 20 May."]]), ("Codebook", wide_rows(BASIC))]), "cb.xlsx"),
     I(items_of(BASIC), sheets=["Codebook"], skipped=["Notes"]), CB + " sheets that are not codebook tables are named in skippedSheets")
case("E113", "the same variable on two sheets: first sheet's description wins", "contract",
     lambda: (xlsx_bytes([("A", [["Variable", "Label", "Values"], ["q1", "From A", "1=Yes, 2=No"], ["q2", "Two", "1=Yes, 2=No"]]), ("B", [["Variable", "Label", "Values"], ["q1", "From B", "1=Oui, 2=Non"], ["q3", "Three", "1=Yes, 2=No"]])]), "cb.xlsx"),
     I([{"v": "q1", "label": "From A", "opts": YN}, {"v": "q2"}, {"v": "q3"}]), CB + " The first description of a variable wins")
case("E114", "a study-information sheet of key/value rows (Project, Fieldwork, Sample)", "question",
     lambda: (xlsx_bytes([("Study", [["Project", "Customer survey 2026"], ["Fieldwork", "1 to 20 May 2026"], ["Sample", "1,000 adults"], ["Method", "Online panel"]]), ("Codebook", wide_rows(BASIC))]), "cb.xlsx"),
     I(items_of(BASIC), sheets=["Codebook"], skipped=["Study"]), CB + " three rows without a header make a codebook", "S3")
case("E115", "SPSS DISPLAY DICTIONARY export: 'Variable Information' + 'Variable Values' sheets", "gap",
     lambda: (xlsx_bytes([("Variable Information", [["Variable", "Position", "Label", "Measurement Level", "Column Width", "Alignment", "Print Format", "Write Format"], ["q1", 1, "Would you recommend us?", "Nominal", 8, "Right", "F1", "F1"], ["q2", 2, "Gender", "Nominal", 8, "Right", "F1", "F1"], ["age", 3, "Age in years", "Scale", 8, "Right", "F3", "F3"]]),
                           ("Variable Values", [["Value", "", "Label"], ["q1", 1, "Yes"], ["", 2, "No"], ["q2", 1, "Male"], ["", 2, "Female"]])]), "cb.xlsx"),
     [I(SPLIT_TRUTH, strict=False), R(r".")], "requirement 1 and 7; SPSS export shape", "S2")

# ---- multiple response, missing, grids, values across columns
MR = [["Variable", "Label", "Values", "Multiple response set"], ["Q13_1", "Sources: Local TV", "1=Selected", "$Q13"], ["Q13_2", "Sources: Newspapers", "1=Selected", "$Q13"], ["Q13_3", "Sources: Podcasts", "1=Selected", "$Q13"], ["Q14", "Direction of the country", "1=Right, 2=Wrong", ""]]
case("E120", "multiple-response set column: variables read correctly", "contract", lambda: (csv_bytes(MR), "cb.csv"),
     I([{"v": "Q13_1", "label": "Sources: Local TV", "opts": [("1", "Selected")]}, {"v": "Q13_2", "opts": [("1", "Selected")]}, {"v": "Q13_3", "opts": [("1", "Selected")]}, {"v": "Q14", "opts": [("1", "Right"), ("2", "Wrong")]}]), CB + " Reading")
case("E121", "multiple-response set column: the set ($Q13 = Q13_1..Q13_3) is kept on the items", "feature", lambda: (csv_bytes(MR), "cb.csv"),
     [("field", "mrSet", {"Q13_1": "$Q13", "Q13_2": "$Q13", "Q13_3": "$Q13"})], "feature gap: multiple-response sets in a codebook", "S3")
case("E122", "multiple-response sets sheet ('$Q13', 'Q13_1 Q13_2 Q13_3') is not read as variables", "contract",
     lambda: (xlsx_bytes([("Codebook", MR[:1] + [r[:3] for r in MR[1:]]), ("MR sets", [["Set", "Label", "Variables"], ["$Q13", "News sources", "Q13_1 Q13_2 Q13_3"], ["$Q45", "Hardship", "Q45_1 Q45_2"], ["$Q47", "Cutbacks", "Q47_1 Q47_2"]])]), "cb.xlsx"),
     I([{"v": "Q13_1"}, {"v": "Q13_2"}, {"v": "Q13_3"}, {"v": "Q14"}], sheets=["Codebook"], skipped=["MR sets"]), CB + " sheets that are not codebook tables are skipped")
MISS = [["Variable", "Label", "Values", "Missing values"], ["q1", "Recommend", "1=Yes, 2=No, 8=Don't know, 9=Refused", "8, 9"], ["q2", "Income band", "1=Low, 2=High, -1=Not asked", "-1"], ["q3", "Return", "1=Yes, 2=No", ""]]
case("E123", "missing-values column: answers still read with their codes", "contract", lambda: (csv_bytes(MISS), "cb.csv"),
     I([{"v": "q1", "opts": [("1", "Yes"), ("2", "No"), ("8", "Don't know"), ("9", "Refused")]}, {"v": "q2", "opts": [("1", "Low"), ("2", "High"), ("-1", "Not asked")]}, {"v": "q3", "opts": YN}]), CB + " Reading")
case("E124", "missing-values column: declared missing codes are kept on the items", "feature", lambda: (csv_bytes(MISS), "cb.csv"),
     [("field", "missing", {"q1": ["8", "9"], "q2": ["-1"]})], "feature gap: declared missing codes in a codebook", "S3")
GRID = [Var(f"Q5_{k}", f"Q5_{k} Is your opinion of each candidate favorable or unfavorable? - {n}", [("1", "Favorable"), ("2", "Unfavorable"), ("3", "Don't know")]) for k, n in enumerate(["Thurmond", "Villaraigosa", "Porter", "Becerra", "Mahan", "Steyer", "Bianco", "Hilton"], 1)]
case("E125", "grid Q5_1..Q5_8 sharing a stem: eight items, each with its own label and codes", "contract", lambda: (csv_bytes(wide_rows(GRID)), "cb.csv"), I(items_of(GRID)), CB + " each variable is an item")
case("E126", "grid Q5_1..Q5_8 grouped as one grid question", "known", lambda: (csv_bytes(wide_rows(GRID)), "cb.csv"), [("field", "grid", {f"Q5_{k}": "Q5" for k in range(1, 9)})], "planned batch 53 (grid detection)", "S3")
case("E127", "values across columns: Variable, Label, 1, 2, 9 (codes as column headers)", "gap",
     lambda: (csv_bytes([["Variable", "Label", "1", "2", "9"], ["q1", "Recommend", "Yes", "No", "Don't know"], ["q2", "Satisfied", "Yes", "No", "Don't know"], ["q3", "Return", "Yes", "No", "Don't know"]]), "cb.csv"),
     [I([{"v": f"q{k}", "opts": [("1", "Yes"), ("2", "No"), ("9", "Don't know")]} for k in (1, 2, 3)], strict=False), R(r"."), F(["q1", "q2", "q3"])], "requirement 1 and 7 (values across columns not described)")
case("E128", "values across columns in code/label pairs: Code1, Label1, Code2, Label2", "gap",
     lambda: (csv_bytes([["Variable", "Label", "Code1", "Label1", "Code2", "Label2"], ["q1", "Recommend", "1", "Yes", "2", "No"], ["q2", "Satisfied", "1", "Yes", "2", "No"], ["q3", "Return", "1", "Yes", "2", "No"]]), "cb.csv"),
     [I([{"v": f"q{k}", "opts": YN} for k in (1, 2, 3)], strict=False), R(r"."), F(["q1", "q2", "q3"])], "requirement 1 and 7")
case("E129", "answer codes one per row under a wide variable row (Excel 'indented' codebook)", "gap",
     lambda: (csv_bytes([["Variable", "Label", "Values"], ["q1", "Recommend", ""], ["", "", "1=Yes"], ["", "", "2=No"], ["q2", "Gender", ""], ["", "", "1=Male"], ["", "", "2=Female"], ["q3", "Return", "1=Yes, 2=No"]]), "cb.csv"),
     [I([{"v": "q1", "opts": YN}, {"v": "q2", "opts": [("1", "Male"), ("2", "Female")]}, {"v": "q3", "opts": YN}]), R(r"."), F(["q1", "q2"])], "requirement 1 and 7")

# ---- refusals
case("E130", "empty CSV is refused", "contract", lambda: (b"", "cb.csv"), R(r"No codebook table was found"), CB + " A file with no codebook table is refused")
case("E131", "header-only CSV is refused", "contract", lambda: (b"Variable,Label,Values\n", "cb.csv"), R(r"No codebook table was found"), CB)
case("E132", "a questionnaire's prose saved as CSV is refused", "contract",
     lambda: (b"Q1. Would you recommend us to a friend?\nYes\nNo\nQ2. How satisfied are you?\nVery\nNot at all\n", "cb.csv"), R(r"No codebook table was found"), CB)
case("E133", "CSV with an unclosed quotation mark is refused with where", "contract",
     lambda: (b'Variable,Label,Values\nq1,"Recommend,1=Yes\nq2,Two,1=Yes\nq3,Three,1=No\n', "cb.csv"), R(r"quoting is not valid"), CB + " B36-F03 (one-reader quoting)")
case("E134", "a PDF sent to the table codebook reader is refused naming the type", "contract", lambda: (b"%PDF-1.4\n", "cb.pdf"), R(r"not \.pdf files"), CB + " Word/PDF codebooks are read by the model")
case("E135", "bytes that are neither UTF-8, UTF-16 nor Windows-1252 text are refused", "contract", lambda: (b"Variable,Label,Values\nq1,\x81\x8d\x8f\x90\x9d,1=Yes\n", "cb.csv"), R(r"."), "claude/one-reader-import.md Encoding: refused, naming the byte")

# ---- XLSForm questionnaires (instrument_sheet, no model)
XLSFORM_SURVEY = [["type", "name", "label", "relevant"], ["start", "start", "", ""], ["begin group", "g1", "Service", ""], ["select_one yn", "q1", "Would you recommend us?", ""], ["select_multiple brands", "q2", "Which brands have you used?", "${q1} = '1'"], ["end group", "", "", ""], ["integer", "age", "Age in years", ""], ["note", "n1", "Thank you", ""], ["calculate", "c1", "", ""]]
XLSFORM_CHOICES = [["list_name", "name", "label"], ["yn", 1, "Yes"], ["yn", 2, "No"], ["brands", "a", "Alpha"], ["brands", "b", "Beta"]]
XLSFORM_ITEMS = [{"v": "q1", "label": "Would you recommend us?", "opts": YN, "section": "Service"}, {"v": "q2", "label": "Which brands have you used?", "opts": [("a", "Alpha"), ("b", "Beta")], "section": "Service", "routing": "${q1} = '1'"}, {"v": "age", "label": "Age in years", "opts": [], "section": None}]
case("X001", "XLSForm: questions, choice lists, group as section, relevant kept, notes/metadata skipped", "contract",
     lambda: (xlsx_bytes([("survey", XLSFORM_SURVEY), ("choices", XLSFORM_CHOICES)]), "form.xlsx", "instrument_sheet"), I(XLSFORM_ITEMS), CB + " XLSForm read without the model")
case("X002", "XLSForm with 'label::English (en)' columns", "contract",
     lambda: (xlsx_bytes([("survey", [["type", "name", "label::English (en)"], ["select_one yn", "q1", "Would you recommend us?"], ["text", "q3", "Any comments?"]]), ("choices", [["list_name", "name", "label::English (en)"], ["yn", "1", "Yes"], ["yn", "2", "No"]])]), "form.xlsx", "instrument_sheet"),
     I([{"v": "q1", "label": "Would you recommend us?", "opts": YN}, {"v": "q3", "label": "Any comments?", "opts": []}]), CB + " XLSForm")
case("X003", "XLSForm select_one naming a list missing from choices", "question",
     lambda: (xlsx_bytes([("survey", [["type", "name", "label"], ["select_one missing_list", "q1", "Would you recommend us?"]]), ("choices", [["list_name", "name", "label"], ["yn", "1", "Yes"]])]), "form.xlsx", "instrument_sheet"),
     F(["q1"]), "requirement 1: unknown never silently empty", "S3")
case("X004", "XLSForm 'select_one yn or_other' keeps the yn list", "contract",
     lambda: (xlsx_bytes([("survey", [["type", "name", "label"], ["select_one yn or_other", "q1", "Recommend?"]]), ("choices", [["list_name", "name", "label"], ["yn", "1", "Yes"], ["yn", "2", "No"]])]), "form.xlsx", "instrument_sheet"),
     I([{"v": "q1", "opts": YN}]), CB + " its choice list as options")
case("X005", "a spreadsheet questionnaire that is not an XLSForm is returned as text for the model reader", "contract",
     lambda: (xlsx_bytes([("Questionnaire", [["Q1", "Would you recommend us?"], ["", "1 Yes"], ["", "2 No"]])]), "q.xlsx", "instrument_sheet"),
     [("text", r"Q1 \| Would you recommend us\?")], CB + " Any other spreadsheet's sheets are turned into text, row by row")


# ---------------------------------------------------------------- IGS (hand reading)

def igs_build():
    with open(os.path.join(IGS_DIR, "May26_codebook.xlsx"), "rb") as fh:
        return fh.read(), "May26_codebook.xlsx"


def igs_csv_build():
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(IGS_DIR, "May26_codebook.xlsx"))
    rows = [["" if c is None else str(c) for c in r] for r in wb.active.iter_rows(values_only=True)]
    return csv_bytes(rows), "May26_codebook.csv"


# Hand reading (28 October 2026) of May26_codebook.xlsx, one sheet "May26_codebook", 268 rows x 3 columns.
# Row 1 is the header, written the wrong way round: "VARIABLE LABEL" heads the names, "VARIABLE NAME" the
# labels. Rows 3, 23 and 42 are section headings. Row 39 lists six names in one cell. 266 variables.
IGS_SECTIONS = {"AGE": "ESSENTIAL VARIABLES FROM PDI RECORD", "URBANICITY": "ESSENTIAL VARIABLES FROM PDI RECORD", "ID": "METADATA VARIABLES", "LA_region": "METADATA VARIABLES", "Q1": "SURVEY ITEMS", "Q48": "SURVEY ITEMS"}
IGS_FIRST = ["AGE", "party_reg", "CNTY", "SEX", "F12", "F13", "F14", "F15", "F16", "F17", "F18", "F19", "F20", "CD", "SD", "AD", "RZIP", "URBANICITY", "ID", "lang", "CompletionDate", "Region", "w1", "w2", "age_tab", "LA", "latlang", "votereg", "likely", "voted", "mode", "votefreqnum", "generation", "white", "black", "latino", "asian", "natam", "other", "LA_region", "Q1"]
FAV = [("1", "Favorable"), ("2", "Unfavorable"), ("3", "Don't know/no opinion")]
IGS_SPOT = [
    {"v": "AGE", "label": "Age (from PDI listing)", "opts": [], "unread": False, "section": "ESSENTIAL VARIABLES FROM PDI RECORD"},
    {"v": "party_reg", "label": "Party registration", "opts": [("1", "Democrat"), ("2", "No Party Preference"), ("3", "Other Party"), ("4", "Republican")]},
    {"v": "CNTY", "label": "CNTY (county)", "opts": [], "unread": False},
    {"v": "RZIP", "label": "Zip code", "opts": [], "unread": False},
    {"v": "ID", "label": "Identification number assigned to respondent", "opts": [], "unread": False, "section": "METADATA VARIABLES"},
    {"v": "w1", "label": "Statewide weights", "opts": []},
    {"v": "votefreqnum", "opts": [], "unread": False},
    {"v": "LA", "opts": [("0", "Non-LA City"), ("1", "LA City")]},
    {"v": "likely", "opts": [("1", "likely voter"), ("0", "not likely voter")]},
    {"v": "Q1", "label": "Q1 California will be holding a statewide primary election on June 2nd. Are you planning or not planning to vote in this election, or have you already voted early or by mail?", "opts": [("1", "Planning to vote"), ("2", "Not planning to vote"), ("3", "Already voted")], "section": "SURVEY ITEMS"},
    {"v": "Q1b", "opts": [("1", "(1) Definitely WILL NOT vote"), ("2", "(2) Probably will NOT vote"), ("3", "(3)"), ("4", "(4)"), ("5", "(5) Might or might not vote"), ("6", "(6) Possibly WILL vote"), ("7", "(7)"), ("8", "(8)"), ("9", "(9) Probably WILL vote"), ("10", "(10) Definitely WILL vote"), ("11", "Not sure")]},
    {"v": "Q8b_6", "label_prefix": "Q8_6 When considering the candidates running for M"},
    {"v": "Q25", "label": "Q25 What is your age?", "opts": [], "unread": False},
    {"v": "Q13_14", "opts": [("1", "None of these")]},
    {"v": "Q48", "opts_n": 5},
] + [{"v": f"Q5_{k}", "opts": FAV, "unread": False} for k in range(1, 9)] + [{"v": f"Q5b_{k}", "opts": [("1", "Favorable"), ("2", "Unfavorable"), ("3", "No opinion")]} for k in range(1, 9)]
IGS_Q3 = {"v": "Q3", "opts_n": 62, "opts_some": [("1", "Akinyemi Agbede, Party: Democratic, Mathematician"), ("12", "Matt Mahan, Party: Democratic, Mayor, San Jose"), ("24", 'Erin "Zez" Zezulak, Party: Democratic, Consultant/Nurse/Businesswoman'), ("31", "Alicia Olivia Lapp, Party: Republican"), ("62", "Blank/did not vote for governor")]}
IGS_Q3B = {"v": "Q3b", "opts_n": 62, "opts_some": [("62", "Undecided/Don't Know")]}
IGS_Q5_LABELS = {"Q5_1": "Tony Thurmond", "Q5_2": "Antonio Villaraigosa", "Q5_3": "Katie Porter", "Q5_4": "Xavier Becerra", "Q5_5": "Matt Mahan", "Q5_6": "Tom Steyer", "Q5_7": "Chad Bianco", "Q5_8": "Steve Hilton"}


def igs_names():
    """The 266 names in order, from the hand reading: column A of every row below the header with a
    label in column B, row 39 split into its six names."""
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(IGS_DIR, "May26_codebook.xlsx"))
    out = []
    for r in wb.active.iter_rows(min_row=2, values_only=True):
        a, b = r[0], r[1]
        if a and b:
            out += [x.strip() for x in a.split(",")] if "," in a else [a.strip()]
    return out


case("IGS01", "IGS codebook: 266 variables in order; header, blank and section rows are not variables", "contract", igs_build, [("igs_names", None)], CB + " IS-F05 (Berkeley IGS layout), IS-F07")
case("IGS02", "IGS codebook: sections, labels, answer lists and non-lists by hand reading", "contract", igs_build, I(IGS_SPOT, strict=False, sheets=["May26_codebook"], skipped=[]), CB + " Reading")
case("IGS03", "IGS codebook: Q3 ballot, 62 codes with labels holding commas and quotation marks", "contract", igs_build, I([IGS_Q3, IGS_Q3B], strict=False), CB + " IS-F04 unquoted label runs to the next ', <code>='")
case("IGS04", "IGS codebook: row 39 'white, black, latino, asian, natam, other' gives six variables", "contract", igs_build,
     I([{"v": n, "label": "Race/ethnicity (summary)"} for n in ["white", "black", "latino", "asian", "natam", "other"]], strict=False), CB + " a name cell listing several names")
case("IGS05", "IGS codebook: row 39 values '1=white and no other race; Black, ...' (a description, not an answer list)", "question", igs_build,
     F(["white", "black", "latino", "asian", "natam", "other"]), CB + " _cb_list reads it as one code '1'", "S3")
case("IGS06", "IGS codebook: grid Q5_1..Q5_8 read as eight items with the candidate in each label", "contract", igs_build,
     [("igs_q5", None)], CB + " each variable is an item")
case("IGS07", "IGS codebook: grid 5A (Q5_1..Q5_8) grouped as one grid question", "known", igs_build, [("field", "grid", {f"Q5_{k}": "Q5" for k in range(1, 9)})], "planned batch 53", "S3")
case("IGS08", "IGS codebook as CSV (same cells) reads the same 266 variables", "contract", igs_csv_build, [("igs_names", None)], "cross-format agreement")
case("IGS09", "IGS codebook: the read is complete (no unread answer lists)", "contract", igs_build, [("no_unread", None)], CB + " IS-F04 incompleteItems")


# ---------------------------------------------------------------- World Bank Findex question map (hand reading)
# indicator-mapping-by-economy.xlsx, read by hand 28 October 2026. Two sheets, same layout: row 1 header
# "Indicator Name | Filter | Question | Key"; then one row per survey question (Findex rows 2-106: 105
# questions; Connectivity rows 2-53: 52 questions); a blank row; a merged note row ("Note: no/dk/r
# indicates ..."); a blank row; a legend table "Key | Countries Surveyed | List of Countries" with keys
# 1-8 (Findex) or 1-4 (Connectivity). Names carry trailing no-break spaces in the source ("Fin11_1\xa0"),
# trimmed as the data reader trims; some hold inner spaces ("Fin 14 B", "Fin14 C"), kept as written.
# There are no answer codes anywhere: Filter is the routing condition (Fin1 or Fin2 = "yes"), Key the
# legend row giving which economies were asked. Truth: 157 questions, label = Question, no options.
FINDEX = "/mnt/user-data/attachments/2029d29d-347f-5d3f-9a5d-a57af1db9390/f429a4db-indicator-mapping-by-economy.xlsx"


def findex_build():
    with open(FINDEX, "rb") as fh:
        return fh.read(), "indicator-mapping-by-economy.xlsx"


def findex_names():
    import openpyxl
    wb = openpyxl.load_workbook(FINDEX)
    out = []
    for ws, last in ((wb["Findex"], 106), (wb["Connectivity"], 53)):
        out += [ws.cell(r, 1).value.strip() for r in range(2, last + 1)]
    return out


FX_LEGEND = {"Indicator Name", "Key", "1", "2", "3", "4", "5", "6", "7", "8"}
case("FX01", "Findex map: the header row 'Indicator Name | Filter | Question | Key' is not a variable", "contract", findex_build,
     [("absent", ["Indicator Name"])], CB + " IS-F07: the header row is never a variable", "S3")
case("FX02", "Findex map: the legend tables (Key, 1..8 / 1..4) and the note rows are not variables", "question", findex_build,
     [("absent", ["Key", "1", "2", "3", "4", "5", "6", "7", "8"])], "contracts silent on footer tables; hand reading: a legend, not variables", "S3")
case("FX03", "Findex map: the 157 question names in order (105 Findex, 52 Connectivity), trimmed, inner spaces kept", "contract", findex_build,
     [("names_without", FX_LEGEND)], CB + " IS-F05 names as written; one-reader names trimmed")
case("FX04", "Findex map: each label is the Question column", "contract", findex_build,
     I([{"v": "Fin1", "label_prefix": "An account can be used to keep or save money,"}, {"v": "Fin43A", "label_prefix": "In the past 12 months, have you received mone"}, {"v": "Con1", "label_prefix": "Do you have a mobile phone that you use to ma"}, {"v": "Fin 14 B", "label_prefix": "Are any of the following statements reasons w"}], strict=False),
     CB + " the label column is the header's, else the longest remaining text")
case("FX05", "Findex map: the Filter column (routing such as Fin42 = \"yes\") is not read as answer codes", "gap", findex_build,
     [("no_options_flag_ok", None), ("no_options", None), R(r".")], "requirement 1 and 6 (a skip condition is not an answer meaning); CB values column by content", "S2")
case("FX06", "Findex map: the Filter column kept as each question's routing text", "feature", findex_build,
     [("field", "routing", {"Fin3": "x", "Fin43A": "x"})], "feature gap: routing column in a codebook (XLSForm 'relevant' is kept)", "S3")


# Severity guesses for cases whose failure has limited reach (extra unlinkable items rather than lost or
# wrong meanings): S3. Every other failing case keeps the S2 it was declared with.
for _c in CASES:
    if _c["id"] in ("E054", "E059", "E060", "E098", "E122"):
        _c["severity"] = "S3"


# ---------------------------------------------------------------- checking

def norm_items(result):
    out = []
    for it in result.get("items", []):
        out.append({
            "v": it.get("variable"),
            "label": it.get("label", ""),
            "opts": [(str(o.get("code")), o.get("label")) for o in it.get("options", [])],
            "unread": bool(it.get("unreadValues")),
            "section": it.get("section"),
            "routing": it.get("routing"),
            "_raw": it,
        })
    return out


def check_items(actual, spec, opts):
    problems = []
    by = {}
    for a in actual:
        by.setdefault(a["v"], a)
    if opts.get("strict"):
        want = [e["v"] for e in spec]
        got = [a["v"] for a in actual]
        if want != got:
            problems.append(f"names: expected {want} got {got}")
    for e in spec:
        a = by.get(e["v"])
        if a is None:
            problems.append(f"{e['v']}: missing")
            continue
        for key in ("label", "opts", "unread", "section", "routing"):
            if key in e and e[key] != a[key]:
                problems.append(f"{e['v']}.{key}: expected {e[key]!r} got {a[key]!r}")
        if "label_prefix" in e and not a["label"].startswith(e["label_prefix"]):
            problems.append(f"{e['v']}.label: expected to start {e['label_prefix']!r} got {a['label'][:60]!r}")
        if "opts_n" in e and len(a["opts"]) != e["opts_n"]:
            problems.append(f"{e['v']}: expected {e['opts_n']} codes got {len(a['opts'])}")
        for c, l in e.get("opts_some", []):
            if (c, l) not in a["opts"]:
                problems.append(f"{e['v']}: expected option {c}={l!r}; got {[o for o in a['opts'] if o[0] == c]}")
    return problems


def evaluate(alt, result, error):
    kind = alt[0]
    if kind == "refused":
        if error is None:
            return ["expected a refusal; the file was read"]
        return [] if re.search(alt[1], error) else [f"refused, but message {error!r} does not match {alt[1]!r}"]
    if error is not None:
        return [f"refused: {error}"]
    actual = norm_items(result)
    if kind == "items":
        spec, opts = alt[1], alt[2]
        p = check_items(actual, spec, opts)
        if opts.get("sheets") is not None and result.get("sheets") != opts["sheets"]:
            p.append(f"sheets: expected {opts['sheets']} got {result.get('sheets')}")
        if opts.get("skipped") is not None and result.get("skippedSheets") != opts["skipped"]:
            p.append(f"skippedSheets: expected {opts['skipped']} got {result.get('skippedSheets')}")
        return p
    if kind == "flagged":
        by = {a["v"]: a for a in actual}
        return [f"{n}: expected flagged unread with no options; got {by[n]['opts'] if n in by else 'missing'} unread={by[n]['unread'] if n in by else None}" for n in alt[1] if n not in by or not by[n]["unread"] or by[n]["opts"]]
    if kind == "field":
        by = {a["v"]: a["_raw"] for a in actual}
        return [f"{n}: no '{alt[1]}' information on the item (keys: {sorted(by.get(n, {}).keys())})" for n, v in alt[2].items() if n not in by or alt[1] not in by[n]]
    if kind == "text":
        t = result.get("text") or ""
        return [] if re.search(alt[1], t) else [f"text does not contain {alt[1]!r}; items={bool(result.get('items'))}"]
    if kind == "igs_names":
        got = [a["v"] for a in actual]
        want = igs_names()
        p = [] if got == want else [f"expected {len(want)} names, got {len(got)}; first difference at {next((i for i, (x, y) in enumerate(zip(want, got)) if x != y), min(len(want), len(got)))}: expected {want[:3]}... got {got[:3]}..."]
        if got[: len(IGS_FIRST)] != IGS_FIRST:
            p.append(f"first {len(IGS_FIRST)} names differ: {got[:len(IGS_FIRST)]}")
        bad = [x for x in got if x.upper() in ("VARIABLE LABEL", "VARIABLE NAME", "SURVEY ITEMS", "METADATA VARIABLES")]
        if bad:
            p.append(f"header or section read as variables: {bad}")
        by = {a["v"]: a for a in actual}
        for n, s in IGS_SECTIONS.items():
            if n in by and by[n]["section"] != s:
                p.append(f"{n}.section expected {s!r} got {by[n]['section']!r}")
        return p
    if kind == "igs_q5":
        by = {a["v"]: a for a in actual}
        p = []
        for n, who in IGS_Q5_LABELS.items():
            a = by.get(n)
            if a is None:
                p.append(f"{n} missing")
                continue
            if not a["label"].startswith(n + " Regardless of whom you supported for Governor") or who not in a["label"]:
                p.append(f"{n}.label {a['label'][:90]!r} lacks stem or {who!r}")
            if a["opts"] != FAV:
                p.append(f"{n}.opts {a['opts']}")
        return p
    if kind == "absent":
        got = [a["v"] for a in actual]
        bad = [n for n in alt[1] if n in got]
        return [f"read as variables: {bad}"] if bad else []
    if kind == "names_without":
        got = [a["v"] for a in actual if a["v"] not in alt[1]]
        want = findex_names()
        if got == want:
            return []
        i = next((k for k, (x, y) in enumerate(zip(want, got)) if x != y), min(len(want), len(got)))
        return [f"expected {len(want)} names, got {len(got)}; first difference at {i}: expected {want[i:i + 3]} got {got[i:i + 3]}"]
    if kind == "no_options":
        bad = [(a["v"], a["opts"][:1]) for a in actual if a["opts"] or a["unread"]]
        return [f"{len(bad)} items given answer codes or flagged from the Filter column, e.g. {bad[:3]}"] if bad else []
    if kind == "no_options_flag_ok":
        bad = [(a["v"], a["opts"][:1]) for a in actual if a["opts"]]
        return [f"{len(bad)} items given answer codes from the Filter column, e.g. {bad[:3]}"] if bad else []
    if kind == "no_unread":
        un = [a["v"] for a in actual if a["unread"]]
        return [f"unread lists: {un}"] if un else []
    return [f"unknown expectation {kind}"]


def run(selected):
    eng = load_engine()
    results = []
    for c in CASES:
        if selected and not any(c["id"].startswith(s) for s in selected):
            continue
        built = c["build"]()
        data, fname = built[0], built[1]
        action = built[2] if len(built) > 2 else "codebook"
        result, error = None, None
        try:
            fn = eng.codebook_items if action == "codebook" else eng.instrument_sheet
            result = fn(data, fname, {})
        except ValueError as exc:
            error = str(exc)
        except Exception as exc:  # an unexpected crash is reported as such, never as a refusal
            error = f"CRASH {type(exc).__name__}: {exc}"
            if os.environ.get("L2_TRACE"):
                traceback.print_exc()
        outcomes = [evaluate(alt, result, error) for alt in c["expect"]]
        ok = any(len(o) == 0 for o in outcomes)
        if error and error.startswith("CRASH"):
            ok = False
        if c["kind"] in ("contract", "gap"):
            status = "PASS" if ok else "FAIL"
        elif c["kind"] == "question":
            status = "QUESTION-agrees" if ok else "QUESTION-differs"
        elif c["kind"] == "feature":
            status = "GAP-supported" if ok else "GAP-unsupported"
        else:
            status = "KNOWN-present" if ok else "KNOWN-absent"
        detail = "" if ok else "; ".join(outcomes[0])[:600]
        results.append({"id": c["id"], "title": c["title"], "kind": c["kind"], "status": status, "clause": c["clause"], "severity": c["severity"] if status == "FAIL" else None, "detail": detail,
                        "actual": None if result is None else [{k: v for k, v in a.items() if k != "_raw"} for a in norm_items(result)][:12], "error": error})
        print(f"{status:17} {c['id']:6} {c['title'][:90]}" + (f"\n{'':24}{detail[:300]}" if detail else ""))
    os.makedirs(os.path.join(HERE, "results"), exist_ok=True)
    with open(os.path.join(HERE, "results", "engine.json"), "w", encoding="utf-8") as fh:
        json.dump(results, fh, ensure_ascii=False, indent=1)
    counts = {}
    for r in results:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    print("\n" + ", ".join(f"{k} {v}" for k, v in sorted(counts.items())) + f"  (total {len(results)})")
    return 1 if counts.get("FAIL") else 0


if __name__ == "__main__":
    sys.exit(run(sys.argv[1:]))
