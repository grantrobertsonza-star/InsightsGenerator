#!/usr/bin/env python3
"""L3 real-dataset conformance for the Insights Elevator data reader (api/parse-stats-file.py).

Checks that the engine reads real research files correctly against expected values that do not come
from the engine:
  1. an independent reading of each file with pyreadstat, pandas, openpyxl or the csv module used
     directly (never through the engine's own functions),
  2. cross-format agreement (the World Bank Global Findex 2025 microdata as .sav, .dta, .csv and the
     fixed-width ASCII .dat + .sps, which this script parses itself),
  3. a few hand-verified cells and workbook-computed anchors written into this file.

The engine is loaded from its path with importlib and called with local file bytes, bypassing storage:
  dictionary(bytes, name, args)       the Variables dictionary (calls _read, the read every calculation uses)
  convert(bytes, name)                the stored preview of an SPSS/Stata file
  read_table(bytes, name, args)       the stored preview of a CSV/TSV/TXT/Excel file
  categories(bytes, name, var)        the answer list with unweighted counts
  crosstab(bytes, name, args)         a weighted one-variable table (no banner)
  codebook_items(bytes, name)         a codebook table read without the model
  _read(bytes, name, selected_sheet)  the data frame and metadata behind all of the above (read-only use)
  _non_answer_codes(labels, ranges, True)  the codes a table treats as non-answers
For the case-limit measurement only, IMPORT_LIMITS["maxCases"] is raised in the loaded module's memory
(never in api/import-limits.json); every result produced that way is labelled "bypass".

Each heavy read runs in its own child process (one large file at a time) with a resident-memory
watchdog, so peak memory (ru_maxrss) and wall time are per read and memory is freed between reads.

Usage:
  python3 conformance/l3/run.py                      all datasets
  python3 conformance/l3/run.py --only findex,igs    some datasets
  python3 conformance/l3/run.py --no-full            skip the full-size Findex reads (limit and measurement)
Options: --work DIR (reduced files and arrays; default a temp folder), --mem-cap-mb N (default 3500),
         --timeout S (default 900), --data-root DIR (or env L3_DATA_ROOT).
Writes conformance/l3/results.json. A dataset whose file is missing is reported "skipped" with the path;
a skipped check is never counted as passed.
"""
import argparse
import collections
import csv
import datetime
import hashlib
import importlib.util
import io
import json
import math
import os
import platform
import re
import resource
import subprocess
import sys
import tempfile
import time
import traceback

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", ".."))
ENGINE = os.path.join(REPO, "api", "parse-stats-file.py")
LIMITS_FILE = os.path.join(REPO, "api", "import-limits.json")
DEFAULT_DATA_ROOT = "/mnt/user-data/uploads/InsighsAI/TestProjects"

FREQ_MAX = 300  # a variable with more distinct values is compared by its cell hash and sums, not a frequency table
REDUCED_N = 5000

# ---------------------------------------------------------------------------------------------
# Dataset definitions (paths relative to the data root)
# ---------------------------------------------------------------------------------------------
FX = "Findex Data"
FINDEX = {
    "sav": f"{FX}/WLD_2024_FINDEX_v01_M_SPSS/findex_microdata_2025_labelled.sav",
    "dta": f"{FX}/WLD_2024_FINDEX_v01_M_STATA14/findex_microdata_2025_labelled.dta",
    "csv": f"{FX}/WLD_2024_FINDEX_v01_M_CSV/findex_microdata_2025_labelled.csv",
    "dat": f"{FX}/WLD_2024_FINDEX_v01_M_ASCII_SPSS/findex_microdata_2025_labelled.dat",
    "sps": f"{FX}/WLD_2024_FINDEX_v01_M_ASCII_SPSS/findex_microdata_2025_labelled.sps",
}
FINDEX_WEIGHT = "wgt"
FINDEX_SAMPLE = ["female", "educ", "inc_q", "account", "fin2", "receive_wages", "con12", "con32", "economycode", "regionwb"]
FINDEX_XTAB = ["educ", "account", "receive_wages", "con32"]
FINDEX_CONTINUOUS = ["wgt", "pop_adult", "wpid_random", "age"]

# Hand-verified cells: read by eye from the fixed-width .dat record and the matching CSV line, both in the
# file order (record 1 is line 1 of the .dat and line 2 of the CSV; record 5000 is .dat line 5000 and CSV
# line 5001). None means a blank (missing) field.
FINDEX_HAND_CELLS = {
    0: {"year": 2024, "economy": "United States", "economycode": "USA", "regionwb": "High income",
        "pop_adult": 275987833, "wpid_random": 128382770, "wgt": 3.26266445566095, "female": 2, "age": 33,
        "educ": 2, "inc_q": 4, "emp_in": 1, "urbanicity": 2, "account_fin": 1, "account_mob": None,
        "account": 1, "domestic_remittances": 4, "fin2": 1},
    4999: {"economy": "Bangladesh", "economycode": "BGD", "regionwb": "South Asia", "wgt": 0.541188983966352,
           "female": 1, "age": 32, "educ": 1, "inc_q": 1, "account": 0, "borrowed": 1, "fin45": 5,
           "con17": 4, "con32": 1, "account_mob": 0},
}

FINSCOPE_SAV = "Finscope2025/FinScope SA 2025.sav"
FINSCOPE_BENCH = "Finscope2025/Finscope25_benchwgt_MY25.xlsx"
IGS_TXT = "IGS/may2026.txt"
IGS_CODEBOOK = "IGS/May26_codebook.xlsx"
EV_XLSX = "EV Tesla/EV Demand Survey South Africa_February 12, 2025_17.33 115.xlsx"
CSM_XLSX = "CSM Tracker/csm_Wave1_Baseline.xlsx"
BRAND_XLSX = "Brand Tracker/brand_tracker_Oct_2025.xlsx"
SYN_XLSX = "Brand Tracker/Synthetic_Brand_Tracker_03_Rotated_Competitors.xlsx"

# Small and mid-size table files: (id, path, data sheet a researcher would choose (None for text), weight, sample vars)
TABLE_DATASETS = [
    ("igs", IGS_TXT, None, "w1", ["Q1", "party_reg", "SEX", "lang", "CompletionDate", "white"]),
    ("finscope_benchwgt", FINSCOPE_BENCH, "Sheet1", None, None),
    ("ev_tesla", EV_XLSX, "Data", None, None),
    ("csm_wave1", CSM_XLSX, "Respondents", None, ["Segment", "UsedIVR", "Transferred"]),
    ("brand_oct2025", BRAND_XLSX, "Respondents", None, ["Gender", "AgeBand", "Province"]),
    ("synthetic_brand", SYN_XLSX, "Data", "weight", ["wave", "gender", "awareness_b1", "wom_b1", "attr_status_b1", "satisfaction_primary"]),
    ("survey_degenerate_banner", "survey_degenerate_banner.xlsx", "Sheet1", None, ["StoreType"]),
    ("survey_many_categories", "survey_many_categories.xlsx", "Sheet1", None, ["Gender", "FavouriteProduct"]),
    ("survey_missing_values", "survey_missing_values.csv", None, None, ["Gender", "Region"]),
]

# Default severity guesses by check family (REPORT.md may refine them).
SEVERITY = {
    "row_count": "S1", "values": "S1", "weighted": "S1", "frequency": "S1", "crosstab": "S1", "hand_cell": "S1",
    "anchor": "S1", "limit_refusal": "S2", "names": "S2", "var_labels": "S2", "value_labels": "S2",
    "missing_decl": "S1", "type": "S2", "dictionary": "S2", "categories": "S2", "preview": "S2",
    "sheet_choice": "S2", "non_answer": "S2", "codebook": "S2", "cross_format": "S1", "label_sheet": "S3",
}


# =============================================================================================
# Shared helpers (used for both the independent and the engine side: they summarise a reading,
# they never read a file)
# =============================================================================================
BOOL_TAG = "\x01bool:"


def canon(v):
    """One text identity per cell: integer-valued numbers without ".0", other numbers as repr, text as
    written, logical values tagged so they never equal the text TRUE/FALSE, calendar values in ISO form.
    None for missing."""
    import numpy as np
    import pandas as pd

    if v is None:
        return None
    if isinstance(v, (bool, np.bool_)):
        return BOOL_TAG + ("TRUE" if v else "FALSE")
    if isinstance(v, str):
        return v
    if isinstance(v, (int, np.integer)):
        return str(int(v))
    if isinstance(v, (float, np.floating)):
        f = float(v)
        if f != f:
            return None
        if f.is_integer() and abs(f) < 2**63:
            return str(int(f))
        return repr(f)
    if isinstance(v, (pd.Timestamp, datetime.datetime)):
        if pd.isna(v):
            return None
        v = v.to_pydatetime() if isinstance(v, pd.Timestamp) else v
        return v.date().isoformat() if v.time() == datetime.time(0) else v.isoformat(timespec="seconds")
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, datetime.time):
        return v.isoformat(timespec="seconds")
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    return str(v)


def as_float(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if math.isfinite(f) else None


class ColumnAcc:
    """Accumulates one variable's summary over chunks of canonical values: missing count, cell-sequence
    hash, frequency of each code, weighted frequency, kind of the present values and numeric sums."""

    def __init__(self):
        self.n = 0
        self.missing = 0
        self.sha = hashlib.sha1()
        self.freq = collections.Counter()
        self.wfreq = collections.defaultdict(float)
        self.kinds = collections.Counter()
        self.sum = 0.0
        self.min = None
        self.max = None
        self.nnum = 0

    def update(self, canon_vals, raw_vals=None, weights=None):
        self.n += len(canon_vals)
        # Each value is followed by a separator, so the hash does not depend on how the rows were chunked.
        self.sha.update("".join(("\x00" if c is None else c) + "\x1f" for c in canon_vals).encode("utf-8", "surrogatepass"))
        present = [c for c in canon_vals if c is not None]
        self.missing += len(canon_vals) - len(present)
        self.freq.update(present)
        if weights is not None:
            wf = self.wfreq
            for c, w in zip(canon_vals, weights):
                if c is not None and w is not None and w > 0:
                    wf[c] += w
        if raw_vals is not None:
            for r in raw_vals:
                if r is None:
                    continue
                k = type(r).__name__
                self.kinds[k] += 1
                if isinstance(r, bool) or k == "bool_":
                    continue
                f = as_float(r) if not isinstance(r, str) else None
                if f is not None:
                    self.nnum += 1
                    self.sum += f
                    self.min = f if self.min is None else min(self.min, f)
                    self.max = f if self.max is None else max(self.max, f)

    def result(self, kind):
        out = {
            "kind": kind,
            "n": self.n,
            "n_missing": self.missing,
            "n_distinct": len(self.freq),
            "hash": self.sha.hexdigest(),
            "value_types": dict(self.kinds),
        }
        if len(self.freq) <= FREQ_MAX:
            out["freq"] = dict(self.freq)
            if self.wfreq:
                out["wfreq"] = {k: round(v, 6) for k, v in self.wfreq.items()}
        if self.nnum:
            out["num"] = {"count": self.nnum, "sum": self.sum, "min": self.min, "max": self.max}
        return out


def frame_kind(series):
    k = series.dtype.kind
    if k in "iuf":
        return "numeric"
    if k == "b":
        return "bool"
    if k == "M":
        return "datetime"
    types = {type(v).__name__ for v in series.dropna().tolist()}
    if not types:
        return "empty"
    if types <= {"str"}:
        return "text"
    if types <= {"int", "float"}:
        return "numeric"
    if types <= {"bool"}:
        return "bool"
    return "mixed:" + ",".join(sorted(types))


def summarize_frame(df, weight=None, arrays=None, arrays_dir=None, tag=None):
    """Per-variable summary of a pandas frame (any reader)."""
    import numpy as np

    w = None
    if weight and weight in df.columns:
        w = [as_float(x) for x in df[weight].tolist()]
    cols = {}
    for name in df.columns:
        s = df[name]
        raw = s.tolist()
        acc = ColumnAcc()
        acc.update([canon(v) for v in raw], raw, w)
        cols[str(name)] = acc.result(frame_kind(s))
        if arrays and name in arrays and arrays_dir:
            arr = np.array([np.nan if as_float(x) is None else float(x) for x in raw], dtype=float)
            np.save(os.path.join(arrays_dir, f"{tag}__{name}.npy"), arr)
    return {"n_rows": int(len(df)), "names": [str(c) for c in df.columns], "columns": cols}


def canon_labels(vl):
    return {str(var): {canon(k): str(v) for k, v in (labs or {}).items()} for var, labs in (vl or {}).items()}


def canon_ranges(mr):
    out = {}
    for var, ranges in (mr or {}).items():
        if not ranges:
            continue
        out[str(var)] = sorted(
            [[canon(r.get("lo")), canon(r.get("hi"))] if isinstance(r, dict) else [canon(r), canon(r)] for r in ranges]
        )
    return out


# Independent classification of a non-answer label: the label says nothing but "don't know", "refused",
# "not applicable" and the like. A label that also carries an answer ("No/DK/Ref") is an answer.
_NONANSWER_WORDS = [
    r"refused\s*to\s*(answer|say|respond)", r"not\s*sure",
    r"don['’`]?t\s*know", r"do\s*not\s*know", r"dont\s*know", r"refused", r"refusal", r"refuse",
    r"not\s*applicable", r"no\s*answer", r"no\s*response", r"prefer\s*not\s*to\s*(say|answer)", r"not\s*stated",
    r"unknown", r"missing", r"\bdk\b", r"\bref\b", r"\bn/?a\b",
]


def independent_non_answer(label):
    s = str(label).lower()
    if not s.strip():
        return False
    for pat in _NONANSWER_WORDS:
        s = re.sub(pat, " ", s)
    return re.sub(r"[\s/\\,.;:()\-_'\"]+", "", s) == ""


# =============================================================================================
# Child process tasks
# =============================================================================================
def peak_mb():
    return resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / 1024.0


def load_engine(bypass=False):
    spec = importlib.util.spec_from_file_location("psf_l3", ENGINE)
    psf = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(psf)
    if bypass:
        # In this process's memory only: lets the read past the case limit for measurement.
        psf.IMPORT_LIMITS["maxCases"] = 10**12
    return psf


# ---- independent readers -------------------------------------------------------------------
def indep_stats(spec):
    """pyreadstat directly. SPSS: read with user_missing=True and the declared missing values masked here
    (SPSS semantics: a user-missing value is not an answer). Stata: Stata's own missing values are missing."""
    import numpy as np
    import pandas as pd
    import pyreadstat

    path, fmt = spec["path"], spec["fmt"]
    if fmt == "sav":
        df, meta = pyreadstat.read_sav(path, user_missing=True)
        for var, ranges in (meta.missing_ranges or {}).items():
            for r in ranges:
                lo, hi = r["lo"], r["hi"]
                col = df[var]
                if col.dtype.kind in "iuf":
                    df.loc[(col >= lo) & (col <= hi), var] = np.nan
                else:
                    df.loc[col == lo, var] = None
    else:
        df, meta = pyreadstat.read_dta(path)
    if spec.get("slice_n"):
        df = df.iloc[: spec["slice_n"]].reset_index(drop=True)
    out = summarize_frame(df, spec.get("weight"), spec.get("arrays"), spec.get("arrays_dir"), spec.get("tag"))
    out["var_labels"] = {k: v for k, v in (meta.column_names_to_labels or {}).items()}
    out["value_labels"] = canon_labels(meta.variable_value_labels)
    out["missing_ranges"] = canon_ranges(getattr(meta, "missing_ranges", None))
    out["string_vars"] = [c for c, t in (meta.original_variable_types or {}).items() if t.startswith("A") or (t.startswith("%") and t.endswith("s"))]
    out["file_encoding"] = getattr(meta, "file_encoding", None)
    out["blank_strings"] = {c: int((df[c].astype(object) == "").sum()) for c in out["string_vars"] if c in df.columns}
    out["rows"] = {str(i): {c: canon(df.iloc[i][c]) for c in df.columns} for i in spec.get("rows", []) if i < len(df)}
    if spec.get("write_reduced"):
        # The reduced copy: first N cases written back with the same names, labels and value labels.
        n = spec["write_reduced"]["n"]
        sub = df.iloc[:n].copy()
        target = spec["write_reduced"]["path"]
        labels = [meta.column_names_to_labels.get(c) or "" for c in sub.columns]
        if fmt == "sav":
            pyreadstat.write_sav(sub, target, column_labels=labels, variable_value_labels=meta.variable_value_labels,
                                 missing_ranges=meta.missing_ranges or None, variable_measure=meta.variable_measure)
        else:
            vvl = {k: {int(c): lab for c, lab in v.items()} for k, v in meta.variable_value_labels.items()}
            pyreadstat.write_dta(sub, target, column_labels=labels, variable_value_labels=vvl, version=14)
        out["reduced_written"] = target
    return out


_STRICT_NUM = re.compile(r"^-?(?:(?:0|[1-9][0-9]*)(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?$")
# The declared missing-value tokens (contract one-reader-import.md, "Missing values": blank or the pandas
# default NA tokens after trimming).
NA_TOKENS = {"", "#N/A", "#N/A N/A", "#NA", "-1.#IND", "-1.#QNAN", "-NaN", "-nan", "1.#IND", "1.#QNAN", "<NA>",
             "N/A", "NA", "NULL", "NaN", "None", "n/a", "nan", "null"}
ID_TOKENS = {"id", "ids", "uuid", "guid"}


def id_like(name):
    s = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", str(name))
    s = re.sub(r"[^A-Za-z0-9]+", " ", s)
    return any(t.lower() in ID_TOKENS for t in s.split())


def strict_number(t):
    """The contract's plain decimal number: no +, no leading zero on a multi-digit whole part, no
    separators or words, and the value must survive as a double (2^53 for whole numbers, round trip
    for decimals)."""
    if not _STRICT_NUM.match(t):
        return None
    f = float(t)
    if not math.isfinite(f):
        return None
    if re.match(r"^-?[0-9]+$", t):
        return f if abs(int(t)) <= 2**53 else None
    if abs(f) > 2**53:
        return None
    if f.is_integer() and "e" not in t.lower() and "." in t:
        return f
    try:
        from decimal import Decimal
        if Decimal(repr(f)) != Decimal(t):
            return None
    except Exception:
        return None
    return f


def indep_delimited(spec):
    """pandas read_csv with every cell as text (no NA conversion), in chunks; missing tokens and the
    contract's number rule applied here. Two passes: the first decides each variable's type."""
    import pandas as pd

    path, sep = spec["path"], spec["sep"]
    raw = open(path, "rb").read()
    if raw.startswith(b"\xef\xbb\xbf"):
        enc = "utf-8-sig"
    else:
        try:
            raw.decode("utf-8")
            enc = "utf-8"
        except UnicodeDecodeError:
            enc = "cp1252"
    del raw
    kw = dict(sep=sep, dtype=str, keep_default_na=False, na_filter=False, encoding=enc, chunksize=20000, skip_blank_lines=True)
    names = None
    numeric_ok = None
    n_rows = 0
    for ch in pd.read_csv(path, **kw):
        if names is None:
            names = [str(c).strip() for c in ch.columns]
            numeric_ok = {c: True for c in names}
        ch.columns = names
        ch = ch[~(ch.apply(lambda col: col.str.strip().isin(NA_TOKENS)).all(axis=1))]
        n_rows += len(ch)
        for c in names:
            if not numeric_ok[c] or id_like(c):
                numeric_ok[c] = False
                continue
            for t in ch[c].tolist():
                if t.strip() in NA_TOKENS:
                    continue
                if strict_number(t) is None:
                    numeric_ok[c] = False
                    break
    accs = {c: ColumnAcc() for c in names}
    weight = spec.get("weight")
    rows_wanted = set(spec.get("rows", []))
    rows = {}
    offset = 0
    for ch in pd.read_csv(path, **kw):
        ch.columns = names
        ch = ch[~(ch.apply(lambda col: col.str.strip().isin(NA_TOKENS)).all(axis=1))]
        typed = {}
        for c in names:
            vals = ch[c].tolist()
            if numeric_ok[c]:
                typed[c] = [None if t.strip() in NA_TOKENS else float(t) for t in vals]
            else:
                typed[c] = [None if t.strip() in NA_TOKENS else t for t in vals]
        w = [as_float(x) for x in typed[weight]] if weight in typed else None
        for c in names:
            accs[c].update([canon(v) for v in typed[c]], typed[c], w)
        for i in rows_wanted:
            if offset <= i < offset + len(ch):
                rows[str(i)] = {c: canon(typed[c][i - offset]) for c in names}
        offset += len(ch)
    return {
        "n_rows": n_rows,
        "names": names,
        "columns": {c: accs[c].result("numeric" if numeric_ok[c] else "text") for c in names},
        "var_labels": {},
        "value_labels": {},
        "missing_ranges": {},
        "encoding": enc,
        "rows": rows,
    }


def parse_sps(sps_text):
    """The SPSS syntax that describes the fixed-width .dat: DATA LIST columns, VARIABLE LABELS, VALUE
    LABELS and MISSING VALUES (none in the Findex file; parsed if present)."""
    dl = sps_text.split("DATA LIST", 1)[1]
    dl = dl.split("/", 1)[1]
    dl = dl[: re.search(r"\s\.\s*\n", dl).start()]
    fields = [(m.group(1), int(m.group(2)), int(m.group(3) or m.group(2)), bool(m.group(4)))
              for m in re.finditer(r"([A-Za-z_][\w.$#@]*)\s+(\d+)(?:-(\d+))?\s*(\(A\))?", dl)]
    var_labels = {}
    vl_text = sps_text.split("VARIABLE LABELS", 1)[1].split("VALUE LABELS", 1)[0]
    q = r"(?:'[^']*'|\"[^\"]*\")"
    for m in re.finditer(rf"^\s*([A-Za-z_][\w.$#@]*)\s+({q}(?:\s*\+\s*{q})*)", vl_text, re.M):
        parts = re.findall(q, m.group(2))
        var_labels[m.group(1)] = "".join(p[1:-1] for p in parts)
    value_labels = {}
    if "VALUE LABELS" in sps_text:
        body = sps_text.split("VALUE LABELS", 1)[1]
        body = re.split(r"\n\s*(?:SAVE OUTFILE|EXECUTE|MISSING VALUES)", body)[0]
        current = []
        expecting_vars = True
        for line in body.splitlines():
            s = line.strip()
            if not s:
                continue
            m = re.match(rf"^(-?[0-9.]+)\s+({q})\s*(/)?\s*(\.)?\s*$", s)
            if m and not expecting_vars:
                code = canon(float(m.group(1)))
                for v in current:
                    value_labels.setdefault(v, {})[code] = m.group(2)[1:-1]
                if m.group(3) or m.group(4):
                    expecting_vars = True
                continue
            current = s.rstrip("/ .").split()
            expecting_vars = False
    missing = {}
    mm = re.search(r"MISSING VALUES(.*?)\.\s*\n", sps_text, re.S)
    if mm:
        for m in re.finditer(r"([\w\s]+?)\s*\(([^)]*)\)", mm.group(1)):
            for v in m.group(1).split():
                missing[v] = m.group(2)
    enc = re.search(r'ENCODING="([^"]+)"', sps_text)
    return fields, var_labels, value_labels, missing, (enc.group(1) if enc else "cp1252")


def indep_fixed_width(spec):
    """The ASCII .dat read by slicing each record at the .sps DATA LIST columns. (A) fields are text with
    trailing blanks removed; numeric fields are numbers, blank is missing."""
    sps_text = open(spec["sps"], encoding="cp1252").read()
    fields, var_labels, value_labels, missing, enc = parse_sps(sps_text)
    names = [f[0] for f in fields]
    accs = {n: ColumnAcc() for n in names}
    weight = spec.get("weight")
    rows_wanted = set(spec.get("rows", []))
    rows = {}
    n = 0
    limit = spec.get("slice_n")
    bad_numbers = collections.Counter()

    def flush(buf, start):
        cols = {}
        for name, s, e, is_text in fields:
            vals = []
            for line in buf:
                t = line[s - 1:e]
                if is_text:
                    t = t.rstrip()
                    vals.append(t)
                else:
                    t = t.strip()
                    if not t:
                        vals.append(None)
                    else:
                        try:
                            vals.append(float(t))
                        except ValueError:
                            bad_numbers[name] += 1
                            vals.append(None)
            cols[name] = vals
        w = [as_float(x) for x in cols[weight]] if weight in cols else None
        for name in names:
            accs[name].update([canon(v) for v in cols[name]], cols[name], w)
        for i in rows_wanted:
            if start <= i < start + len(buf):
                rows[str(i)] = {c: canon(cols[c][i - start]) for c in names}

    buf = []
    start = 0
    with open(spec["dat"], encoding=enc, newline="") as fh:
        for line in fh:
            line = line.rstrip("\r\n")
            if not line.strip():
                continue
            buf.append(line)
            n += 1
            if len(buf) == 20000:
                flush(buf, start)
                start += len(buf)
                buf = []
            if limit and n >= limit:
                break
    if buf:
        flush(buf, start)
    kinds = {f[0]: ("text" if f[3] else "numeric") for f in fields}
    return {
        "n_rows": n,
        "names": names,
        "columns": {c: accs[c].result(kinds[c]) for c in names},
        "var_labels": var_labels,
        "value_labels": value_labels,
        "missing_ranges": missing,
        "unparsed_numbers": dict(bad_numbers),
        "rows": rows,
        "sps_encoding": enc,
    }


def excel_cell_value(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if v.time() == datetime.time(0) else v.isoformat(timespec="seconds")
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, datetime.time):
        return v.isoformat(timespec="seconds")
    if isinstance(v, str) and v.strip() in NA_TOKENS:
        return None
    return v


def indep_excel(spec):
    """openpyxl directly (cached values, data_only=True): cells as the workbook holds them, calendar cells
    as ISO text, blank and NA-token cells missing, wholly blank rows skipped. Also the label sheets read
    by hand rules (exact name matches only) and every sheet's row count."""
    import openpyxl

    wb = openpyxl.load_workbook(spec["path"], data_only=True, read_only=True)
    out = {"sheets": {}}
    for ws in wb.worksheets:
        nonblank = 0
        first = True
        for r in ws.iter_rows(values_only=True):
            if first:
                first = False
                continue
            if any(c is not None and not (isinstance(c, str) and c.strip() in NA_TOKENS) for c in r):
                nonblank += 1
        out["sheets"][ws.title] = nonblank
    ws = wb[spec["sheet"]]
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    header = grid[0] if grid else []
    width = max((len(r) for r in grid), default=0)
    while width and all((len(r) < width or r[width - 1] is None) for r in grid):
        width -= 1

    def name_text(v):
        if v is None:
            return ""
        if isinstance(v, bool):
            return "TRUE" if v else "FALSE"
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        if isinstance(v, (datetime.datetime, datetime.date)):
            return excel_cell_value(v)
        return str(v).strip()

    names = [name_text(header[i] if i < len(header) else None) for i in range(width)]
    body = []
    for r in grid[1:]:
        r = list(r) + [None] * (width - len(r))
        vals = [excel_cell_value(v) for v in r[:width]]
        if all(v is None for v in vals):
            continue
        body.append(vals)
    weight = spec.get("weight")
    cols = {}
    widx = names.index(weight) if weight in names else None
    w = [as_float(r[widx]) for r in body] if widx is not None else None
    for i, nm in enumerate(names):
        vals = [r[i] for r in body]
        if id_like(nm):
            vals = [None if v is None else (canon(v) if not isinstance(v, str) else v) for v in vals]
        present = [v for v in vals if v is not None]
        types = {type(v).__name__ for v in present}
        if id_like(nm) or types <= {"str"}:
            kind = "text" if present else "empty"
        elif types <= {"int", "float"}:
            kind = "numeric"
        elif types <= {"bool"}:
            kind = "bool"
        else:
            kind = "mixed:" + ",".join(sorted(types))
        acc = ColumnAcc()
        acc.update([canon(v) for v in vals], vals, w)
        cols[nm] = acc.result(kind)
    out.update({"n_rows": len(body), "names": names, "columns": cols, "var_labels": {}, "value_labels": {}, "missing_ranges": {}})
    # Label sheets, hand rules: a row whose name cell equals a data column exactly gives that column's label.
    expected_labels, expected_vlabels, ambiguous = {}, {}, {}
    for ws2 in wb.worksheets:
        if ws2.title == spec["sheet"]:
            continue
        rows = [list(r) for r in ws2.iter_rows(values_only=True)]
        if not rows:
            continue
        hdr = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
        if not hdr or hdr[0] not in ("column", "variable"):
            continue
        label_col = next((i for i, h in enumerate(hdr) if h in ("meaning", "question / label", "label")), None)
        vl_col = next((i for i, h in enumerate(hdr) if h.startswith("value labels")), None)
        for r in rows[1:]:
            if not r or r[0] is None:
                continue
            nm = str(r[0]).strip()
            if nm not in names:
                continue
            if label_col is not None and r[label_col] is not None:
                expected_labels[nm] = str(r[label_col])
            if vl_col is not None and r[vl_col]:
                text = str(r[vl_col])
                entries = [e.strip() for e in text.split(";")]
                parsed, odd = {}, []
                for e in entries:
                    m = re.match(r"^(-?\d+)=(.*)$", e)
                    if m and not re.search(r"\bto\s+-?\d+=", m.group(2)):
                        parsed[m.group(1)] = m.group(2).strip()
                    else:
                        odd.append(e)
                if odd:
                    ambiguous[nm] = text
                if parsed and not odd:
                    expected_vlabels[nm] = parsed
    out["label_sheet_expect"] = {"labels": expected_labels, "value_labels": expected_vlabels, "ambiguous": ambiguous}
    # Workbook-computed anchors (the author's own formulas, cached values).
    if "Metric checks" in wb.sheetnames:
        mc = [list(r) for r in wb["Metric checks"].iter_rows(values_only=True)]
        out["metric_checks"] = [dict(zip([str(h) for h in mc[0]], r)) for r in mc[1:] if r and r[0] is not None]
    wb.close()
    return out


# ---- engine calls ----------------------------------------------------------------------------
def reduce_dictionary(d):
    return {
        "nCases": d.get("nCases"),
        "nVariables": len(d.get("variables", [])),
        "variables": {v["name"]: {k: v.get(k) for k in ("label", "numeric", "nValid", "labelled", "weightLike", "text", "nDistinct")} for v in d.get("variables", [])},
        "audit_weights": (d.get("audit") or {}).get("weights"),
        "audit_dk": {k: v for k, v in ((d.get("audit") or {}).get("dk") or {}).items() if k != "top"},
        "sheets": d.get("sheets"),
        "labelConflicts": len(d.get("labelConflicts") or []),
        "unmatchedLabelVariables": d.get("unmatchedLabelVariables"),
        "nSets": len(d.get("sets") or []),
    }


def reduce_convert(res, sample_labels):
    """The stored SPSS/Stata preview: headers are variable labels and values are decoded to their labels."""
    headers = res["headers"]
    out = {"nHeaders": len(headers), "nRows": len(res["rows"]), "totalVariables": res.get("totalVariables"),
           "keptVariables": res.get("keptVariables"), "reader": res.get("reader"), "decoded_freq": {}}
    for var, header in sample_labels.items():
        if header in headers:
            out["decoded_freq"][var] = dict(collections.Counter(canon(r.get(header)) for r in res["rows"] if r.get(header) is not None))
        else:
            out["decoded_freq"][var] = None
    out["headers_sample"] = headers[:5]
    return out


def reduce_read_table(res):
    if "sheetSelectionLost" in res:
        return {"sheetSelectionLost": res["sheetSelectionLost"]}
    headers = res["headers"]
    cols = {}
    for h in headers:
        acc = ColumnAcc()
        vals = [r.get(h) for r in res["rows"]]
        acc.update([canon(v) for v in vals], vals)
        cols[h] = acc.result("preview")
    return {"nHeaders": len(headers), "nRows": len(res["rows"]), "sheetName": res.get("sheetName"),
            "encoding": res.get("encoding"), "dateColumns": res.get("dateColumns"),
            "unsupportedIdentifierColumns": res.get("unsupportedIdentifierColumns"), "reader": res.get("reader"),
            "columns": cols}


def app_task(spec):
    """One engine session over one file. The first action is the measured one (peak memory is the
    process peak, so later actions in the same process are not separately measured)."""
    import gc

    t_import = time.time()
    psf = load_engine(spec.get("bypass", False))
    out = {"baseline_mb": round(peak_mb(), 1), "import_s": round(time.time() - t_import, 2), "actions": {}}
    name = spec.get("name") or os.path.basename(spec["path"])
    with open(spec["path"], "rb") as fh:
        b = fh.read()
    sheet = spec.get("selected_sheet")
    targs = {"selectedSheet": sheet} if sheet else {}
    for action in spec["actions"]:
        t0 = time.time()
        rec = {}
        try:
            if action == "dictionary":
                rec["result"] = reduce_dictionary(psf.dictionary(b, name, dict(targs)))
            elif action == "convert":
                res = psf.convert(b, name)
                rec["result"] = reduce_convert(res, spec.get("sample_headers") or {})
                del res
            elif action == "read_table":
                res = psf.read_table(b, name, dict(targs))
                rec["result"] = reduce_read_table(res)
                del res
            elif action == "codebook":
                rec["result"] = psf.codebook_items(b, name)
            else:
                raise ValueError(action)
            rec["status"] = "ok"
        except Exception as exc:  # the refusal is the expected result in the limit checks
            rec["status"] = "error"
            rec["error_type"] = type(exc).__name__
            rec["error"] = str(exc)[:1000]
        rec["wall_s"] = round(time.time() - t0, 2)
        rec["peak_mb_after"] = round(peak_mb(), 1)
        out["actions"][action] = rec
        gc.collect()
    if spec.get("summary"):
        t0 = time.time()
        try:
            df, meta = psf._read(b, name, selected_sheet=sheet)
            s = summarize_frame(df, spec.get("weight"), spec.get("arrays"), spec.get("arrays_dir"), spec.get("tag"))
            s["var_labels"] = {k: v for k, v in (meta.column_names_to_labels or {}).items()}
            s["value_labels"] = canon_labels(meta.variable_value_labels)
            s["missing_ranges"] = canon_ranges(getattr(meta, "missing_ranges", None))
            s["data_sheet"] = getattr(meta, "data_sheet", None)
            s["date_columns"] = list(getattr(meta, "date_columns", None) or [])
            s["encoding"] = getattr(meta, "encoding", None)
            s["rows"] = {str(i): {c: canon(df.iloc[i][c]) for c in df.columns} for i in spec.get("rows", []) if i < len(df)}
            na = {}
            for var, vl in (meta.variable_value_labels or {}).items():
                codes = psf._non_answer_codes(vl, (getattr(meta, "missing_ranges", None) or {}).get(var), True)
                na[var] = sorted(canon(c) for c in codes)
            s["non_answer_codes"] = na
            out["summary"] = s
            vlabs = meta.variable_value_labels or {}
            cats, xtabs = {}, {}
            for var in spec.get("sample_vars") or []:
                if var not in df.columns:
                    continue
                try:
                    cats[var] = psf.categories(b, name, var, selected_sheet=sheet)["categories"]
                except Exception as exc:
                    cats[var] = {"error": str(exc)[:300]}
            for var in spec.get("xtab_vars") or []:
                if var not in df.columns:
                    continue
                try:
                    args = {"columnVars": [], "rowItems": [{"key": var, "label": var, "vars": [var]}]}
                    if spec.get("weight"):
                        args["weightVar"] = spec["weight"]
                    if sheet:
                        args["selectedSheet"] = sheet
                    r = psf.crosstab(b, name, args)
                    q = r["questions"][0]
                    rows = []
                    for row in q["rows"]:
                        tot = row.get("total") or {}
                        rows.append({"label": row.get("label"), "isBase": bool(row.get("isBase")), "isEffBase": bool(row.get("isEffBase")),
                                     "rawCode": row.get("rawCode"), "count": tot.get("count"), "nCell": tot.get("nCell"),
                                     "base": tot.get("base")})
                    xtabs[var] = {"kind": q.get("kind"), "rows": rows, "total": r.get("total"), "weightVar": r.get("weightVar")}
                except Exception as exc:
                    xtabs[var] = {"error": f"{type(exc).__name__}: {str(exc)[:300]}"}
            out["categories"] = cats
            out["crosstabs"] = xtabs
            out["labelled_vars"] = sorted(vlabs)
        except Exception as exc:
            out["summary_error"] = f"{type(exc).__name__}: {exc}\n{traceback.format_exc()[-1500:]}"
        out["summary_s"] = round(time.time() - t0, 2)
    out["peak_mb"] = round(peak_mb(), 1)
    return out


CHILD_TASKS = {
    "indep_stats": indep_stats,
    "indep_delimited": indep_delimited,
    "indep_fixed_width": indep_fixed_width,
    "indep_excel": indep_excel,
    "app": app_task,
}


def child_main(task, spec_path, out_path):
    spec = json.load(open(spec_path))
    t0 = time.time()
    try:
        res = CHILD_TASKS[task](spec)
        res["_status"] = "ok"
    except Exception as exc:
        res = {"_status": "error", "_error": f"{type(exc).__name__}: {exc}", "_trace": traceback.format_exc()[-3000:]}
    res["_wall_s"] = round(time.time() - t0, 2)
    res["_peak_mb"] = round(peak_mb(), 1)
    with open(out_path, "w") as fh:
        json.dump(res, fh, default=str)


# =============================================================================================
# Parent: run children with a memory watchdog
# =============================================================================================
def rss_mb(pid):
    try:
        with open(f"/proc/{pid}/status") as fh:
            for line in fh:
                if line.startswith("VmRSS:"):
                    return int(line.split()[1]) / 1024.0
    except OSError:
        return None
    return None


class Runner:
    def __init__(self, work, mem_cap_mb, timeout):
        self.work = work
        self.mem_cap = mem_cap_mb
        self.timeout = timeout
        self.n = 0

    def run(self, task, spec, label):
        self.n += 1
        spec_path = os.path.join(self.work, f"spec_{self.n:03d}.json")
        out_path = os.path.join(self.work, f"out_{self.n:03d}.json")
        json.dump(spec, open(spec_path, "w"))
        if os.path.exists(out_path):
            os.unlink(out_path)
        print(f"  [{time.strftime('%H:%M:%S')}] {label} ...", flush=True)
        t0 = time.time()
        proc = subprocess.Popen([sys.executable, os.path.abspath(__file__), "--child", task, spec_path, out_path],
                                stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
        peak_polled = 0.0
        killed = None
        while proc.poll() is None:
            r = rss_mb(proc.pid)
            if r:
                peak_polled = max(peak_polled, r)
                if r > self.mem_cap:
                    killed = f"resident memory passed the harness cap of {self.mem_cap} MB"
                    proc.kill()
                    break
            if time.time() - t0 > self.timeout:
                killed = f"did not finish within the harness timeout of {self.timeout} s"
                proc.kill()
                break
            time.sleep(0.2)
        proc.wait()
        wall = round(time.time() - t0, 2)
        logs = proc.stdout.read().decode("utf-8", "replace")[-2000:] if proc.stdout else ""
        if killed:
            print(f"      killed: {killed} (after {wall} s, polled peak {peak_polled:.0f} MB)", flush=True)
            return {"_status": "killed", "_killed": killed, "_wall_s": wall, "_peak_polled_mb": round(peak_polled, 1), "_log": logs}
        if not os.path.exists(out_path):
            return {"_status": "crashed", "_returncode": proc.returncode, "_wall_s": wall, "_peak_polled_mb": round(peak_polled, 1), "_log": logs}
        res = json.load(open(out_path))
        res["_parent_wall_s"] = wall
        res["_peak_polled_mb"] = round(peak_polled, 1)
        print(f"      {res.get('_status')} in {wall} s, peak {res.get('_peak_mb')} MB", flush=True)
        return res


# =============================================================================================
# Comparisons
# =============================================================================================
class Report:
    def __init__(self, dataset):
        self.dataset = dataset
        self.checks = []
        self.discrepancies = []
        self.observations = []
        self.observe_only = False

    def check(self, family, name, ok, detail=None, compared=None):
        # Independent readings of different formats are compared to describe the files: a difference is an observation.
        status = "pass" if ok else ("observation" if family.startswith("xf_indep") else "fail")
        self.checks.append({"family": family, "check": name, "status": status, "compared": compared, "detail": detail})
        return ok

    def skip(self, family, name, reason):
        self.checks.append({"family": family, "check": name, "status": "skipped", "detail": reason})

    def disc(self, family, variable, expected, actual, note=None, severity=None):
        if self.observe_only:
            # Two independent readings of different formats disagreeing is a property of the files, not an engine fault.
            severity = "OBS"
        self.discrepancies.append({"dataset": self.dataset, "family": family, "variable": variable,
                                   "expected": expected, "actual": actual, "note": note,
                                   "severity_guess": severity or SEVERITY.get(family, "S2")})

    def obs(self, text, data=None):
        self.observations.append({"text": text, "data": data})

    def as_dict(self):
        c = collections.Counter(ch["status"] for ch in self.checks)
        return {"counts": dict(c), "checks": self.checks, "discrepancies": self.discrepancies, "observations": self.observations}


def freq_diff(e, a, limit=8):
    keys = sorted(set(e) | set(a), key=lambda k: (len(k), k))
    diffs = [(k, e.get(k, 0), a.get(k, 0)) for k in keys if e.get(k, 0) != a.get(k, 0)]
    return diffs[:limit], len(diffs)


def wfreq_diff(e, a, rel=1e-9, absol=1e-6):
    out = []
    for k in set(e) | set(a):
        x, y = e.get(k, 0.0), a.get(k, 0.0)
        if abs(x - y) > max(absol, rel * max(abs(x), abs(y))):
            out.append((k, x, y))
    return out


def compare_readings(rep, exp, act, label, skip_hash_vars=(), num_tol=None, check_types=True, family_prefix=""):
    """Expected (independent) summary against actual (engine) summary, variable by variable."""
    fp = family_prefix
    rep.observe_only = fp.startswith("xf_indep")
    try:
        _compare_readings(rep, exp, act, label, skip_hash_vars, num_tol, check_types, fp)
    finally:
        rep.observe_only = False


def _compare_readings(rep, exp, act, label, skip_hash_vars, num_tol, check_types, fp):
    if exp.get("_status") not in (None, "ok"):
        rep.skip(fp + "values", f"{label}: independent reading", f"independent reading failed: {exp.get('_error') or exp.get('_killed')}")
        return
    if act is None:
        rep.skip(fp + "values", f"{label}: engine reading", "no engine reading")
        return
    rep.check(fp + "row_count", f"{label}: row count", exp["n_rows"] == act["n_rows"], {"expected": exp["n_rows"], "actual": act["n_rows"]})
    if exp["n_rows"] != act["n_rows"]:
        rep.disc(fp + "row_count" if not fp else "cross_format", "(cases)", exp["n_rows"], act["n_rows"], label)
    en, an = exp["names"], act["names"]
    ok = en == an
    rep.check(fp + "names", f"{label}: variable names and order", ok, None if ok else {"missing": [n for n in en if n not in an][:20], "extra": [n for n in an if n not in en][:20], "n_expected": len(en), "n_actual": len(an)}, len(en))
    if not ok:
        missing = [n for n in en if n not in an]
        extra = [n for n in an if n not in en]
        if missing or extra:
            rep.disc(fp + "names" if not fp else "cross_format", "(variable list)", {"count": len(en), "only_expected": missing[:20]}, {"count": len(an), "only_actual": extra[:20]}, label)
        else:
            rep.disc(fp + "names" if not fp else "cross_format", "(variable order)", en[:10], an[:10], label, "S3")
    common = [n for n in en if n in an]
    bad_vals, bad_types, bad_w = 0, 0, 0
    for n in common:
        e, a = exp["columns"][n], act["columns"][n]
        # An all-blank variable has no type of its own: any type the reader gives it is accepted.
        if check_types and e["kind"] != a["kind"] and "empty" not in (e["kind"], a["kind"]):
            bad_types += 1
            rep.disc(fp + "type" if not fp else "cross_format", n, e["kind"], a["kind"], f"{label}: variable type")
        if n in skip_hash_vars:
            en_, an_ = e.get("num") or {}, a.get("num") or {}
            tol = num_tol or 1e-9
            same = e["n_missing"] == a["n_missing"] and abs((en_.get("sum") or 0) - (an_.get("sum") or 0)) <= tol * max(1.0, abs(en_.get("sum") or 0))
            if not same:
                bad_vals += 1
                rep.disc(fp + "values" if not fp else "cross_format", n, {"missing": e["n_missing"], **en_}, {"missing": a["n_missing"], **an_}, f"{label}: continuous variable (sum within relative {tol})")
            continue
        if e["hash"] == a["hash"]:
            pass
        else:
            bad_vals += 1
            detail = {"missing_expected": e["n_missing"], "missing_actual": a["n_missing"], "distinct_expected": e["n_distinct"], "distinct_actual": a["n_distinct"]}
            if "freq" in e and "freq" in a:
                diffs, nd = freq_diff(e["freq"], a["freq"])
                detail["freq_differences"] = nd
                if nd == 0:
                    detail["note"] = "same frequencies, different cell order or cell values with equal counts"
                rep.disc(fp + ("frequency" if nd else "values") if not fp else "cross_format", n,
                         {k: x for k, x, _ in diffs} or detail, {k: y for k, _, y in diffs} or detail, f"{label}: cells differ; {json.dumps(detail)}")
            else:
                rep.disc(fp + "values" if not fp else "cross_format", n, {"missing": e["n_missing"], "distinct": e["n_distinct"], "num": e.get("num")},
                         {"missing": a["n_missing"], "distinct": a["n_distinct"], "num": a.get("num")}, f"{label}: cells differ")
        if "wfreq" in e and "wfreq" in a:
            wd = wfreq_diff(e["wfreq"], a["wfreq"])
            if wd:
                bad_w += 1
                rep.disc(fp + "weighted" if not fp else "cross_format", n, {k: x for k, x, _ in wd[:6]}, {k: y for k, _, y in wd[:6]}, f"{label}: weighted totals")
    rep.check(fp + "values", f"{label}: cell values identical per variable", bad_vals == 0, {"variables_differing": bad_vals}, len(common))
    if check_types:
        rep.check(fp + "type", f"{label}: variable types", bad_types == 0, {"variables_differing": bad_types}, len(common))
    nw = sum(1 for n in common if "wfreq" in exp["columns"][n] and "wfreq" in act["columns"][n])
    if nw:
        rep.check(fp + "weighted", f"{label}: weighted totals per code", bad_w == 0, {"variables_differing": bad_w}, nw)


def compare_metadata(rep, exp, act, label, labels=True):
    if labels:
        el, al = exp.get("var_labels") or {}, act.get("var_labels") or {}
        bad = [n for n in exp["names"] if (el.get(n) or None) != (al.get(n) or None)]
        rep.check("var_labels", f"{label}: variable labels", not bad, {"differing": len(bad)}, len(exp["names"]))
        for n in bad[:30]:
            rep.disc("var_labels", n, el.get(n), al.get(n), label)
    ev, av = exp.get("value_labels") or {}, act.get("value_labels") or {}
    bad = [n for n in set(ev) | set(av) if ev.get(n) != av.get(n)]
    rep.check("value_labels", f"{label}: value labels (raw codes)", not bad, {"differing": len(bad)}, len(set(ev) | set(av)))
    for n in sorted(bad)[:30]:
        e, a = ev.get(n) or {}, av.get(n) or {}
        rep.disc("value_labels", n, {k: e.get(k) for k in set(e) | set(a) if e.get(k) != a.get(k)},
                 {k: a.get(k) for k in set(e) | set(a) if e.get(k) != a.get(k)}, label)
    em, am = exp.get("missing_ranges") or {}, act.get("missing_ranges") or {}
    rep.check("missing_decl", f"{label}: declared user-missing values", em == am, {"expected": em, "actual": am} if em != am else {"declared": len(em)})
    if em != am:
        rep.disc("missing_decl", "(declared missing)", em, am, label)


def compare_dictionary(rep, exp, d, label):
    if not d or d.get("status") != "ok":
        rep.check("dictionary", f"{label}: dictionary built", False, d.get("error") if d else "not run")
        rep.disc("dictionary", "(dictionary)", "a dictionary", (d or {}).get("error"), label)
        return
    r = d["result"]
    rep.check("dictionary", f"{label}: nCases", r["nCases"] == exp["n_rows"], {"expected": exp["n_rows"], "actual": r["nCases"]})
    if r["nCases"] != exp["n_rows"]:
        rep.disc("row_count", "(dictionary nCases)", exp["n_rows"], r["nCases"], label)
    rep.check("dictionary", f"{label}: variable count", r["nVariables"] == len(exp["names"]), {"expected": len(exp["names"]), "actual": r["nVariables"]})
    bad = []
    for n in exp["names"]:
        v = r["variables"].get(n)
        want = exp["columns"][n]["n"] - exp["columns"][n]["n_missing"]
        if v is None or v["nValid"] != want:
            bad.append((n, want, v and v["nValid"]))
    rep.check("dictionary", f"{label}: valid (answered) count per variable", not bad, {"differing": len(bad)}, len(exp["names"]))
    for n, want, got in bad[:30]:
        rep.disc("dictionary", n, want, got, f"{label}: dictionary nValid")


def expected_categories(exp_col, vl):
    """Unweighted count per answer label, every declared label included (zero when unobserved)."""
    out = collections.Counter()
    for code, lab in (vl or {}).items():
        out[lab] += 0
    for code, n in (exp_col.get("freq") or {}).items():
        out[(vl or {}).get(code, code)] += n
    return dict(out)


def compare_categories(rep, exp, app, label, value_labels):
    cats = app.get("categories") or {}
    if not cats:
        return
    bad = 0
    for var, got in cats.items():
        if isinstance(got, dict):
            bad += 1
            rep.disc("categories", var, "an answer list", got.get("error"), label)
            continue
        if "freq" not in exp["columns"][var]:
            rep.skip("categories", f"{label}: categories() of {var}", f"{exp['columns'][var]['n_distinct']} distinct values: compared by cell hash only")
            continue
        have_labels = None
        if exp["columns"][var]["n_distinct"] > 200:
            rep.obs(f"{label}: categories() of {var} lists at most 200 answers ({exp['columns'][var]['n_distinct']} distinct values in the file); compared on the answers it lists")
            have_labels = {c["label"] for c in got}
        want = expected_categories(exp["columns"][var], value_labels.get(var))
        if have_labels is not None:
            want = {k: v for k, v in want.items() if str(k) in have_labels}
        have = collections.Counter()
        for c in got:
            have[c["label"]] += c["n"]
        if dict(have) != want:
            bad += 1
            diffs, nd = freq_diff({str(k): v for k, v in want.items()}, {str(k): v for k, v in have.items()})
            strip0 = lambda d: {re.sub(r"\.0$", "", str(k)): v for k, v in d.items() if v}
            wording_only = strip0(want) == strip0(dict(have))
            rep.disc("categories", var, {k: x for k, x, _ in diffs if x}, {k: y for k, _, y in diffs if y},
                     f"{label}: categories() counts ({nd} differ)" + ("; counts agree, only the shown answer text differs: an unlabelled whole-number code is written with '.0' because the column also holds blanks (stored as decimals)" if wording_only else ""),
                     "S3" if wording_only else None)
    rep.check("categories", f"{label}: categories() unweighted counts", bad == 0, {"variables_differing": bad}, len(cats))


def compare_crosstabs(rep, exp, app, label, value_labels, weight):
    xt = app.get("crosstabs") or {}
    if not xt:
        return
    bad = 0
    for var, t in xt.items():
        if "error" in t:
            bad += 1
            rep.disc("crosstab", var, "a table", t["error"], label)
            continue
        col = exp["columns"][var]
        vl = value_labels.get(var) or {}
        nonans = {c for c, lab in vl.items() if independent_non_answer(lab)}
        if not vl:
            # No value labels: the answer text is its own label (a CSV/Excel "DontKnow" answer).
            nonans = {c for c in (col.get("freq") or {}) if independent_non_answer(c)}
        freq = col.get("freq") or {}
        wfreq = col.get("wfreq") or {}
        want_n = {c: n for c, n in freq.items() if c not in nonans}
        want_w = {c: w for c, w in wfreq.items() if c not in nonans} if weight else {c: float(n) for c, n in want_n.items()}
        have_n, have_w = {}, {}
        for r in t["rows"]:
            if r["isBase"] or r["rawCode"] is None:
                continue
            if r["nCell"]:
                have_n[r["rawCode"]] = r["nCell"]
            if r["count"]:
                have_w[r["rawCode"]] = r["count"]
        base = next((r for r in t["rows"] if r["isBase"] and not r["isEffBase"]), None)
        problems = []
        if have_n != want_n:
            d, nd = freq_diff(want_n, have_n)
            problems.append({"unweighted_cells": d})
        wd = [(k, want_w.get(k, 0), have_w.get(k, 0)) for k in set(want_w) | set(have_w)
              if abs(want_w.get(k, 0) - have_w.get(k, 0)) > max(0.051, 1e-6 * abs(want_w.get(k, 0)))]
        if wd:
            problems.append({"weighted_cells": wd[:6]})
        if base and base.get("base"):
            bw = base["base"].get("w")
            if bw is not None and abs(bw - sum(want_w.values())) > max(0.051, 1e-6 * sum(want_w.values())):
                problems.append({"base_w": (round(sum(want_w.values()), 3), bw)})
            bn = base["base"].get("n")
            if bn is not None and bn != sum(want_n.values()):
                problems.append({"base_n": (sum(want_n.values()), bn)})
        if problems:
            bad += 1
            rep.disc("crosstab", var, {"non_answer_codes_excluded": sorted(nonans)}, problems, f"{label}: one-variable weighted table")
    rep.check("crosstab", f"{label}: crosstab() weighted and unweighted cells and base", bad == 0, {"variables_differing": bad}, len(xt))


def compare_non_answers(rep, app, value_labels, label):
    s = app.get("summary") or {}
    na = s.get("non_answer_codes") or {}
    if not value_labels:
        return
    disagree = []
    for var, vl in value_labels.items():
        mine = sorted(c for c, lab in vl.items() if independent_non_answer(lab))
        theirs = na.get(var, [])
        if mine != theirs:
            disagree.append((var, {c: vl.get(c) for c in mine}, {c: vl.get(c) for c in theirs}))
    # Grouped by the label pattern so one decision covers many variables.
    groups = collections.defaultdict(list)
    for var, m, t in disagree:
        groups[(json.dumps(m, sort_keys=True), json.dumps(t, sort_keys=True))].append(var)
    rep.check("non_answer", f"{label}: non-answer codes (engine rule against an independent label rule)", not disagree, {"variables_differing": len(disagree), "patterns": len(groups)}, len(value_labels))
    for (m, t), vars_ in sorted(groups.items(), key=lambda kv: -len(kv[1]))[:15]:
        rep.disc("non_answer", f"{vars_[0]} (+{len(vars_) - 1} more)" if len(vars_) > 1 else vars_[0], json.loads(m), json.loads(t),
                 f"{label}: codes excluded as non-answers; variables: {', '.join(vars_[:12])}", "S2")


def compare_rows(rep, rows_by_idx, hand, label, wtol=1e-12):
    bad = 0
    n = 0
    for idx, cells in hand.items():
        row = (rows_by_idx or {}).get(str(idx))
        if row is None:
            rep.skip("hand_cell", f"{label}: record {idx + 1}", "record not in this reading")
            continue
        for var, want in cells.items():
            n += 1
            got = row.get(var)
            if isinstance(want, float) and not float(want).is_integer():
                ok = got is not None and abs(float(got) - want) <= wtol * max(1, abs(want)) + 1e-14
            else:
                ok = got == canon(want)
            if not ok:
                bad += 1
                rep.disc("hand_cell", f"{var} (record {idx + 1})", canon(want), got, label)
    if n:
        rep.check("hand_cell", f"{label}: hand-verified cells", bad == 0, {"differing": bad}, n)


# =============================================================================================
# Datasets
# =============================================================================================
def data_path(root, rel):
    return os.path.join(root, rel)


def measure_record(label, res, action):
    a = (res.get("actions") or {}).get(action) or {}
    return {
        "read": label,
        "status": res.get("_status") if res.get("_status") != "ok" else a.get("status"),
        "wall_s_action": a.get("wall_s"),
        "wall_s_process": res.get("_wall_s") or res.get("_parent_wall_s"),
        "peak_mb_after_action": a.get("peak_mb_after"),
        "peak_mb_process": res.get("_peak_mb") or res.get("_peak_polled_mb"),
        "baseline_mb": res.get("baseline_mb"),
        "error": a.get("error") or res.get("_killed") or res.get("_error"),
    }


def run_findex(root, runner, work, results, full=True):
    rep = Report("findex")
    paths = {k: data_path(root, v) for k, v in FINDEX.items()}
    missing = [p for k, p in paths.items() if not os.path.exists(p)]
    if missing:
        for p in missing:
            rep.skip("files", "Findex file present", f"missing: {p}")
    arrays_dir = os.path.join(work, "arrays")
    os.makedirs(arrays_dir, exist_ok=True)
    limits = json.load(open(LIMITS_FILE))
    measurements = []
    indep = {}
    rows = [0, 4999]
    # ---- independent readings, full files --------------------------------------------------
    red_paths = {"sav": os.path.join(work, "findex_first5000.sav"), "dta": os.path.join(work, "findex_first5000.dta"),
                 "csv": os.path.join(work, "findex_first5000.csv")}
    for fmt in ("sav", "dta"):
        if os.path.exists(paths[fmt]):
            indep[fmt] = runner.run("indep_stats", {"path": paths[fmt], "fmt": fmt, "weight": FINDEX_WEIGHT, "arrays": FINDEX_CONTINUOUS,
                                                    "arrays_dir": arrays_dir, "tag": f"indep_{fmt}", "rows": rows,
                                                    "write_reduced": {"n": REDUCED_N, "path": red_paths[fmt]}}, f"Findex {fmt}: independent reading (pyreadstat)")
    if os.path.exists(paths["csv"]):
        indep["csv"] = runner.run("indep_delimited", {"path": paths["csv"], "sep": ",", "weight": FINDEX_WEIGHT, "rows": rows}, "Findex csv: independent reading (pandas, text)")
        # The reduced CSV: the original header and first 5,000 records, byte for byte.
        with open(paths["csv"], "rb") as src, open(red_paths["csv"], "wb") as dst:
            for i, line in enumerate(src):
                if i > REDUCED_N:
                    break
                dst.write(line)
    if os.path.exists(paths["dat"]) and os.path.exists(paths["sps"]):
        indep["dat"] = runner.run("indep_fixed_width", {"dat": paths["dat"], "sps": paths["sps"], "weight": FINDEX_WEIGHT, "rows": rows}, "Findex dat+sps: independent fixed-width parse")
    # Independent truth for the reduced files: the first 5,000 cases of each original, read independently.
    indep_red = {}
    for fmt in ("sav", "dta"):
        if os.path.exists(paths[fmt]):
            indep_red[fmt] = runner.run("indep_stats", {"path": paths[fmt], "fmt": fmt, "weight": FINDEX_WEIGHT, "slice_n": REDUCED_N, "rows": rows}, f"Findex {fmt}: independent reading, first {REDUCED_N}")
    if os.path.exists(red_paths["csv"]):
        indep_red["csv"] = runner.run("indep_delimited", {"path": red_paths["csv"], "sep": ",", "weight": FINDEX_WEIGHT, "rows": rows}, f"Findex csv: independent reading, first {REDUCED_N}")
    if "dat" in indep:
        indep_red["dat"] = runner.run("indep_fixed_width", {"dat": paths["dat"], "sps": paths["sps"], "weight": FINDEX_WEIGHT, "slice_n": REDUCED_N, "rows": rows}, f"Findex dat: independent parse, first {REDUCED_N}")

    # ---- hand cells against every independent reading --------------------------------------
    for fmt, r in indep.items():
        if r.get("_status") == "ok":
            compare_rows(rep, r.get("rows"), FINDEX_HAND_CELLS, f"independent {fmt} (full)", wtol=1e-13)

    # ---- cross-format agreement among the independent readings ------------------------------
    base = indep.get("sav")
    if base and base.get("_status") == "ok":
        for fmt in ("dta", "csv", "dat"):
            other = indep.get(fmt)
            if other and other.get("_status") == "ok":
                compare_readings(rep, base, other, f"independent sav vs independent {fmt}", skip_hash_vars=("wgt",), num_tol=1e-12,
                                 check_types=False, family_prefix="xf_indep_")
                if fmt in ("dta", "dat"):
                    ev, ov = base.get("value_labels", {}), other.get("value_labels", {})
                    diffv = [v for v in set(ev) | set(ov) if ev.get(v) != ov.get(v)]
                    rep.check("xf_indep_value_labels", f"independent sav vs {fmt}: value labels", not diffv, {"differing": sorted(diffv)[:20]}, len(set(ev) | set(ov)))
                    for v in sorted(diffv)[:20]:
                        rep.observe_only = True
                        rep.disc("cross_format", v, ev.get(v), ov.get(v), f"independent sav vs {fmt}: value labels")
                        rep.observe_only = False
                    el, ol = base.get("var_labels", {}), other.get("var_labels", {})
                    diffl = [v for v in base["names"] if (el.get(v) or None) != (ol.get(v) or None)]
                    rep.check("xf_indep_var_labels", f"independent sav vs {fmt}: variable labels", not diffl, {"differing": diffl[:20], "examples": {v: [el.get(v), ol.get(v)] for v in diffl[:5]}}, len(base["names"]))
                    for v in diffl[:20]:
                        rep.observe_only = True
                        rep.disc("cross_format", v, el.get(v), ol.get(v), f"independent sav vs {fmt}: variable label")
                        rep.observe_only = False
        # weights to full precision
        import numpy as np
        for fmt in ("dta",):
            p1 = os.path.join(arrays_dir, "indep_sav__wgt.npy")
            p2 = os.path.join(arrays_dir, f"indep_{fmt}__wgt.npy")
            if os.path.exists(p1) and os.path.exists(p2):
                a1, a2 = np.load(p1), np.load(p2)
                md = float(np.nanmax(np.abs(a1 - a2))) if a1.shape == a2.shape else None
                rep.check("xf_indep_values", f"independent sav vs {fmt}: wgt per case (max abs diff)", md is not None and md <= 1e-12, {"max_abs_diff": md})

    # ---- engine: limit refusals (no bypass) --------------------------------------------------
    app_full = {}
    if full:
        for fmt, actions in (("sav", ["dictionary"]), ("sav", ["convert"]), ("dta", ["dictionary"]), ("dta", ["convert"]), ("csv", ["dictionary"])):
            if not os.path.exists(paths[fmt]):
                rep.skip("limit_refusal", f"{fmt} {actions[0]}", f"missing {paths[fmt]}")
                continue
            res = runner.run("app", {"path": paths[fmt], "actions": actions}, f"Findex {fmt}: engine {actions[0]} at the {limits['maxCases']:,}-case limit")
            a = (res.get("actions") or {}).get(actions[0]) or {}
            msg = a.get("error") or ""
            ok = a.get("status") == "error" and f"{limits['maxCases']:,}" in msg and "case" in msg.lower()
            rep.check("limit_refusal", f"Findex {fmt} ({indep.get('sav', {}).get('n_rows', '144,090')} cases): {actions[0]} refuses, naming the limit", ok, {"message": msg, "status": a.get("status")})
            if not ok:
                rep.disc("limit_refusal", f"{fmt} {actions[0]}", "a refusal naming the 100,000 case limit", a.get("status") + ": " + msg[:200] if a else res.get("_status"))
            measurements.append(measure_record(f"Findex full .{fmt} {actions[0]}: refusal at the limit (no bypass)", res, actions[0]))
        # ---- engine: full reads with the limit bypassed in memory (measurement + full content) ----
        for fmt, action in (("sav", "dictionary"), ("dta", "dictionary"), ("csv", "dictionary"), ("sav", "convert"), ("dta", "convert"), ("csv", "read_table")):
            if not os.path.exists(paths[fmt]):
                continue
            spec = {"path": paths[fmt], "actions": [action], "bypass": True}
            if action == "dictionary":
                spec.update({"summary": True, "weight": FINDEX_WEIGHT, "arrays": FINDEX_CONTINUOUS, "arrays_dir": arrays_dir,
                             "tag": f"app_{fmt}", "rows": rows, "sample_vars": FINDEX_SAMPLE, "xtab_vars": FINDEX_XTAB})
            if action == "convert" and base and base.get("_status") == "ok":
                spec["sample_headers"] = {v: (base["var_labels"].get(v) or v) for v in FINDEX_SAMPLE}
            res = runner.run("app", spec, f"Findex {fmt}: engine {action}, case limit bypassed in memory (measurement)")
            measurements.append(measure_record(f"Findex full .{fmt} {action} (limit bypassed)", res, action))
            app_full[(fmt, action)] = res

    # ---- engine: reduced files (no bypass) ---------------------------------------------------
    app_red = {}
    for fmt in ("sav", "dta", "csv"):
        if not os.path.exists(red_paths[fmt]):
            rep.skip("values", f"reduced {fmt}", "reduced file not written (source missing)")
            continue
        acts = ["dictionary", "convert" if fmt != "csv" else "read_table"]
        spec = {"path": red_paths[fmt], "actions": acts, "summary": True, "weight": FINDEX_WEIGHT, "rows": rows,
                "sample_vars": FINDEX_SAMPLE, "xtab_vars": FINDEX_XTAB, "tag": f"red_{fmt}", "arrays": ["wgt"], "arrays_dir": arrays_dir}
        if fmt != "csv" and base and base.get("_status") == "ok":
            spec["sample_headers"] = {v: (base["var_labels"].get(v) or v) for v in FINDEX_SAMPLE}
        app_red[fmt] = runner.run("app", spec, f"Findex reduced {fmt} ({REDUCED_N} cases): engine")
        measurements.append(measure_record(f"Findex first {REDUCED_N} .{fmt} dictionary", app_red[fmt], "dictionary"))

    # ---- compare engine readings with the independent ones -----------------------------------
    def engine_checks(app, exp, label, fmt, value_labels):
        s = app.get("summary")
        if app.get("_status") != "ok" or not s:
            rep.skip("values", label, f"engine reading unavailable: {app.get('_status')} {app.get('_killed') or app.get('_error') or app.get('summary_error') or ''}")
            return
        compare_readings(rep, exp, s, label, skip_hash_vars=("wgt",) if fmt == "csv" else (), num_tol=1e-12)
        compare_metadata(rep, exp, s, label, labels=(fmt != "csv"))
        compare_dictionary(rep, exp, (app.get("actions") or {}).get("dictionary"), label)
        compare_categories(rep, exp, app, label, value_labels)
        compare_crosstabs(rep, exp, app, label, value_labels, FINDEX_WEIGHT)
        if fmt != "csv":
            compare_non_answers(rep, app, value_labels, label)
        compare_rows(rep, s.get("rows"), FINDEX_HAND_CELLS, label, wtol=(1e-13 if fmt != "csv" else 1e-13))

    sav_vl = (indep.get("sav") or {}).get("value_labels") or {}
    for fmt in ("sav", "dta", "csv"):
        r = app_full.get((fmt, "dictionary"))
        if r is not None and fmt in indep:
            engine_checks(r, indep[fmt], f"engine {fmt} full (limit bypassed) vs independent {fmt}", fmt, indep[fmt].get("value_labels") or {})
        if fmt in app_red and fmt in indep_red:
            # The CSV carries codes only; its categories and tables are compared with codes as labels.
            engine_checks(app_red[fmt], indep_red[fmt], f"engine {fmt} first {REDUCED_N} vs independent {fmt} first {REDUCED_N}", fmt, indep_red[fmt].get("value_labels") or {})
            # For the reduced SPSS/Stata copies the metadata must still be the original file's.
            if fmt != "csv" and app_red[fmt].get("summary"):
                compare_metadata(rep, indep[fmt], app_red[fmt]["summary"], f"engine {fmt} first {REDUCED_N} vs original file metadata")
    # Previews: stored preview against the calculation read
    for key, res in list(app_full.items()) + [((f, "preview_red"), app_red[f]) for f in app_red]:
        fmt, action = key
        acts = res.get("actions") or {}
        if "convert" in acts and acts["convert"].get("status") == "ok":
            pr = acts["convert"]["result"]
            exp = indep.get(fmt) if action != "preview_red" else indep_red.get(fmt)
            if not exp or exp.get("_status") != "ok":
                continue
            label = f"engine {fmt} convert preview ({'full, bypass' if action != 'preview_red' else f'first {REDUCED_N}'})"
            rep.check("preview", f"{label}: rows", pr["nRows"] == exp["n_rows"], {"expected": exp["n_rows"], "actual": pr["nRows"]})
            if pr.get("keptVariables"):
                rep.obs(f"{label}: the preview kept {pr['keptVariables']} of {pr['totalVariables']} variables (size budget)")
            bad = 0
            for var, got in (pr.get("decoded_freq") or {}).items():
                vl = exp["value_labels"].get(var) or {}
                want = collections.Counter()
                for code, n in (exp["columns"][var].get("freq") or {}).items():
                    want[vl.get(code, code)] += n
                if got is None:
                    rep.obs(f"{label}: {var} not in the preview")
                    continue
                if "freq" not in exp["columns"][var]:
                    rep.skip("preview", f"{label}: {var}", "too many distinct values for a frequency comparison")
                    continue
                if dict(want) != got:
                    bad += 1
                    d, nd = freq_diff(dict(want), got)
                    rep.disc("preview", var, {k: x for k, x, _ in d}, {k: y for k, _, y in d}, label)
            rep.check("preview", f"{label}: decoded answer counts of sample variables", bad == 0, {"differing": bad}, len(pr.get("decoded_freq") or {}))
        if "read_table" in acts and acts["read_table"].get("status") == "ok":
            pr = acts["read_table"]["result"]
            calc = (app_full.get(("csv", "dictionary")) or {}).get("summary") if action != "preview_red" else res.get("summary")
            label = f"engine csv read_table preview ({'full, bypass' if action != 'preview_red' else f'first {REDUCED_N}'})"
            if calc:
                bad = [h for h in calc["names"] if (pr["columns"].get(h) or {}).get("hash") != calc["columns"][h]["hash"]]
                rep.check("preview", f"{label}: preview cells identical to the calculation read", not bad and pr["nRows"] == calc["n_rows"], {"differing": bad[:20], "rows": [pr["nRows"], calc["n_rows"]]}, len(calc["names"]))
                for h in bad[:10]:
                    rep.disc("preview", h, "same cells as the calculation read", "different", label)

    # ---- cross-format agreement among the engine readings -------------------------------------
    def xf(sumA, sumB, label):
        if sumA and sumB:
            compare_readings(rep, sumA, sumB, label, skip_hash_vars=("wgt",), num_tol=1e-12, check_types=False, family_prefix="xf_engine_")

    full_s = {f: (app_full.get((f, "dictionary")) or {}).get("summary") for f in ("sav", "dta", "csv")}
    xf(full_s["sav"], full_s["dta"], "engine sav vs engine dta (full, bypass)")
    xf(full_s["sav"], full_s["csv"], "engine sav vs engine csv (full, bypass)")
    red_s = {f: (app_red.get(f) or {}).get("summary") for f in ("sav", "dta", "csv")}
    xf(red_s["sav"], red_s["dta"], f"engine sav vs engine dta (first {REDUCED_N})")
    xf(red_s["sav"], red_s["csv"], f"engine sav vs engine csv (first {REDUCED_N})")
    # Genuine format differences, recorded as observations
    if full_s["sav"] and full_s["csv"]:
        rep.obs("The CSV ('labelled' in its name) carries numeric codes, not labels, and no variable or value labels; "
                "the engine reads it with names only. Its wgt has 15 significant digits against 16-17 in .sav/.dta.")
    if indep.get("sav", {}).get("_status") == "ok":
        rep.obs("Declared user-missing values in the .sav: " + json.dumps(indep["sav"].get("missing_ranges")) +
                "; non-answers are coded as labelled codes (8/9, 98/99, 'DK/ref' 5) instead.")
        rep.obs("String variables (.sav): " + ", ".join(indep["sav"].get("string_vars") or []) +
                f"; blank strings: {indep['sav'].get('blank_strings')}")
    if indep.get("dat", {}).get("_status") == "ok":
        rep.obs("The ASCII .dat + .sps (Stat/Transfer SPSS syntax, CP1252) parsed independently: "
                f"{indep['dat']['n_rows']} records, {len(indep['dat']['names'])} variables, unparsed numeric fields {indep['dat'].get('unparsed_numbers')}.")
    results["datasets"]["findex"] = rep.as_dict()
    results["measurements"].extend(measurements)
    results["raw"]["findex"] = {
        "indep_sav_meta": {k: (indep.get("sav") or {}).get(k) for k in ("n_rows", "file_encoding", "missing_ranges", "string_vars", "blank_strings")},
        "engine_full_dictionary": {f: ((app_full.get((f, "dictionary")) or {}).get("actions") or {}).get("dictionary") for f in ("sav", "dta", "csv")},
        "engine_reduced_dictionary": {f: ((app_red.get(f) or {}).get("actions") or {}).get("dictionary") for f in app_red},
        "engine_reduced_crosstabs": {f: (app_red.get(f) or {}).get("crosstabs") for f in app_red},
        "engine_full_errors": {f"{k[0]}:{k[1]}": v.get("summary_error") or v.get("_error") or v.get("_killed") for k, v in app_full.items() if v.get("summary_error") or v.get("_error") or v.get("_killed")},
    }


def run_finscope(root, runner, work, results):
    rep = Report("finscope2025")
    path = data_path(root, FINSCOPE_SAV)
    if not os.path.exists(path):
        rep.skip("files", "FinScope 2025 .sav present", f"missing: {path}")
        results["datasets"]["finscope2025"] = rep.as_dict()
        return
    weight = "PERSON_WEIGHTS"
    sample = ["Province", "Metro", "Race", "A1", "B1", "Geo_Area", "SPName"]
    exp = runner.run("indep_stats", {"path": path, "fmt": "sav", "weight": weight}, "FinScope 2025: independent reading (pyreadstat)")
    if exp.get("_status") != "ok":
        rep.skip("values", "FinScope independent reading", exp.get("_error") or exp.get("_killed"))
        results["datasets"]["finscope2025"] = rep.as_dict()
        return
    labelled = [v for v in sample if v in exp["value_labels"]]
    app = runner.run("app", {"path": path, "actions": ["dictionary"], "summary": True, "weight": weight, "sample_vars": sample,
                             "xtab_vars": labelled[:5]}, "FinScope 2025: engine dictionary + calculation read")
    results["measurements"].append(measure_record("FinScope SA 2025 .sav dictionary (5,562 cases x 3,960 variables, 80.8 MB)", app, "dictionary"))
    vl = exp["value_labels"]
    if app.get("summary"):
        compare_readings(rep, exp, app["summary"], "engine vs independent")
        compare_metadata(rep, exp, app["summary"], "engine vs independent")
        compare_dictionary(rep, exp, app["actions"].get("dictionary"), "engine vs independent")
        compare_categories(rep, exp, app, "engine vs independent", vl)
        compare_crosstabs(rep, exp, app, "engine vs independent", vl, weight)
        compare_non_answers(rep, app, vl, "engine vs independent")
        d = (app["actions"].get("dictionary") or {}).get("result") or {}
        wl = sorted(n for n, v in (d.get("variables") or {}).items() if v.get("weightLike"))
        rep.obs(f"Variables the dictionary marks weight-like: {wl}")
        for w in (d.get("audit_weights") or []):
            col = exp["columns"].get(w["name"]) or {}
            num = col.get("num") or {}
            if w.get("sum") is not None and num:
                rep.check("dictionary", f"weight audit sum of {w['name']}", abs(w["sum"] - num["sum"]) <= 0.01 + 1e-9 * abs(num["sum"]), {"expected": round(num["sum"], 2), "actual": w["sum"]})
        blanks = exp.get("blank_strings") or {}
        heavy = {k: v for k, v in blanks.items() if v}
        if heavy:
            dvars = d.get("variables") or {}
            examples = sorted(heavy.items(), key=lambda kv: -kv[1])[:5]
            rep.obs("SPSS string variables with blank (empty) values: blank is a valid SPSS string value unless declared "
                    f"missing, so the engine counts it as an answer. {len(heavy)} of {len(blanks)} string variables hold blanks; "
                    "examples (blank cases, engine nValid): " + json.dumps({k: [v, (dvars.get(k) or {}).get('nValid')] for k, v in examples}))
    else:
        rep.skip("values", "FinScope engine reading", app.get("_error") or app.get("_killed") or app.get("summary_error"))
    # The stored preview, separately measured
    headers = {v: (exp["var_labels"].get(v) or v) for v in sample}
    prev = runner.run("app", {"path": path, "actions": ["convert"], "sample_headers": headers}, "FinScope 2025: engine convert (stored preview)")
    results["measurements"].append(measure_record("FinScope SA 2025 .sav convert (stored preview)", prev, "convert"))
    a = (prev.get("actions") or {}).get("convert") or {}
    if a.get("status") == "ok":
        pr = a["result"]
        rep.check("preview", "convert preview: rows", pr["nRows"] == exp["n_rows"], {"expected": exp["n_rows"], "actual": pr["nRows"]})
        if pr.get("keptVariables"):
            rep.obs(f"The stored preview kept {pr['keptVariables']} of {pr['totalVariables']} variables (the 100 MB result budget); calculations read every variable.")
        bad = 0
        for var, got in pr["decoded_freq"].items():
            if got is None:
                rep.obs(f"convert preview: {var} was left out of the preview")
                continue
            if "freq" not in exp["columns"][var]:
                rep.skip("preview", f"convert preview: {var}", "too many distinct values for a frequency comparison")
                continue
            want = collections.Counter()
            for code, n in (exp["columns"][var].get("freq") or {}).items():
                want[vl.get(var, {}).get(code, code)] += n
            if dict(want) != got:
                bad += 1
                d, nd = freq_diff(dict(want), got)
                rep.disc("preview", var, {k: x for k, x, _ in d}, {k: y for k, _, y in d}, "convert preview decoded counts")
        rep.check("preview", "convert preview: decoded counts of sample variables", bad == 0, {"differing": bad})
    else:
        rep.check("preview", "convert preview produced", False, a.get("error") or prev.get("_killed") or prev.get("_error"))
        rep.disc("preview", "(convert)", "a preview", a.get("error") or prev.get("_killed") or prev.get("_error"))
    results["datasets"]["finscope2025"] = rep.as_dict()
    results["raw"]["finscope2025"] = {"indep": {k: exp.get(k) for k in ("n_rows", "file_encoding", "missing_ranges")},
                                      "engine_dictionary": (app.get("actions") or {}).get("dictionary"),
                                      "engine_crosstabs": app.get("crosstabs"), "engine_summary_error": app.get("summary_error")}


def run_table_dataset(root, runner, work, results, ds):
    did, rel, sheet, weight, sample = ds
    rep = Report(did)
    path = data_path(root, rel)
    if not os.path.exists(path):
        rep.skip("files", f"{rel} present", f"missing: {path}")
        results["datasets"][did] = rep.as_dict()
        return
    is_text = path.lower().endswith((".csv", ".txt", ".tsv"))
    if is_text:
        sep = "\t" if did == "igs" else ","
        exp = runner.run("indep_delimited", {"path": path, "sep": sep, "weight": weight}, f"{did}: independent reading (pandas, text)")
    else:
        exp = runner.run("indep_excel", {"path": path, "sheet": sheet, "weight": weight}, f"{did}: independent reading (openpyxl)")
    if exp.get("_status") != "ok":
        rep.skip("values", f"{did} independent reading", exp.get("_error") or exp.get("_killed"))
        results["datasets"][did] = rep.as_dict()
        return
    sample = sample or []
    # First read: no persisted sheet, as on upload.
    app = runner.run("app", {"path": path, "actions": ["dictionary", "read_table"], "summary": True, "weight": weight,
                             "sample_vars": sample, "xtab_vars": sample[:4]}, f"{did}: engine (first read, no sheet chosen)")
    results["measurements"].append(measure_record(f"{did} dictionary", app, "dictionary"))
    chosen = (app.get("summary") or {}).get("data_sheet")
    if not is_text:
        rep.check("sheet_choice", f"sheet chosen on first read is the respondent data sheet '{sheet}'", chosen == sheet,
                  {"chosen": chosen, "non-blank rows per sheet (independent)": exp.get("sheets")})
        if chosen != sheet:
            rep.disc("sheet_choice", "(data sheet)", sheet, chosen, f"non-blank rows per sheet: {exp.get('sheets')}")
            app = runner.run("app", {"path": path, "actions": ["dictionary", "read_table"], "summary": True, "weight": weight, "selected_sheet": sheet,
                                     "sample_vars": sample, "xtab_vars": sample[:4]}, f"{did}: engine with sheet '{sheet}' selected")
    s = app.get("summary")
    if not s:
        rep.skip("values", f"{did} engine reading", app.get("summary_error") or app.get("_error") or app.get("_killed"))
        results["datasets"][did] = rep.as_dict()
        return
    compare_readings(rep, exp, s, "engine vs independent")
    compare_dictionary(rep, exp, app["actions"].get("dictionary"), "engine vs independent")
    vl_engine = s.get("value_labels") or {}
    # Categories and tables: expected values use the labels the engine shows only for wording; counts are by raw code.
    compare_categories(rep, exp, app, "engine vs independent", vl_engine)
    compare_crosstabs(rep, exp, app, "engine vs independent", vl_engine, weight)
    rt = (app["actions"].get("read_table") or {})
    if rt.get("status") == "ok":
        pr = rt["result"]
        bad = [h for h in s["names"] if (pr["columns"].get(h) or {}).get("hash") != s["columns"][h]["hash"]]
        rep.check("preview", "read_table preview cells identical to the calculation read", not bad and pr["nRows"] == s["n_rows"], {"differing": bad[:20]}, len(s["names"]))
        for h in bad[:10]:
            rep.disc("preview", h, "same cells as the calculation read", "different")
        if is_text:
            rep.check("preview", "encoding recorded", pr.get("encoding") == {"utf-8-sig": "utf-8-sig"}.get(exp.get("encoding"), pr.get("encoding")), {"engine": pr.get("encoding"), "independent": exp.get("encoding")})
        if pr.get("dateColumns"):
            rep.obs(f"Date columns reported: {pr['dateColumns']}")
    else:
        rep.check("preview", "read_table preview produced", False, rt.get("error"))
        rep.disc("preview", "(read_table)", "a preview", rt.get("error"))
    # Label sheets (Excel): expected labels by exact name match in a Column/Variable sheet
    lse = exp.get("label_sheet_expect") or {}
    if lse.get("labels") or lse.get("value_labels"):
        el, al = lse.get("labels") or {}, s.get("var_labels") or {}
        badl = [n for n in el if (al.get(n) or None) != el[n]]
        rep.check("label_sheet", "variable labels from the workbook's dictionary/codebook sheet (exact name matches)", not badl, {"differing": len(badl)}, len(el))
        for n in badl[:15]:
            rep.disc("label_sheet", n, el[n], al.get(n), "label from the dictionary sheet", "S3")
        ev = lse.get("value_labels") or {}
        badv = [n for n in ev if (vl_engine.get(n) or {}) != ev[n]]
        rep.check("label_sheet", "value labels from the codebook sheet (plain code=label lists)", not badv, {"differing": len(badv)}, len(ev))
        for n in badv[:15]:
            rep.disc("label_sheet", n, ev[n], vl_engine.get(n), "value labels from the codebook sheet", "S2")
        amb = lse.get("ambiguous") or {}
        if amb:
            by_cell = {}
            for n, t in amb.items():
                by_cell.setdefault(t, {"cell": t, "engine": vl_engine.get(n), "variables": []})["variables"].append(n)
            rep.obs("Codebook value-label cells that are not plain code=label lists (one entry per distinct cell), with what the engine made of them",
                    list(by_cell.values()))
            glued = [e for e in by_cell.values() if any("=" in str(lab) for lab in (e["engine"] or {}).values())]
            rep.check("label_sheet", "codebook cells that are not plain lists are not read into labels holding another code's text", not glued, {"cells": len(glued)}, len(by_cell))
            for e in glued:
                rep.disc("label_sheet", ", ".join(e["variables"][:3]), {"cell": e["cell"], "expected": "either the plain entries only, or the list reported unread (codebook-import.md: an unread list keeps no answers and is reported)"},
                         e["engine"], "in-workbook codebook sheet: a label swallowed the following entry", "S2")
    d = (app["actions"].get("dictionary") or {}).get("result") or {}
    if d.get("unmatchedLabelVariables"):
        rep.obs(f"Label-sheet entries matching no data column (engine): {d['unmatchedLabelVariables'][:15]}")
    if weight:
        wl = (d.get("variables") or {}).get(weight, {}).get("weightLike")
        rep.check("dictionary", f"'{weight}' recognised as a weight", bool(wl), {"weightLike": wl})
    results["datasets"][did] = rep.as_dict()
    results["raw"][did] = {"indep_sheets": exp.get("sheets"), "engine_sheet": chosen, "engine_dictionary": app["actions"].get("dictionary"),
                           "engine_read_table": {k: v for k, v in (rt.get("result") or {}).items() if k != "columns"},
                           "engine_crosstabs": app.get("crosstabs"), "engine_categories": app.get("categories"),
                           "engine_types": {n: c["kind"] for n, c in s["columns"].items()},
                           "metric_checks": exp.get("metric_checks")}
    return exp, app


def hand_anchors(results, root, table_out):
    """Workbook-computed and hand-counted anchors, checked against the engine's calculation read."""
    rep = Report("hand_anchors")
    # survey_missing_values.csv, counted by hand (awk over the 51 lines)
    t = table_out.get("survey_missing_values")
    if t:
        s = t[1].get("summary") or {}
        cols = s.get("columns") or {}
        want = {"n_rows": 50, "Gender": ({"Male": 21, "Female": 23}, 6), "Region": ({"North": 22, "South": 22}, 6), "Satisfaction_missing": 6}
        rep.check("anchor", "survey_missing_values: 50 cases", s.get("n_rows") == 50, {"actual": s.get("n_rows")})
        for v in ("Gender", "Region"):
            ok = cols.get(v, {}).get("freq") == want[v][0] and cols.get(v, {}).get("n_missing") == want[v][1]
            rep.check("anchor", f"survey_missing_values: {v} counts and 6 blanks", ok, {"freq": cols.get(v, {}).get("freq"), "missing": cols.get(v, {}).get("n_missing")})
            if not ok:
                rep.disc("anchor", v, want[v], [cols.get(v, {}).get("freq"), cols.get(v, {}).get("n_missing")], "survey_missing_values.csv")
        sm = cols.get("Satisfaction", {})
        ok = sm.get("n_missing") == 6 and sm.get("kind") == "numeric"
        rep.check("anchor", "survey_missing_values: Satisfaction numeric with 6 blanks", ok, {"missing": sm.get("n_missing"), "kind": sm.get("kind")})
    else:
        rep.skip("anchor", "survey_missing_values", "dataset not run")
    # EV: the workbook's own pivot tables (Tables, Multiple Response Q's), counts out of 115 respondents
    t = table_out.get("ev_tesla")
    if t:
        s = t[1].get("summary") or {}
        cols = s.get("columns") or {}
        own = cols.get("Do you currently own a vehicle?", {}).get("freq") or {}
        rep.check("anchor", "EV: own a vehicle Yes 91 / No 24 of 115 (Tables sheet 0.7913 and 0.2087)", own == {"Yes": 91, "No": 24}, {"engine": own})
        if own != {"Yes": 91, "No": 24}:
            rep.disc("anchor", "Do you currently own a vehicle?", {"Yes": 91, "No": 24}, own, "EV workbook pivot")
        typ = cols.get("What type of vehicle do you own?", {}).get("freq") or {}
        want = {"Petrol": 59, "Diesel": 28, "Electric": 3, "Hybrid": 1}
        rep.check("anchor", "EV: vehicle type Petrol 59, Diesel 28, Electric 3, Hybrid 1 of 91 (Tables sheet shares)", typ == want, {"engine": typ})
        if typ != want:
            rep.disc("anchor", "What type of vehicle do you own?", want, typ, "EV workbook pivot")
        ben = {n: c for n, c in cols.items() if n.startswith("What do you think are the 3 biggest benefits") and "Selected Ch" in n}
        counts = {}
        for n, c in ben.items():
            for k, v in (c.get("freq") or {}).items():
                counts[k] = counts.get(k, 0) + v
        want = {"Lower fuel costs": 92, "Environmental benefits": 89, "Reduced maintenance costs": 66, "Government incentives": 13, "Status symbol": 12, "No benefit": 6}
        got = {k: counts.get(k) for k in want}
        rep.check("anchor", "EV: benefit mentions (Multiple Response Q's sheet: 92, 89, 66, 13, 12, 6)", got == want, {"engine": got})
        if got != want:
            rep.disc("anchor", "benefits multiple response", want, got, "EV workbook MR sheet")
    results["datasets"]["hand_anchors"] = rep.as_dict()
    return rep


def synthetic_wave_anchor(root, runner, results):
    """Metric checks of the synthetic brand tracker (the author's COUNTIFS formulas, cached values) against
    the engine's weighted crosstab of each funnel variable by wave (weight is 1 for everyone)."""
    rep_d = results["datasets"].get("hand_anchors")
    path = data_path(root, SYN_XLSX)
    if not os.path.exists(path):
        return
    exp = runner.run("indep_excel", {"path": path, "sheet": "Data"}, "synthetic_brand: metric checks (cached formula values)")
    mc = exp.get("metric_checks") or []
    # A dedicated child: crosstab of the funnel variables with wave as the banner.
    code = r'''
import json, sys, importlib.util
spec = importlib.util.spec_from_file_location("psf_l3", sys.argv[1]); psf = importlib.util.module_from_spec(spec); spec.loader.exec_module(psf)
b = open(sys.argv[2], "rb").read()
out = {}
for k in range(1, 7):
    for var in (f"awareness_b{k}", f"consideration_b{k}", f"usage_p3m_b{k}", f"frequent_b{k}"):
        r = psf.crosstab(b, "syn.xlsx", {"columnVars": ["wave"], "rowItems": [{"key": var, "label": var, "vars": [var]}], "selectedSheet": "Data"})
        cols = [c["name"] for c in r["columns"]]
        q = r["questions"][0]
        out[var] = {"columns": cols, "rows": [{"rawCode": row.get("rawCode"), "nCell": row.get("nCell"), "isBase": row.get("isBase"), "base": row.get("base")} for row in q["rows"]]}
print(json.dumps(out))
'''
    p = subprocess.run([sys.executable, "-c", code, ENGINE, path], capture_output=True, text=True, timeout=600)
    rep = Report("hand_anchors")
    if p.returncode != 0:
        rep.check("anchor", "synthetic brand tracker metric checks", False, p.stderr[-800:])
    else:
        xt = json.loads(p.stdout)
        brands = ["Kora Coffee", "Metro Bean", "Daily Cup", "Summit Coffee", "Route Roast", "Neighbourhood Brew"]
        bad = 0
        n = 0
        for row in mc:
            if row.get("Brand") not in brands or row.get("Sample n") in (None, ""):
                continue
            k = brands.index(row["Brand"]) + 1
            wave = str(int(row["Wave"]))
            size = int(row["Sample n"])
            for metric, var in (("Awareness %", f"awareness_b{k}"), ("Consideration %", f"consideration_b{k}"), ("P3M %", f"usage_p3m_b{k}"), ("Frequent %", f"frequent_b{k}")):
                want = round(float(row[metric]) * size)
                t = xt[var]
                ci = next((i for i, c in enumerate(t["columns"]) if str(c).strip() in (wave, f"{wave}.0")), None)
                yes = next((r for r in t["rows"] if r["rawCode"] == "1"), None)
                got = yes["nCell"][ci] if (yes and ci is not None) else None
                n += 1
                if got != want:
                    bad += 1
                    rep.disc("anchor", f"{var} wave {wave}", want, got, "synthetic brand tracker Metric checks (COUNTIFS)")
        rep.check("anchor", "synthetic brand tracker: funnel counts by wave against the workbook's Metric checks", bad == 0, {"differing": bad}, n)
    d = results["datasets"]["hand_anchors"]
    d["checks"].extend(rep.checks)
    d["discrepancies"].extend(rep.discrepancies)
    d["counts"] = dict(collections.Counter(c["status"] for c in d["checks"]))


def run_igs_codebook(root, runner, results):
    rep = Report("igs_codebook")
    path = data_path(root, IGS_CODEBOOK)
    if not os.path.exists(path):
        rep.skip("files", "IGS codebook present", f"missing {path}")
        results["datasets"]["igs_codebook"] = rep.as_dict()
        return
    res = runner.run("app", {"path": path, "actions": ["codebook"]}, "IGS codebook: engine codebook_items")
    a = (res.get("actions") or {}).get("codebook") or {}
    if a.get("status") != "ok":
        rep.check("codebook", "codebook read", False, a.get("error"))
        rep.disc("codebook", "(codebook)", "items", a.get("error"))
        results["datasets"]["igs_codebook"] = rep.as_dict()
        return
    items = {it["variable"]: it for it in a["result"]["items"]}

    def opts(it):
        return {str(o.get("code")): o.get("label") for o in (it.get("options") or [])}

    # Read by eye from May26_codebook.xlsx (column A holds the names although headed "VARIABLE LABEL").
    want = {
        "party_reg": ("Party registration", {"1": "Democrat", "2": "No Party Preference", "3": "Other Party", "4": "Republican"}),
        "SEX": ("sex (from PDI listing)", {"1": "Female", "2": "Male"}),
        "lang": ("Language of survey", {"1": "English", "2": "Spanish"}),
        "CompletionDate": ("Date survey completion", {"1": "May 19-20", "2": "May 21-22", "3": "May 23-24", "4": "May 25"}),
        "URBANICITY": ("Urbanicity (from PDI listing)", {"1": "Rural", "2": "Suburb/mix", "3": "Urban"}),
        "Q1": (None, {"1": "Planning to vote", "2": "Not planning to vote", "3": "Already voted"}),
        "CNTY": ("CNTY (county)", {}),
        "ID": ("Identification number assigned to respondent", {}),
        "w1": ("Statewide weights", {}),
    }
    bad = 0
    for var, (label, options) in want.items():
        it = items.get(var)
        if it is None:
            bad += 1
            rep.disc("codebook", var, {"label": label, "options": options}, None, "not read from the codebook")
            continue
        got_label = it.get("text") or it.get("label")
        got_opts = opts(it)
        ok = (label is None or got_label == label) and got_opts == options
        if not ok:
            bad += 1
            rep.disc("codebook", var, {"label": label, "options": options}, {"label": got_label, "options": got_opts}, "IGS codebook item")
    rep.check("codebook", "IGS codebook: hand-read items (names, labels, answer lists)", bad == 0, {"differing": bad}, len(want))
    race = [v for v in ("white", "black", "latino", "asian", "natam", "other") if v in items]
    rep.check("codebook", "IGS codebook: one row naming six race variables gives six items", len(race) == 6, {"found": race})
    rep.obs(f"IGS codebook: {len(items)} items read; sheets used {a['result'].get('sheets')}; names described in the codebook "
            "but absent from may2026.txt include w2 (LA City weights).")
    results["datasets"]["igs_codebook"] = rep.as_dict()
    results["raw"]["igs_codebook"] = {"n_items": len(items), "sample": {k: items.get(k) for k in list(want)[:4]}}


# =============================================================================================
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--child", nargs=3)
    ap.add_argument("--only", default="")
    ap.add_argument("--work", default=os.path.join(tempfile.gettempdir(), "insights-l3-work"))
    ap.add_argument("--data-root", default=os.environ.get("L3_DATA_ROOT", DEFAULT_DATA_ROOT))
    ap.add_argument("--mem-cap-mb", type=int, default=3500)
    ap.add_argument("--timeout", type=int, default=900)
    ap.add_argument("--no-full", action="store_true")
    ap.add_argument("--out", default=os.path.join(HERE, "results.json"))
    args = ap.parse_args()
    if args.child:
        child_main(*args.child)
        return
    os.makedirs(args.work, exist_ok=True)
    import pandas
    import pyreadstat

    only = {x.strip() for x in args.only.split(",") if x.strip()}
    want = lambda k: not only or k in only
    results = {
        "generated": datetime.datetime.now().isoformat(timespec="seconds"),
        "engine": ENGINE,
        "engine_sha1": hashlib.sha1(open(ENGINE, "rb").read()).hexdigest(),
        "import_limits": json.load(open(LIMITS_FILE)),
        "data_root": args.data_root,
        "work_dir": args.work,
        "environment": {"python": sys.version.split()[0], "pandas": pandas.__version__, "pyreadstat": pyreadstat.__version__,
                        "platform": platform.platform(), "cpus": os.cpu_count(),
                        "mem_total_mb": round(os.sysconf("SC_PAGE_SIZE") * os.sysconf("SC_PHYS_PAGES") / 2**20),
                        "harness_mem_cap_mb": args.mem_cap_mb, "harness_timeout_s": args.timeout},
        "datasets": {}, "measurements": [], "raw": {},
    }
    if not os.path.isdir(args.data_root):
        print(f"Data root not found: {args.data_root}; every dataset is skipped.")
    runner = Runner(args.work, args.mem_cap_mb, args.timeout)
    t0 = time.time()
    if want("findex"):
        print("Findex")
        run_findex(args.data_root, runner, args.work, results, full=not args.no_full)
    if want("finscope"):
        print("FinScope 2025")
        run_finscope(args.data_root, runner, args.work, results)
    table_out = {}
    for ds in TABLE_DATASETS:
        if want(ds[0]) or (want("tables")):
            print(ds[0])
            r = run_table_dataset(args.data_root, runner, args.work, results, ds)
            if r:
                table_out[ds[0]] = r
    if want("igs") or want("igs_codebook"):
        run_igs_codebook(args.data_root, runner, results)
    if want("anchors") or not only:
        hand_anchors(results, args.data_root, table_out)
        synthetic_wave_anchor(args.data_root, runner, results)
    total = collections.Counter()
    for d in results["datasets"].values():
        total.update(d["counts"])
    results["totals"] = {"checks": dict(total), "discrepancies": sum(len(d["discrepancies"]) for d in results["datasets"].values()),
                         "by_severity": dict(collections.Counter(x["severity_guess"] for d in results["datasets"].values() for x in d["discrepancies"]))}
    results["run_seconds"] = round(time.time() - t0, 1)
    with open(args.out, "w") as fh:
        json.dump(results, fh, indent=1, default=str)
    print(json.dumps(results["totals"], indent=1))
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
