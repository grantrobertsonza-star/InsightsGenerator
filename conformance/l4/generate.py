"""L4 fixture generator. Writes every synthetic fixture FROM a declared truth, never the other way round,
and writes cases.json (the truth and expected readings) beside them.

Run: python3 -I conformance/l4/generate.py   (from the worktree root; needs openpyxl, python-docx,
python-pptx, reportlab, Pillow). Output: conformance/l4/fixtures/synthetic/*, conformance/l4/cases.json.
Real files are referenced read-only (see REAL below); the two small Findex indicator CSVs supplied by
the lead are copied into fixtures/real so the library is self-contained.
"""
import csv
import io
import json
import os
import shutil

HERE = os.path.dirname(os.path.abspath(__file__))
SYN = os.path.join(HERE, "fixtures", "synthetic")
REALDIR = os.path.join(HERE, "fixtures", "real")
os.makedirs(SYN, exist_ok=True)
os.makedirs(REALDIR, exist_ok=True)

# ---------------------------------------------------------------------------------------------
# Declared truths
# ---------------------------------------------------------------------------------------------

# Banner truth: Q2 overall satisfaction by Total, Gender, Age. Percentages are column percentages.
COLS = [
    {"id": "A", "group": "Total", "label": "Total"},
    {"id": "B", "group": "Gender", "label": "Male"},
    {"id": "C", "group": "Gender", "label": "Female"},
    {"id": "D", "group": "Age", "label": "18-34"},
    {"id": "E", "group": "Age", "label": "35-54"},
    {"id": "F", "group": "Age", "label": "55+"},
]
UNW_BASE = [1000, 480, 520, 400, 350, 250]
WTD_BASE = [1000, 490, 510, 380, 360, 260]
ANSWERS = [
    ("Very satisfied", [30, 34, 26, 38, 28, 20]),
    ("Satisfied", [40, 38, 42, 37, 42, 42]),
    ("Dissatisfied", [20, 18, 22, 17, 20, 25]),
    ("Very dissatisfied", [10, 10, 10, 8, 10, 13]),
]
NET = ("Net: Satisfied", [70, 72, 68, 75, 70, 62])
# Significance letters declared by the (simulated) tab program: column id -> letters per answer row.
LETTERS = {("Very satisfied", "B"): "C", ("Very satisfied", "D"): "EF", ("Very satisfied", "E"): "F",
           ("Dissatisfied", "F"): "D", ("Very dissatisfied", "F"): "D"}
for i in range(6):
    assert sum(a[1][i] for a in ANSWERS) == 100
    assert ANSWERS[0][1][i] + ANSWERS[1][1][i] == NET[1][i]

def banner_cells(unit="percent"):
    cells = []
    for label, vals in ANSWERS + [NET]:
        for c, v in zip(COLS, vals):
            cells.append({"row": label, "col": c["label"], "group": c["group"], "value": v, "unit": unit,
                          "letters": LETTERS.get((label, c["id"]), "")})
    return cells

BANNER_TRUTH = {
    "title": "Q2 Overall satisfaction",
    "label_column": "row labels (first column)",
    "columns": COLS,
    "row_kinds": {**{a[0]: "answer" for a in ANSWERS}, NET[0]: "net", "Unweighted base": "base_unweighted",
                  "Weighted base": "base_weighted", "Base": "base_unweighted"},
    "base_unweighted": {c["label"]: n for c, n in zip(COLS, UNW_BASE)},
    "base_weighted": {c["label"]: n for c, n in zip(COLS, WTD_BASE)},
    "cells": banner_cells(),
}

# Frequency truth: Q1 awareness.
FREQ = [("Yes", 600, 60), ("No", 350, 35), ("Don't know", 50, 5)]
FREQ_TRUTH = {
    "title": "Q1 Have you heard of Brand X?",
    "columns": [{"group": "", "label": "Count"}, {"group": "", "label": "Percent"}],
    "row_kinds": {"Yes": "answer", "No": "answer", "Don't know": "answer", "Total": "total"},
    "base_unweighted": {"Count": 1000, "Percent": 1000},
    "cells": [c for r, n, p in FREQ for c in ({"row": r, "col": "Count", "value": n, "unit": "count"},
                                              {"row": r, "col": "Percent", "value": p, "unit": "percent"})],
}

# Two stacked tables (one sheet). Columns Total / Male / Female.
STACK = [
    {"title": "Table 1: Q1 Aware of Brand X", "rows": [("Yes", [60, 64, 56]), ("No", [35, 31, 39]), ("Don't know", [5, 5, 5])],
     "base": [1000, 480, 520]},
    {"title": "Table 2: Q3 Would recommend Brand X", "rows": [("Yes", [45, 41, 49]), ("No", [40, 44, 36]), ("Don't know", [15, 15, 15])],
     "base": [600, 307, 293]},
]
SCOLS = ["Total", "Male", "Female"]

# Counts banner: Total / Male / Female counts with a Total row.
COUNTS = [("Very satisfied", [300, 163, 137]), ("Satisfied", [400, 182, 218]), ("Dissatisfied", [200, 86, 114]),
          ("Very dissatisfied", [100, 49, 51])]
COUNT_TOTAL = [1000, 480, 520]
for i in range(3):
    assert sum(r[1][i] for r in COUNTS) == COUNT_TOTAL[i]

# Dash and star conventions: "-" means no respondent (0), "*" means more than 0 but under 0.5%.
DASH = [("Brand A", ["52%", "55%", "49%"]), ("Brand B", ["47%", "45%", "51%"]), ("Brand C", ["*", "*", "-"]),
        ("Brand D", ["1%", "-", "*"])]
DASH_TRUTH_VALUES = {("Brand C", "Total"): (0, 0.5), ("Brand C", "Male"): (0, 0.5), ("Brand C", "Female"): (0, 0),
                     ("Brand D", "Male"): (0, 0), ("Brand D", "Female"): (0, 0.5), ("Brand A", "Male"): (55, 55)}

# ---------------------------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------------------------

def write_csv(name, rows):
    with open(os.path.join(SYN, name), "w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows(rows)

def pct(v):
    return f"{v}%"

def sig(label, col, v):
    return pct(v) + LETTERS.get((label, col["id"]), "")

# AGG-01 frequency CSV
write_csv("agg01_frequency.csv", [["Response", "Count", "Percent"]] + [[r, n, pct(p)] for r, n, p in FREQ] + [["Total", 1000, "100%"]])

# AGG-04 clean banner CSV (one header row, a Base row, text percentages, a net)
write_csv("agg04_banner_clean.csv",
          [[""] + [c["label"] for c in COLS], ["Base"] + UNW_BASE] +
          [[label] + [pct(v) for v in vals] for label, vals in ANSWERS + [NET]])

# AGG-05 banner with significance letters attached to the cells
write_csv("agg05_banner_sig.csv",
          [[""] + [c["label"] for c in COLS], ["Base"] + UNW_BASE] +
          [[label] + [sig(label, c, v) for c, v in zip(COLS, vals)] for label, vals in ANSWERS] +
          [[NET[0]] + [pct(v) for v in NET[1]]])

# AGG-08 counts banner with a Total row
write_csv("agg08_counts.csv", [["", "Total", "Male", "Female"]] + [[r] + v for r, v in COUNTS] + [["Total"] + COUNT_TOTAL])

# AGG-09 dash and star conventions, plus a Base row
write_csv("agg09_dash_star.csv", [["", "Total", "Male", "Female"], ["Base", 1000, 480, 520]] + [[r] + v for r, v in DASH])

# AGG-16 full banner as a CSV export: title, group row, header row, letters row, unweighted and weighted
# bases, sig letters, net, footnote.
write_csv("agg16_banner_export.csv",
          [["Table 2: " + BANNER_TRUTH["title"]] + [""] * 6,
           ["Base: All respondents"] + [""] * 6,
           ["", "Total", "Gender", "", "Age", "", ""],
           [""] + [c["label"] for c in COLS],
           [""] + [f"({c['id']})" for c in COLS],
           ["Unweighted base"] + UNW_BASE,
           ["Weighted base"] + WTD_BASE] +
          [[label] + [sig(label, c, v) for c, v in zip(COLS, vals)] for label, vals in ANSWERS] +
          [[NET[0]] + [pct(v) for v in NET[1]]] +
          [["Column proportions tested at 95%; letters mark a column significantly higher."] + [""] * 6])

# AGG-17 clean banner with the WEIGHTED base row first, then the unweighted base row
write_csv("agg17_banner_weighted_first.csv",
          [[""] + [c["label"] for c in COLS], ["Weighted base"] + WTD_BASE, ["Unweighted base"] + UNW_BASE] +
          [[label] + [pct(v) for v in vals] for label, vals in ANSWERS + [NET]])

# AGG-18 banner whose bases are given only in the column headers, "Male (n=480)"
write_csv("agg18_banner_n_in_header.csv",
          [[""] + [f"{c['label']} (n={n})" for c, n in zip(COLS, UNW_BASE)]] +
          [[label] + [pct(v) for v in vals] for label, vals in ANSWERS + [NET]])

import openpyxl
from openpyxl.styles import Font

def wb_save(wb, name):
    wb.save(os.path.join(SYN, name))

# AGG-02 frequency xlsx: percentages stored as fractions with a 0% format (how Excel holds them)
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Q1"
ws.append(["Response", "Count", "Percent"])
for r, n, p in FREQ:
    ws.append([r, n, p / 100]); ws.cell(ws.max_row, 3).number_format = "0%"
ws.append(["Total", 1000, 1.0]); ws.cell(ws.max_row, 3).number_format = "0%"
wb_save(wb, "agg02_frequency_pctformat.xlsx")

# AGG-03 frequency xlsx: percentages written as text "60%"
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Q1"
ws.append(["Response", "Count", "Percent"])
for r, n, p in FREQ:
    ws.append([r, n, pct(p)])
ws.append(["Total", 1000, "100%"])
wb_save(wb, "agg03_frequency_textpct.xlsx")

# AGG-06 banner xlsx: title, merged group header, header row, letters row, WEIGHTED base before
# unweighted, fractions with % format, a blank separator column between Gender and Age, footnote.
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Q2"
layout = ["", "Total", "Male", "Female", None, "18-34", "35-54", "55+"]  # None = blank separator column
def row_with_sep(first, vals):
    out = [first]; vi = iter(vals)
    for h in layout[1:]:
        out.append("" if h is None else next(vi))
    return out
ws.append(["Table 2: " + BANNER_TRUTH["title"]]); ws["A1"].font = Font(bold=True)
ws.append(["", "Total", "Gender", "", "", "Age", "", ""])
ws.merge_cells("C2:D2"); ws.merge_cells("F2:H2")
ws.append([h if h is not None else "" for h in layout])
ws.append(row_with_sep("", [f"({c['id']})" for c in COLS]))
ws.append(row_with_sep("Weighted base", WTD_BASE))
ws.append(row_with_sep("Unweighted base", UNW_BASE))
for label, vals in ANSWERS + [NET]:
    ws.append(row_with_sep(label, [v / 100 for v in vals]))
    for col in range(2, 9):
        if ws.cell(ws.max_row, col).value != "":
            ws.cell(ws.max_row, col).number_format = "0%"
ws.append(["Source: simulated tracker wave 3. Weighted to population; bases unweighted."])
wb_save(wb, "agg06_banner_titled.xlsx")

# AGG-07 two tables stacked on one sheet with titles, base rows, and a blank row between.
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Tables"
for i, t in enumerate(STACK):
    if i:
        ws.append([])
    ws.append([t["title"]])
    ws.append([""] + SCOLS)
    ws.append(["Base"] + t["base"])
    for r, vals in t["rows"]:
        ws.append([r] + [pct(v) for v in vals])
wb_save(wb, "agg07_stacked.xlsx")

# AGG-10 numbers stored as text cells
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Q1"
ws.append(["Response", "Count", "Percent"])
for r, n, p in FREQ:
    ws.append([r, str(n), str(p)])
    ws.cell(ws.max_row, 2).data_type = "s"; ws.cell(ws.max_row, 3).data_type = "s"
ws.append(["Total", "1000", "100"])
wb_save(wb, "agg10_numbers_as_text.xlsx")

# AGG-11 one table per sheet (two sheets)
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Q1 Awareness"
ws.append(["", "Total", "Male", "Female"]); ws.append(["Base", 1000, 480, 520])
for r, vals in STACK[0]["rows"]:
    ws.append([r] + [pct(v) for v in vals])
ws2 = wb.create_sheet("Q2 Satisfaction")
ws2.append([""] + [c["label"] for c in COLS]); ws2.append(["Base"] + UNW_BASE)
for label, vals in ANSWERS + [NET]:
    ws2.append([label] + [pct(v) for v in vals])
wb_save(wb, "agg11_two_sheets.xlsx")

# ---------------------------------------------------------------------------------------------
# Report documents
# ---------------------------------------------------------------------------------------------
AWARE = {"caption": "Table 1: Awareness of Brand X by gender (column %)",
         "headers": ["Response", "Total", "Male", "Female"],
         "rows": [["Yes", "60%", "64%", "56%"], ["No", "35%", "31%", "39%"], ["Don't know", "5%", "5%", "5%"],
                  ["Base", "1000", "480", "520"]]}
PROSE = ("Awareness of Brand X reached 60% of adults this wave. Men were more aware than women "
         "(64% against 56%). Table 1 gives the full results.")

from docx import Document
from docx.enum.text import WD_BREAK

def docx_table(doc, spec):
    doc.add_paragraph(spec["caption"])
    t = doc.add_table(rows=1 + len(spec["rows"]), cols=len(spec["headers"]))
    t.style = "Table Grid"
    for j, h in enumerate(spec["headers"]):
        t.cell(0, j).text = h
    for i, r in enumerate(spec["rows"], 1):
        for j, v in enumerate(r):
            t.cell(i, j).text = str(v)
    return t

# DOC-01: prose on page 1, table on page 2
d = Document()
d.add_heading("Brand X tracker, wave 3", 1)
d.add_paragraph(PROSE)
d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
docx_table(d, AWARE)
d.save(os.path.join(SYN, "doc01_report_table.docx"))

# DOC-02: four pages of prose, then a 70-row table that starts on page 5 and runs onto page 6+
REGIONS = [[f"Region {i:02d}", str(400 + i), f"{30 + (i % 40)}%"] for i in range(1, 71)]
d = Document()
for p in range(4):
    d.add_paragraph(f"Section {p + 1}. " + PROSE)
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
LONG = {"caption": "Table 5: Account ownership by region (row %)", "headers": ["Region", "Base", "Has account"], "rows": REGIONS}
tbl = docx_table(d, LONG)
d.save(os.path.join(SYN, "doc02_multipage_table.docx"))

# DOC-03: PDF with a two-row-header banner on page 1 and a simple table on page 2
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, Paragraph, PageBreak, Spacer, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
styles = getSampleStyleSheet()
WAVES = {"caption": "Table 3: Uses mobile money, wave 1 and wave 2 (row %)",
         "group_row": ["", "Wave 1", "", "Wave 2", ""],
         "headers": ["Province", "Yes", "No", "Yes", "No"],
         "rows": [["Gauteng", "48%", "52%", "55%", "45%"], ["Limpopo", "31%", "69%", "39%", "61%"],
                  ["Western Cape", "44%", "56%", "47%", "53%"]]}
story = [Paragraph("Mobile money report", styles["Title"]), Paragraph(WAVES["caption"], styles["Normal"]), Spacer(1, 6)]
t = Table([WAVES["group_row"], WAVES["headers"]] + WAVES["rows"])
t.setStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("SPAN", (1, 0), (2, 0)), ("SPAN", (3, 0), (4, 0))])
story += [t, PageBreak(), Paragraph(AWARE["caption"], styles["Normal"]), Spacer(1, 6)]
t2 = Table([AWARE["headers"]] + AWARE["rows"]); t2.setStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey)])
story.append(t2)
SimpleDocTemplate(os.path.join(SYN, "doc03_tables.pdf"), pagesize=A4).build(story)

# A picture of a table (for the image-only cases)
from PIL import Image, ImageDraw
img = Image.new("RGB", (900, 260), "white"); dr = ImageDraw.Draw(img)
grid = [AWARE["headers"]] + AWARE["rows"]
for i, r in enumerate(grid):
    for j, v in enumerate(r):
        dr.rectangle([20 + j * 210, 20 + i * 44, 230 + j * 210, 64 + i * 44], outline="black")
        dr.text((30 + j * 210, 34 + i * 44), str(v), fill="black")
IMG = os.path.join(SYN, "table_picture.png"); img.save(IMG)

# DOC-05: PDF whose only page is a picture of a table
SimpleDocTemplate(os.path.join(SYN, "doc05_image_only.pdf"), pagesize=A4).build([RLImage(IMG, width=450, height=130)])

# DOC-04: deck: title, native table, bar chart with data labels, picture of a table
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[0]); s.shapes.title.text = "Brand X tracker deck"
s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = AWARE["caption"]
gt = s.shapes.add_table(1 + len(AWARE["rows"]), 4, Inches(0.5), Inches(1.6), Inches(9), Inches(3)).table
for j, h in enumerate(AWARE["headers"]):
    gt.cell(0, j).text = h
for i, r in enumerate(AWARE["rows"], 1):
    for j, v in enumerate(r):
        gt.cell(i, j).text = v
CHART = {"title": "Very satisfied by gender (%)", "categories": ["Male", "Female"], "values": [34, 26]}
s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = CHART["title"]
cd = CategoryChartData(); cd.categories = CHART["categories"]; cd.add_series("Very satisfied", CHART["values"])
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(8), Inches(4.5), cd).chart
ch.plots[0].has_data_labels = True
s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = "Awareness (picture of a table)"
s.shapes.add_picture(IMG, Inches(0.5), Inches(2), width=Inches(9))
prs.save(os.path.join(SYN, "doc04_deck.pptx"))

# ---------------------------------------------------------------------------------------------
# Real files
# ---------------------------------------------------------------------------------------------
REAL = "/mnt/user-data/uploads/InsighsAI/TestProjects"
FINDEX_SRC = "/tmp/claude-0/-home-claude/2029d29d-347f-5d3f-9a5d-a57af1db9390/scratchpad/findex-tab"
for n in ("plain.csv", "bytopic.csv"):
    src = os.path.join(FINDEX_SRC, n)
    if os.path.exists(src):
        shutil.copyfile(src, os.path.join(REALDIR, "findex_" + n))
for n in ("survey_degenerate_banner.xlsx", "survey_many_categories.xlsx"):
    shutil.copyfile(os.path.join(REAL, n), os.path.join(REALDIR, n))

def findex_truth(path, label_col, group_cols):
    with open(path, newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    cells = []
    for r in rows:
        v = float(r["Value"])
        base = r["Indicator name"][r["Indicator name"].rfind("(") + 1:-1]
        cells.append({"row": r[label_col], "col": "Value", "value": round(v * 100, 6), "unit": "percent",
                      "stored_as": "proportion", "base_description": base,
                      "groups": [r[g] for g in group_cols]})
    return {"title": None, "label_column": label_col, "group_columns": group_cols,
            "columns": [{"group": "", "label": "Value"}], "rows": len(rows), "base_unweighted": None,
            "cells": cells}

def xlsx_hand_reading(path):
    wb = openpyxl.load_workbook(path, data_only=True); ws = wb.active
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    return {"headers": [str(h) for h in grid[0]], "rows": [dict(zip([str(h) for h in grid[0]], r)) for r in grid[1:]]}

# ---------------------------------------------------------------------------------------------
# cases.json
# ---------------------------------------------------------------------------------------------
def atom_level(row, col, stated):
    return {"kind": "level", "row": row, "col": col, "statedText": stated, "threshold": None, "direction": None}
def atom_cmp(row, a, b, direction="a_higher"):
    return {"kind": "comparison", "row": row, "colA": a, "colB": b, "direction": direction, "claimsSignificant": True}

BANNER_CLAIMS = [
    {"id": "c_level_ok", "text": "34% of men are very satisfied.", "atom": atom_level("Very satisfied", "Male", "34%"), "expected": "supports"},
    {"id": "c_level_wrong", "text": "40% of women are very satisfied.", "atom": atom_level("Very satisfied", "Female", "40%"), "expected": "against"},
    {"id": "c_gap", "text": "Men are significantly more likely than women to be very satisfied.",
     "atom": atom_cmp("Very satisfied", "Male", "Female"), "expected": "supports",
     "why": "34% of 480 against 26% of 520: gap 8 points, se 2.89 (unpooled, tableClaims method), 95% interval 2.3 to 13.7"},
    {"id": "c_gap_wrong", "text": "Women are more likely than men to be very satisfied.",
     "atom": atom_cmp("Very satisfied", "Male", "Female", "b_higher"), "expected": "against"},
    {"id": "c_net", "text": "70% are satisfied overall.", "atom": atom_level("Net: Satisfied", "Total", "70%"), "expected": "supports"},
]

cases = []
def add(**kw):
    cases.append(kw)

add(id="AGG-01", file="synthetic/agg01_frequency.csv", path="A", desc="simple frequency table, CSV, percent text, Total row",
    truth={"tables": [FREQ_TRUTH], "suggestion": "tables"},
    claims=[{"id": "c_level_ok", "text": "60% have heard of Brand X.", "atom": atom_level("Yes", "Percent", "60%"), "expected": "supports"}])
add(id="AGG-02", file="synthetic/agg02_frequency_pctformat.xlsx", path="A", desc="frequency table, Excel, percentages stored as fractions with a 0% format",
    truth={"tables": [FREQ_TRUTH]},
    claims=[{"id": "c_level_ok", "text": "60% have heard of Brand X.", "atom": atom_level("Yes", "Percent", "60%"), "expected": "supports"}])
add(id="AGG-03", file="synthetic/agg03_frequency_textpct.xlsx", path="A", desc="frequency table, Excel, percentages written as text",
    truth={"tables": [FREQ_TRUTH]},
    claims=[{"id": "c_level_ok", "text": "60% have heard of Brand X.", "atom": atom_level("Yes", "Percent", "60%"), "expected": "supports"}])
add(id="AGG-04", file="synthetic/agg04_banner_clean.csv", path="A", desc="banner, one header row, Base row, text percentages, net",
    truth={"tables": [BANNER_TRUTH], "suggestion": "tables"}, claims=BANNER_CLAIMS)
add(id="AGG-05", file="synthetic/agg05_banner_sig.csv", path="A", desc="banner with significance letters attached (34%C)",
    truth={"tables": [BANNER_TRUTH], "suggestion": "tables", "letters_in_file": True}, claims=BANNER_CLAIMS)
add(id="AGG-06", file="synthetic/agg06_banner_titled.xlsx", path="A",
    desc="banner, Excel: title row, merged group header, letters row, weighted base before unweighted, fractions with % format, blank separator column, footnote",
    truth={"tables": [BANNER_TRUTH]}, claims=BANNER_CLAIMS)
add(id="AGG-07", file="synthetic/agg07_stacked.xlsx", path="A", desc="two tables stacked on one sheet, each with title and Base row",
    truth={"tables": [
        {"title": t["title"], "columns": [{"group": "", "label": c} for c in SCOLS],
         "base_unweighted": dict(zip(SCOLS, t["base"])),
         "cells": [{"row": r, "col": c, "value": v, "unit": "percent", "table": i} for r, vals in t["rows"] for c, v in zip(SCOLS, vals)]}
        for i, t in enumerate(STACK)]},
    claims=[{"id": "c_table2", "text": "49% of women would recommend Brand X.", "atom": atom_level("Yes", "Female", "49%"), "expected": "supports",
             "why": "Table 2 row Yes, Female = 49%. Table 1 also has a row Yes (Female 56%)."}])
add(id="AGG-08", file="synthetic/agg08_counts.csv", path="A", desc="counts banner with a Total row, no percentages",
    truth={"tables": [{"title": None, "columns": [{"group": "", "label": c} for c in SCOLS],
                       "base_unweighted": dict(zip(SCOLS, COUNT_TOTAL)),
                       "cells": [{"row": r, "col": c, "value": v, "unit": "count"} for r, vals in COUNTS for c, v in zip(SCOLS, vals)]}],
           "suggestion": "tables"},
    claims=[{"id": "c_pct_from_counts", "text": "34% of men are very satisfied.", "atom": atom_level("Very satisfied", "Male", "34%"),
             "expected": "supports_or_silent", "why": "163 of 480 men = 33.96%, rounds to 34%. A count cell must not be compared with a percentage as if it were one."}])
add(id="AGG-09", file="synthetic/agg09_dash_star.csv", path="A", desc="'-' for none and '*' for under 0.5%",
    truth={"tables": [{"title": None, "columns": [{"group": "", "label": c} for c in SCOLS], "base_unweighted": dict(zip(SCOLS, [1000, 480, 520])),
                       "dash_star": {f"{r}|{c}": list(v) for (r, c), v in DASH_TRUTH_VALUES.items()}}], "suggestion": "tables"},
    claims=[{"id": "c_dash_zero", "text": "No man mentioned Brand D.", "atom": atom_level("Brand D", "Male", "0%"), "expected": "supports_or_silent",
             "why": "'-' means no respondent. Never 'against'."}])
add(id="AGG-10", file="synthetic/agg10_numbers_as_text.xlsx", path="A", desc="counts and percentages stored as Excel text cells",
    truth={"tables": [{**FREQ_TRUTH, "columns": [{"group": "", "label": "Count"}, {"group": "", "label": "Percent"}]}]},
    claims=[{"id": "c_level_ok", "text": "60% have heard of Brand X.", "atom": atom_level("Yes", "Percent", "60"), "expected": "supports"}])
add(id="AGG-11", file="synthetic/agg11_two_sheets.xlsx", path="A", desc="one table per sheet, two sheets",
    truth={"tables": [{"title": "Q1 Awareness", "sheet": "Q1 Awareness"}, {"title": "Q2 Satisfaction", "sheet": "Q2 Satisfaction"}],
           "cells_sheet1": [{"row": r, "col": c, "value": v} for r, vals in STACK[0]["rows"] for c, v in zip(SCOLS, vals)]},
    claims=[{"id": "c_sheet1", "text": "64% of men are aware of Brand X.", "atom": atom_level("Yes", "Male", "64%"), "expected": "supports"}])
add(id="AGG-16", file="synthetic/agg16_banner_export.csv", path="A", desc="banner CSV export: title, group row, header row, letters row, two bases, sig letters, net, footnote",
    truth={"tables": [BANNER_TRUTH], "suggestion": "tables", "letters_in_file": True}, claims=BANNER_CLAIMS)

add(id="AGG-17", file="synthetic/agg17_banner_weighted_first.csv", path="A", desc="banner with the weighted base row above the unweighted base row",
    truth={"tables": [BANNER_TRUTH], "suggestion": "tables"}, claims=BANNER_CLAIMS)
add(id="AGG-18", file="synthetic/agg18_banner_n_in_header.csv", path="A", desc="banner with bases only in the column headers (Male (n=480))",
    truth={"tables": [{**BANNER_TRUTH, "columns": [{**c, "label": c["label"]} for c in COLS]}], "suggestion": "tables", "headers_carry_n": True}, claims=BANNER_CLAIMS)

fp = os.path.join(REALDIR, "findex_plain.csv")
if os.path.exists(fp):
    add(id="AGG-12", file="real/findex_plain.csv", path="A", desc="World Bank Findex indicator table (Indicator name, Value), proportions 0 to 1, base in the label",
        truth={"tables": [findex_truth(fp, "Indicator name", [])], "suggestion": "tables"},
        claims=[{"id": "c_account", "text": "81% of adults have an account.", "atom": atom_level("Account (%, age 15+)", "Value", "81%"), "expected": "supports",
                 "why": "0.8113 is a proportion; the label says %; 81.1% rounds to 81%."}])
    add(id="AGG-13", file="real/findex_bytopic.csv", path="A", desc="Findex indicators with Topic/Sub Topic grouping columns before the label column",
        truth={"tables": [findex_truth(os.path.join(REALDIR, "findex_bytopic.csv"), "Indicator name", ["Topic", "Sub Topic 1", "Sub Topic 2"])], "suggestion": "tables"},
        claims=[{"id": "c_phone", "text": "87% of adults own a mobile phone.", "atom": atom_level("Own a mobile phone (%, age 15+)", "Value", "87%"), "expected": "supports"}])
# World Bank Data360 DHS survey-based indicators (lead, 28 October 2026): SDMX long, wide-by-year, data dictionary.
DHS_SRC = "/mnt/user-data/attachments/2029d29d-347f-5d3f-9a5d-a57af1db9390"
for src, dst in (("4d6745cb-DHS_SD.csv", "dhs_sd.csv"), ("d7110d2d-DHS_SD_WIDEF.csv", "dhs_sd_widef.csv"), ("e69b846a-DHS_SD_DATADICT.csv", "dhs_sd_datadict.csv")):
    if os.path.exists(os.path.join(DHS_SRC, src)):
        shutil.copyfile(os.path.join(DHS_SRC, src), os.path.join(REALDIR, dst))
DIMS = ["REF_AREA", "INDICATOR", "SEX", "AGE", "URBANISATION", "UNIT_MEASURE", "COMP_BREAKDOWN_1", "COMP_BREAKDOWN_2", "COMP_BREAKDOWN_3", "TIME_PERIOD"]
if os.path.exists(os.path.join(REALDIR, "dhs_sd.csv")):
    with open(os.path.join(REALDIR, "dhs_sd.csv"), newline="", encoding="utf-8-sig") as fh:
        rd = list(csv.DictReader(fh))
    const = [k for k in rd[0] if len({x[k] for x in rd}) == 1]
    obs = [{"dims": {d: x[d] for d in DIMS}, "labels": {d: x.get(d + "_LABEL") for d in DIMS if d + "_LABEL" in x}, "value": float(x["OBS_VALUE"])} for x in rd[:5]]
    add(id="AGG-19", file="real/dhs_sd.csv", path="A", desc="Data360 DHS indicators, SDMX long: one observation per country/breakdown/year, value OBS_VALUE in percent",
        truth={"dhs": True, "rows": len(rd), "columns": len(rd[0]), "identity_columns": DIMS, "distinct_identities": len({tuple(x[d] for d in DIMS) for x in rd}),
               "value_column": "OBS_VALUE", "unit": rd[0]["UNIT_MEASURE_LABEL"], "constant_columns": const, "first_observations": obs, "suggestion": "tables"},
        claims=[{"id": "c_second_obs", "text": "In Afghanistan (2015), 42.4% of women received iron during antenatal care (ANCC_IRN_05).",
                 "atom": atom_level("datastructure", "OBS_VALUE", "42.4%"), "expected": "supports_or_silent",
                 "why": "Row 2 of the file is 42.4. The view collapses every row onto the constant first column, so the only row a model can name is 'datastructure', which holds row 1 (42.3)."}])
    with open(os.path.join(REALDIR, "dhs_sd_widef.csv"), newline="", encoding="utf-8-sig") as fh:
        wd = list(csv.DictReader(fh))
    years = [k for k in wd[0] if k.isdigit()]
    filled = [(i, y, x[y]) for i, x in enumerate(wd) for y in years if x[y].strip()]
    add(id="AGG-20", file="real/dhs_sd_widef.csv", path="A", desc="Data360 DHS indicators, wide: one column per year 1999 to 2025, mostly blank",
        truth={"dhs": True, "rows": len(wd), "columns": len(wd[0]), "identity_columns": [d for d in DIMS if d != "TIME_PERIOD"],
               "distinct_identities": len({tuple(x[d] for d in DIMS if d != "TIME_PERIOD") for x in wd}), "value_columns": years,
               "filled_cells": len(filled), "first_filled": [{"row": i, "year": y, "value": float(v)} for i, y, v in filled[:5]],
               "unit": wd[0]["UNIT_MEASURE_LABEL"], "constant_columns": [k for k in wd[0] if not k.isdigit() and len({x[k] for x in wd}) == 1], "suggestion": "tables"}, claims=[])
    with open(os.path.join(REALDIR, "dhs_sd_datadict.csv"), newline="", encoding="utf-8-sig") as fh:
        dd = list(csv.DictReader(fh))
    add(id="CB-01", file="real/dhs_sd_datadict.csv", path="A", desc="Data360 data dictionary of the DHS columns, read as a codebook",
        truth={"codebook": True, "variables": [x["VARIABLE_ID"] for x in dd], "labels": {x["VARIABLE_ID"]: x["VARIABLE_LABEL"] for x in dd},
               "label_column": "VARIABLE_LABEL", "suggestion": "codebook"}, claims=[])

for cid, n in (("AGG-14", "survey_degenerate_banner.xlsx"), ("AGG-15", "survey_many_categories.xlsx")):
    add(id=cid, file="real/" + n, path="A", desc=f"case-level file {n} uploaded as Tables",
        truth={"hand_reading": xlsx_hand_reading(os.path.join(REALDIR, n)), "case_level": True}, claims=[])

add(id="DOC-01", file="synthetic/doc01_report_table.docx", path="BC", desc="Word report: prose page 1, table page 2",
    truth={"pages": 2, "tables": [{**AWARE, "page": 2}], "prose": PROSE, "prose_page": 1})
add(id="DOC-02", file="synthetic/doc02_multipage_table.docx", path="BC", desc="Word report: 70-row table starting on page 5 and continuing",
    truth={"tables": [{**LONG, "page_first": 5}], "pages_min": 6})
add(id="DOC-03", file="synthetic/doc03_tables.pdf", path="BC", desc="PDF: two-level header banner page 1 (Wave 1 Yes/No, Wave 2 Yes/No), simple table page 2",
    truth={"pages": 2, "tables": [{**WAVES, "page": 1}, {**AWARE, "page": 2}]})
add(id="DOC-04", file="synthetic/doc04_deck.pptx", path="BC", desc="PowerPoint: native table slide 2, bar chart with data labels slide 3, picture of a table slide 4",
    truth={"pages": 4, "tables": [{**AWARE, "page": 2}], "chart": {**CHART, "page": 3}, "picture_table": {**AWARE, "page": 4}})
add(id="DOC-05", file="synthetic/doc05_image_only.pdf", path="BC", desc="PDF whose only content is a picture of a table",
    truth={"pages": 1, "picture_table": {**AWARE, "page": 1}})

with open(os.path.join(HERE, "cases.json"), "w", encoding="utf-8") as fh:
    json.dump({"generated_by": "conformance/l4/generate.py", "cases": cases}, fh, indent=1, default=str)
print(f"{len(cases)} cases written")
