"""L4 real-file preparation and independent readings (read-only on the originals).

Writes conformance/l4/work/databook_subset.pptx (slides 10, 15 and 60 of the FinScope 2025 databook: the
whole 652-slide, 28 MB deck was killed for memory by LibreOffice in this sandbox, see REPORT.md) and
conformance/l4/real_expected.json: the expected tables and chart series for chosen slides/pages, read
independently of the app:
  - native PowerPoint tables and chart series: read from the file itself with python-pptx;
  - hand readings (BRC page 17, Tesla slide 13 first chart, databook spot checks): typed from the
    rendered page images by Claude on 28 October 2026 and recorded here.
Run from the worktree root: python3 -I conformance/l4/real_prep.py
"""
import json
import os

from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
WORK = os.path.join(HERE, "work")
os.makedirs(WORK, exist_ok=True)
REAL = "/mnt/user-data/uploads/InsighsAI/TestProjects"
DATABOOK = os.path.join(REAL, "Finscope2025", "FInScope SA 2025 Survey - Databook 16 March 2026.pptx")
TESLA = os.path.join(REAL, "EV Tesla", "QMDA Tesla Survey Presentation Final.pptx")
BRC = os.path.join(REAL, "BRC", "BRC ES Webinar - ASK AFRICA - Still Eating Together 2026.pdf")
KEEP = [10, 15, 60]


def subset(src, dst, keep):
    p = Presentation(src)
    lst = p.slides._sldIdLst
    for i, sid in enumerate(list(lst), 1):
        if i not in keep:
            p.part.drop_rel(sid.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"))
            lst.remove(sid)
    p.save(dst)


def shapes(sh):
    for s in sh:
        if s.shape_type == 6:
            yield from shapes(s.shapes)
        else:
            yield s


def slide_objects(path, slide_no):
    p = Presentation(path)
    s = p.slides[slide_no - 1]
    tables, charts = [], []
    for sh in shapes(s.shapes):
        if getattr(sh, "has_table", False) and sh.has_table:
            tables.append([[c.text.strip() for c in r.cells] for r in sh.table.rows])
        if getattr(sh, "has_chart", False) and sh.has_chart:
            ch = sh.chart
            cats = [str(c) for c in ch.plots[0].categories]
            charts.append({"categories": cats, "series": [{"name": se.name, "values": list(se.values)} for se in ch.plots[0].series]})
    return tables, charts


sub = os.path.join(WORK, "databook_subset.pptx")
if not os.path.exists(sub):
    subset(DATABOOK, sub, KEEP)

expected = {"files": {}}

# Databook subset: page n of the converted PDF is KEEP[n-1].
db = {"source": DATABOOK, "file": "work/databook_subset.pptx", "pages": {}}
for page, orig in enumerate(KEEP, 1):
    t, c = slide_objects(DATABOOK, orig)
    db["pages"][page] = {"original_slide": orig, "native_tables": t, "charts": c}
db["hand_spot_checks"] = {
    "1": ["Race: Black 2025 80, 2024 80, 2023 79, 2022 79", "Gender: Male 48 in all four years",
          "Personal monthly income: No income 2, 3, 3, 2; R2,000 - R2,999 25, 20, 20, 12",
          "Column headers carry the base: 2025 (n=5600), 2022 (n=5604)"],
    "2": ["Chart only: Gauteng 2023 11,491,171 and 2022 12,080,646; Northern Cape 953,555 and 909,578"],
}
expected["files"]["databook"] = db

tesla = {"source": TESLA, "file": TESLA, "pages": {}}
for slide in (13, 16, 18):
    t, c = slide_objects(TESLA, slide)
    tesla["pages"][slide] = {"native_tables": t, "charts": c}
tesla["hand"] = {"13": {"chart": "How likely would you be to buy an electric vehicle within the next 5 years?",
                        "values": {"Unlikely": "37%", "Very unlikely": "30%", "Likely": "23%", "Very likely": "10%"},
                        "claim_on_slide": "33% would buy an EV within 5 years",
                        "claim_derivation": "Likely 23% + Very likely 10% = 33%; the figure is not printed in the chart"}}
expected["files"]["tesla"] = tesla

expected["files"]["brc"] = {"source": BRC, "file": BRC, "hand": {"17": {
    "tables": [
        {"title": "HOW INTERNET & WI-FI CHANGED THE WAY THEY WATCH.",
         "rows": [["Mostly traditional TV, little change", "39%"], ["Now stream more than traditional TV", "20%"],
                  ["Combine TV and streaming to find content", "12%"], ["Mostly online, rarely traditional TV", "9%"]]},
        {"title": "HOW DATA COST & AVAILABILITY AFFECT THEIR VIEWING.",
         "rows": [["Mostly free-to-air, little change", "38%"], ["Both pay-TV & streaming, adjust by budget", "15%"],
                  ["Stream until data runs out, then TV", "14%"], ["Only stream, no traditional TV", "10%"],
                  ["Mostly streaming, some TV by content", "7%"]]}],
    "notes": "Two tables side by side; some rows carry an italic significance note on a second line (*SEM Supergroup H significantly higher)."}},
    "infographic_pages": [10, 15, 18, 22, 23],
    "infographic_note": "Most figures in this deck are infographic tiles (a large percentage and a sentence), not tables."}

with open(os.path.join(HERE, "real_expected.json"), "w", encoding="utf-8") as fh:
    json.dump(expected, fh, indent=1, default=str)
print("written", os.path.join(HERE, "real_expected.json"))
