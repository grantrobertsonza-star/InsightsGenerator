"""Writes a modified copy of a CSM pack workbook for the R6 failure and invariance tests. Only the first sheet (the
data sheet) is changed; the Codebook, Labels and MR sets sheets are kept as they are.

Usage: python3 csm_modify.py <in.xlsx> <out.xlsx> '<json ops>'
ops: {"scaleWeight": 3.7}                      every weight multiplied by the factor
     {"set": [{"row": 1, "column": "journey_task", "value": 2}]}   data row 1 is the first record
     {"setAll": {"journey_recent30": 0}}       the column set in every record
     {"duplicateRow": 1}                       the record appended again at the end
"""
import json
import sys

import openpyxl

src, dst, ops = sys.argv[1], sys.argv[2], json.loads(sys.argv[3])
wb = openpyxl.load_workbook(src)
ws = wb.worksheets[0]
header = [c.value for c in ws[1]]
col = {name: i + 1 for i, name in enumerate(header)}
if "scaleWeight" in ops:
    for r in range(2, ws.max_row + 1):
        c = ws.cell(row=r, column=col["weight"])
        if c.value is not None:
            c.value = c.value * ops["scaleWeight"]
for s in ops.get("set", []):
    ws.cell(row=s["row"] + 1, column=col[s["column"]]).value = s["value"]
for name, value in ops.get("setAll", {}).items():
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=col[name]).value = value
if "duplicateRow" in ops:
    values = [c.value for c in ws[ops["duplicateRow"] + 1]]
    ws.append(values)
wb.save(dst)
